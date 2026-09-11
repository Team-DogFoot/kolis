"""② 자료 검수: 원문 이미지 폴더 검사.

작품(콘텐츠) 폴더 하나 또는 상위 폴더(작품별 하위 폴더) 전체를 검사한다.
- 파일 형식: jpg/jpeg 외는 경고(지침: JPG, JPEG)
- 깨짐: PIL verify + 실제 디코딩
- 컷 중복: 완전 동일(md5) + 유사(dHash 해밍거리 ≤ 임계값)
- 개수·용량·해상도 요약 → 반입용 extent 문구 생성
결과: JSON + 요약 xlsx (문제 셀 노란색)
"""
from __future__ import annotations
import hashlib, json
from dataclasses import dataclass, field, asdict
from pathlib import Path
from PIL import Image, ImageFile
from .common import list_images, IMAGE_EXT_ALLOWED, extent_string

ImageFile.LOAD_TRUNCATED_IMAGES = False


def dhash(img: Image.Image, size: int = 8) -> int:
    g = img.convert("L").resize((size + 1, size), Image.Resampling.LANCZOS)
    px = list(g.getdata())
    bits = 0
    for r in range(size):
        for c in range(size):
            bits = (bits << 1) | (1 if px[r * (size + 1) + c] > px[r * (size + 1) + c + 1] else 0)
    return bits


def hamming(a: int, b: int) -> int:
    return bin(a ^ b).count("1")


@dataclass
class FileResult:
    name: str
    size: int
    ext: str
    ok: bool = True
    width: int | None = None
    height: int | None = None
    mode: str | None = None
    md5: str | None = None
    dhash: int | None = None
    problems: list[str] = field(default_factory=list)


@dataclass
class FolderResult:
    folder: str
    count: int = 0
    total_bytes: int = 0
    extent: str = ""
    color: str = "천연색"
    files: list[FileResult] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    exact_duplicates: list[list[str]] = field(default_factory=list)
    near_duplicates: list[list[str]] = field(default_factory=list)

    @property
    def has_problem(self) -> bool:
        return bool(self.problems) or any(not f.ok for f in self.files)


def inspect_folder(folder: Path, near_threshold: int = 4) -> FolderResult:
    res = FolderResult(folder=str(folder))
    images = list_images(folder)
    others = [p for p in folder.iterdir() if p.is_file() and not p.name.startswith(".") and p not in images]
    if others:
        res.problems.append("이미지 외 파일 존재: " + ", ".join(p.name for p in others[:5]))
    if not images:
        res.problems.append("이미지 파일 없음")
        return res
    md5map: dict[str, list[str]] = {}
    hashes: list[tuple[str, int]] = []
    grayscale_all = True
    for p in images:
        fr = FileResult(name=p.name, size=p.stat().st_size, ext=p.suffix.lower())
        if fr.ext not in IMAGE_EXT_ALLOWED:
            fr.problems.append(f"형식 {fr.ext} (지침: jpg/jpeg)")
        if fr.size == 0:
            fr.ok = False; fr.problems.append("0바이트")
        data = p.read_bytes()
        fr.md5 = hashlib.md5(data).hexdigest()
        md5map.setdefault(fr.md5, []).append(p.name)
        try:
            with Image.open(p) as im:
                im.verify()
            with Image.open(p) as im:
                im.load()
                fr.width, fr.height, fr.mode = im.width, im.height, im.mode
                fr.dhash = dhash(im)
                if im.mode not in ("L", "1"):
                    grayscale_all = False
            if fr.width and fr.height and (fr.width < 100 or fr.height < 100):
                fr.problems.append(f"해상도 작음 {fr.width}x{fr.height}")
        except Exception as e:  # noqa: BLE001
            fr.ok = False
            fr.problems.append(f"디코딩 실패: {type(e).__name__}: {e}")
        if fr.problems and any("디코딩" in x or "0바이트" in x for x in fr.problems):
            fr.ok = False
        res.files.append(fr)
        if fr.dhash is not None:
            hashes.append((p.name, fr.dhash))
    res.count = len(images)
    res.total_bytes = sum(f.size for f in res.files)
    res.color = "단색" if grayscale_all else "천연색"
    res.extent = extent_string(res.count, res.total_bytes, res.color)
    res.exact_duplicates = [names for names in md5map.values() if len(names) > 1]
    # near duplicates (O(n^2) but n은 화당 수십~수백 장)
    seen: set[tuple[str, str]] = set()
    for i in range(len(hashes)):
        for j in range(i + 1, len(hashes)):
            a, b = hashes[i], hashes[j]
            if hamming(a[1], b[1]) <= near_threshold and a[0] != b[0]:
                key = (a[0], b[0])
                if key not in seen:
                    seen.add(key); res.near_duplicates.append([a[0], b[0]])
    exact_names = {n for g in res.exact_duplicates for n in g}
    res.near_duplicates = [pair for pair in res.near_duplicates if not (pair[0] in exact_names and pair[1] in exact_names)]
    if res.exact_duplicates:
        res.problems.append(f"완전 동일 컷 {len(res.exact_duplicates)}묶음")
    if res.near_duplicates:
        res.problems.append(f"유사 컷 {len(res.near_duplicates)}쌍 (사람 확인)")
    bad = [f.name for f in res.files if not f.ok]
    if bad:
        res.problems.append(f"깨진 파일 {len(bad)}개: " + ", ".join(bad[:5]))
    return res


def inspect_root(root: Path, recursive: bool = True) -> list[FolderResult]:
    root = Path(root)
    if list_images(root):
        return [inspect_folder(root)]
    results = []
    for sub in sorted(p for p in root.iterdir() if p.is_dir() and not p.name.startswith(".")):
        if list_images(sub):
            results.append(inspect_folder(sub))
        elif recursive:
            results.extend(inspect_root(sub, recursive))
    return results


def write_reports(results: list[FolderResult], out_dir: Path) -> tuple[Path, Path]:
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    out_dir.mkdir(parents=True, exist_ok=True)
    jpath = out_dir / "inspect.json"
    jpath.write_text(json.dumps([asdict(r) for r in results], ensure_ascii=False, indent=1), encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "요약"
    ws.append(["폴더", "파일 수", "용량(바이트)", "extent 문구", "색상", "문제"])
    yellow = PatternFill("solid", fgColor="FFFF00"); red = Font(color="FF0000")
    for r in results:
        ws.append([r.folder, r.count, r.total_bytes, r.extent, r.color, " / ".join(r.problems)])
        if r.has_problem:
            c = ws.cell(row=ws.max_row, column=6); c.fill = yellow; c.font = red
    ws2 = wb.create_sheet("파일별")
    ws2.append(["폴더", "파일", "바이트", "확장자", "가로", "세로", "모드", "정상", "문제"])
    for r in results:
        for f in r.files:
            ws2.append([r.folder, f.name, f.size, f.ext, f.width, f.height, f.mode, "Y" if f.ok else "N", " / ".join(f.problems)])
            if f.problems:
                c = ws2.cell(row=ws2.max_row, column=9); c.fill = yellow; c.font = red
    ws3 = wb.create_sheet("중복")
    ws3.append(["폴더", "종류", "파일들"])
    for r in results:
        for g in r.exact_duplicates:
            ws3.append([r.folder, "완전동일", ", ".join(g)])
        for g in r.near_duplicates:
            ws3.append([r.folder, "유사(확인)", ", ".join(g)])
    xpath = out_dir / "inspect.xlsx"
    wb.save(xpath)
    return jpath, xpath
