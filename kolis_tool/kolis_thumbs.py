"""7단계: 썸네일 등록 반복 (수정 팝업에서 '다음'으로 넘기며 건별 처리). 유저 정의 절차(2026-09-18), 가이드 [썸네일 등록]과 다름.

건별 절차(모두 상태를 읽어 확인, 어긋나면 그 자리에서 멈춤):
  확인창 없음 → (첫 건 제외) '다음' → 편/권차 읽기 → 반입용 엑셀에서 같은 편/권차 행의 thum_files 로 파일 결정
  → 파일 목록 분석: 열람 CNTS 행 정확히 1개, 나머지는 비정상(자리표시자 등) → 비정상만 체크(픽셀 검증)·원문삭제·저장
  → 열람 행 체크 → 원문등록 팝업: 파일추가('열기' 대화상자) → 전송하기 → 원문유형 select 를 방향키로 '썸네일'(값 검증) → 확인
  → 저장 → 썸네일 행에 등록일 붙었는지 확인.
이미 썸네일이 등록·저장된 건은 건너뛴다(재실행 안전).

기술 메모: 파일 목록 체크박스(jqxGrid)는 UIA 상태가 없어 픽셀로 판정. IE select 는 UIA select() 가 값을 못 바꾸므로 키보드+값 읽기.
알림창이 열려 있으면 뒤 클릭이 모두 무시되므로 단계마다 확인창부터 본다.
"""
from __future__ import annotations
import re, time
from pathlib import Path
from .logutil import UiLog

CHECK_X = 600            # 수정 팝업 파일 목록 체크박스 열 x (팝업 위치 기준으로 보정)
ROW0_Y, ROW_H = 743, 25  # 첫 행 중심 y, 행 높이
UP_CELL_DX, UP_CELL_DY = 152, 355   # 등록 팝업 IE 영역 좌상단 기준 원문유형 셀 위치
TYPE_ORDER = ["열람", "표지", "목차", "초록", "원본", "썸네일"]


class Stop(Exception):
    """사람이 봐야 하는 상황. 화면은 그대로 둔다."""


def _desktop():
    from pywinauto import Desktop
    return Desktop(backend="uia")


def edit_popup():
    from . import kolis_ui as k
    for w in _desktop().windows():
        if w.class_name() == "Alternate Modal Top Most":
            try:
                ie = k.ie_content(w)
                if k._button(ie, "원문삭제") and k._edit(ie, "편/권차"):
                    return w, ie
            except Exception:  # noqa: BLE001
                pass
    return None, None


def upload_ie():
    from . import kolis_ui as k
    for w in _desktop().windows():
        try:
            for ie in w.descendants(class_name="Internet Explorer_Server"):
                if k._button(ie, "파일추가"):
                    return ie
        except Exception:  # noqa: BLE001
            pass
    return None


def dialog_text():
    from . import kolis_ui as k
    d = k.confirm_dialog_now()
    return (d, " ".join(t.window_text() for t in d.descendants(class_name="Static") if t.window_text())) if d else (None, "")


def wait(fn, sec: float, what: str, step: float = 0.3):
    t0 = time.time()
    while time.time() - t0 < sec:
        r = fn()
        if r:
            return r
        time.sleep(step)
    raise Stop(f"기다리다 못 찾음: {what}")


def confirm(expect: str, sec: float = 20) -> str:
    """기대 문구의 확인창을 기다려 '확인'. 다른 문구면 Stop."""
    t0 = time.time()
    while time.time() - t0 < sec:
        d, t = dialog_text()
        if d:
            if expect not in t:
                raise Stop(f"예상과 다른 알림창: '{t[:80]}' (기대 '{expect}')")
            d.child_window(title="확인", class_name="Button").click_input()
            # 같은 제목의 다음 알림창("저장 되었습니다")이 바로 이어질 수 있어 "사라지거나 문구가 바뀜"을 기다린다
            wait(lambda: (lambda dd, tt: dd is None or tt != t)(*dialog_text()), 5, "알림창 닫힘", 0.15)
            return t
        time.sleep(0.2)
    raise Stop(f"알림창이 뜨지 않음(기대 '{expect}')")


def checked(cx: int, cy: int, half: int = 7) -> bool:
    from PIL import ImageGrab
    im = ImageGrab.grab(bbox=(cx - half, cy - half, cx + half, cy + half)).convert("L")
    px = list(im.get_flattened_data()) if hasattr(im, "get_flattened_data") else list(im.getdata())
    return sum(1 for p in px if p < 110) / len(px) > 0.03


def set_check(row: dict, want: bool, x: int, label: str = ""):
    """행의 체크박스: 행 DataItem 의 top 에서 y 를 잡는다(팝업 위치·스크롤 무관)."""
    from pywinauto import mouse
    y = row["top"] + 13
    for _ in range(3):
        if checked(x, y) == want:
            return
        mouse.click(coords=(x, y)); time.sleep(0.5)
    raise Stop(f"{label or row.get('name', '')} 행 체크박스를 {'체크' if want else '해제'} 상태로 만들지 못함")


def file_rows(ie) -> list[dict]:
    """수정 팝업 파일 목록. 열 위치는 IE 영역 좌표 기준 상대값으로 구분."""
    base = ie.rectangle()
    cells = []
    for e in ie.descendants(control_type="DataItem"):
        r = e.rectangle()
        if r.width() > 0 and r.top - base.top > 520:
            cells.append((r.top, r.left - base.left, e.window_text().strip()))
    rows: dict[int, dict] = {}
    for top, dx, txt in cells:
        row = rows.setdefault(round(top / ROW_H), {"top": top})
        if 80 <= dx < 180: row["type"] = txt
        elif 180 <= dx < 340: row["name"] = txt
        elif 340 <= dx < 440: row["size"] = txt
        elif 440 <= dx < 540: row["date"] = txt
        elif dx >= 540: row["path"] = txt
    return [rows[k] for k in sorted(rows) if rows[k].get("name")]


def load_map(import_xlsx: Path) -> dict[str, str]:
    """반입용 엑셀: 편/권차(partNumber 문자열 정규화) → thum_files."""
    import openpyxl
    wb = openpyxl.load_workbook(import_xlsx, read_only=True)
    ws = wb["Contents"]
    hdr = [str(c or "") for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    pi, ti = hdr.index("/mods/titleInfo/partNumber"), hdr.index("thum_files")
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[pi] is not None and r[ti]:
            out[_norm_part(r[pi])] = str(r[ti]).strip()
    return out


def _norm_part(v) -> str:
    return re.sub(r"\s+", "", str(v or "")).strip()


def thumb_for(part_text: str, mapping: dict[str, str], thumb_dir: Path, log) -> Path:
    key = _norm_part(part_text)
    name = mapping.get(key)
    if not name:
        m = re.match(r"\s*(\d+)", part_text or "")
        if not m:
            raise Stop(f"편/권차 '{part_text}' 에 맞는 썸네일을 엑셀에서 찾지 못함(앞 숫자도 없음)")
        num = int(m.group(1))
        cands = [v for k, v in mapping.items() if re.match(rf"{num}(\D|$)", k)]
        if len(cands) != 1:
            raise Stop(f"편/권차 '{part_text}' 에 맞는 썸네일을 정하지 못함(후보 {cands})")
        name = cands[0]
        log.warn(f"편/권차 '{part_text}' 는 엑셀과 정확히 일치하지 않아 앞 숫자 {num} 로 '{name}' 선택")
    f = Path(thumb_dir) / name
    if not f.exists():
        raise Stop(f"썸네일 파일 없음: {f}")
    return f


def select_type_thumbnail(cell_xy, log):
    """등록 팝업 표의 원문유형 select 를 '썸네일'로. 옵션 순서가 고정이라 DOWN 5번 후 값 검증, 안 맞으면 한 칸씩 보정."""
    from pywinauto import mouse, keyboard
    def combo_val():
        ie = upload_ie()
        c = [x.window_text() for x in ie.descendants(control_type="ComboBox") if x.rectangle().top - ie.rectangle().top > 330]
        return c[0] if c else None
    def cell_val():
        ie = upload_ie(); b = ie.rectangle()
        return [e.window_text().strip() for e in ie.descendants(control_type="DataItem")
                if 335 <= e.rectangle().top - b.top <= 375 and 90 <= e.rectangle().left - b.left <= 135]
    mouse.click(coords=cell_xy); time.sleep(0.4)
    v = wait(combo_val, 5, "원문유형 select")
    idx = TYPE_ORDER.index(v) if v in TYPE_ORDER else 0
    keyboard.send_keys("{DOWN}" * (len(TYPE_ORDER) - 1 - idx), pause=0.05); time.sleep(0.3)
    for _ in range(6):
        v = combo_val()
        if v == "썸네일":
            break
        keyboard.send_keys("{DOWN}" if (v in TYPE_ORDER and TYPE_ORDER.index(v) < 5) else "{UP}"); time.sleep(0.25)
    if combo_val() != "썸네일":
        raise Stop(f"원문유형 select 값이 '썸네일'이 아님: {combo_val()}")
    keyboard.send_keys("{ENTER}"); time.sleep(0.3)
    mouse.click(coords=(cell_xy[0] + 180, cell_xy[1] + 90)); time.sleep(0.5)   # 셀 밖 클릭으로 확정
    if "썸네일" not in cell_val():
        raise Stop(f"원문유형 셀 확정값이 '썸네일'이 아님: {cell_val()}")
    log("원문유형 = 썸네일 확인")


def process_one(mapping: dict, thumb_dir: Path, log: UiLog, go_next: bool, expect_part: str | None = None) -> dict:
    from . import kolis_ui as k
    from pywinauto import mouse
    d, t = dialog_text()
    if d:
        raise Stop(f"시작 전 알림창이 열려 있음: '{t[:80]}'")
    w, ie = edit_popup()
    if not ie:
        raise Stop("수정 팝업이 없습니다(납본자료접수에서 전체 선택 → 수정 으로 여세요)")
    if go_next:
        k._button(ie, "다음").click_input()
        prev = expect_part
        wait(lambda: _norm_part(k._value(k._edit(edit_popup()[1], "편/권차"))) not in ("", _norm_part(prev)), 15, "'다음' 후 편/권차 변경", 0.3)
        d, t = dialog_text()
        if d:
            raise Stop(f"'다음' 후 알림창: '{t[:80]}'")
    w, ie = edit_popup()
    base = ie.rectangle()
    x_check = base.left + 82      # 실측: IE 영역 좌측 + 82 = 체크박스 열 중심
    part = k._value(k._edit(ie, "편/권차")).strip()
    thumb = thumb_for(part, mapping, thumb_dir, log)
    rows = file_rows(ie)
    log(f"편/권차 {part!r} → {thumb.name} | 목록 {[(r.get('type'), r.get('name', '')[:24]) for r in rows]}")
    view = [i for i, r in enumerate(rows) if r.get("type") == "열람" and r.get("name", "").startswith("CNTS-")]
    if len(view) != 1:
        raise Stop(f"열람 CNTS 행이 {len(view)}개 — 사람이 확인")
    done = [r for r in rows if r.get("type") == "썸네일" and r.get("name") == thumb.name and r.get("date")]
    if done:
        log("이미 썸네일 등록·저장됨 → 건너뜀")
        return {"part": part, "status": "already", "thumb": thumb.name}
    bad = [i for i in range(len(rows)) if i not in view]
    if bad:
        for i in range(len(rows)):
            set_check(rows[i], i in bad, x_check)
        k._button(edit_popup()[1], "원문삭제").click_input()
        confirm("삭제하시겠습니까")
        d2, t2 = dialog_text()
        if d2:
            d2.child_window(title="확인", class_name="Button").click_input(); time.sleep(0.3)
        k._button(edit_popup()[1], "저장").click_input()
        confirm("저장하시겠습니까"); confirm("저장 되었습니다")
        rows = wait(lambda: (lambda rr: rr if len(rr) == 1 else None)(file_rows(edit_popup()[1])), 10, "삭제 후 목록 1행")
        log(f"비정상 {len(bad)}행 삭제·저장")
    set_check(rows[0], True, x_check, "열람")
    # 원문등록
    k._button(edit_popup()[1], "원문등록").click_input()
    uie = wait(upload_ie, 15, "등록 팝업")
    k._button(uie, "파일추가").click_input()
    k.select_file_dialog(thumb, k.Log(lambda m: None), title="열기")
    wait(lambda: any(thumb.name in t.window_text() for t in upload_ie().descendants(control_type="Text")), 8, "파일 추가 확인")
    k._button(upload_ie(), "전송하기").click_input()
    def row():
        r = [e.window_text().strip() for e in upload_ie().descendants(control_type="DataItem") if e.window_text().strip()]
        return r if r else None
    r = wait(row, 60, "전송 결과 행")
    if thumb.name not in r:
        raise Stop(f"전송 결과에 {thumb.name} 없음: {r}")
    ub = upload_ie().rectangle()
    select_type_thumbnail((ub.left + UP_CELL_DX, ub.top + UP_CELL_DY), log)
    k._button(upload_ie(), "확인").click_input()
    confirm("등록 하시겠습니까")
    wait(lambda: any(r.get("type") == "썸네일" and r.get("name") == thumb.name for r in file_rows(edit_popup()[1])), 10, "목록에 썸네일 행")
    k._button(edit_popup()[1], "저장").click_input()
    confirm("저장하시겠습니까"); confirm("저장 되었습니다")
    final = wait(lambda: (lambda rr: rr if any(x.get("type") == "썸네일" and x.get("date") for x in rr) else None)(file_rows(edit_popup()[1])), 10, "저장 후 썸네일 행 등록일")
    log(f"완료 편/권차 {part}: {[(x.get('type'), x.get('name'), x.get('date')) for x in final]}")
    return {"part": part, "status": "done", "thumb": thumb.name}


def run(import_xlsx: Path, thumb_dir: Path, count: int, log: UiLog, handle: dict, progress=None) -> dict:
    """현재 열린 수정 팝업부터 count 건(0=끝까지). handle['cancel'] 로 중단. progress(i, result) 콜백."""
    mapping = load_map(import_xlsx)
    log(f"썸네일 등록 시작: 엑셀 행 {len(mapping)}, 폴더 {thumb_dir}, 목표 {count or '끝까지'}건")
    results = []; i = 0; last_part = None
    t0 = time.time()
    while count == 0 or i < count:
        if handle.get("cancel"):
            log("중단 요청으로 멈춤"); break
        t1 = time.time()
        try:
            r = process_one(mapping, thumb_dir, log, go_next=i > 0, expect_part=last_part)
        except Stop as e:
            log.error(f"멈춤: {e}")
            results.append({"status": "stop", "error": str(e)})
            if progress: progress(i, results[-1])
            break
        r["sec"] = round(time.time() - t1, 1); results.append(r); last_part = r["part"]
        if progress: progress(i, r)
        i += 1
        if last_part and _norm_part(last_part).startswith("1") and re.match(r"1(\D|$)", _norm_part(last_part)):
            log("편/권차 1 처리 → 마지막으로 보고 종료"); break
    done = sum(1 for r in results if r["status"] == "done"); skip = sum(1 for r in results if r["status"] == "already")
    log(f"썸네일 등록 종료: 완료 {done}, 건너뜀 {skip}, 멈춤 {sum(1 for r in results if r['status']=='stop')} ({int(time.time()-t0)}초)")
    return {"done": done, "skipped": skip, "results": results, "stopped": any(r["status"] == "stop" for r in results)}
