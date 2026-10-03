"""반입용 엑셀(83열) 쓰기와 검사 — 형태가 정해진 부분이라 코드가 한다.

칸에 들어갈 **값은 에이전트가 정한다**(`import.json`). 이 파일이 하는 일:
  - 값을 양식의 정해진 열에 옮겨 쓴다(열 위치는 템플릿 1행에서 읽는다).
  - 웹툰 사업 고정값을 채운다(CONSTANTS: 양식과 완료 사례에서 항상 같은 값).
  - 이미지 수·용량을 실제 폴더에서 세어 수량(extent) 문구를 만든다.
  - 에이전트가 "사람이 확인할 칸"으로 지정한 칸을 노란색으로 칠하고 이유를 메모로 단다.
  - 검사한다: 파일 증거와 맞는가(행 수 = 폴더 수, 폴더·썸네일 파일이 실제로 있는가), 형식이 맞는가(날짜 8자리, 필수 칸).
    판단은 검사하지 않는다.
시트는 양식에 있는 Sample, Contents 둘뿐이다. 다른 시트를 덧붙이지 않는다(유저 확정 2026-09-29).
"""
from __future__ import annotations
import json, re, shutil
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from . import arrange
from .common import REGION_CODE, extent_string

HERE = Path(__file__).parent
TEMPLATE_83 = HERE / "templates" / "import_template_83.xlsx"
YELLOW = PatternFill("solid", fgColor="FFFF00")
Key = tuple[str, int]      # (열 이름, 같은 이름 중 몇 번째)

# 웹툰 사업 고정값(예시 양식·완료 사례에서 동일하게 확인된 값)
CONSTANTS: dict[Key, str] = {
    ("/mods/originInfo[@eventType]", 0): "publication",
    ("/mods/originInfo/place/placeTerm[@type]", 0): "text",
    ("/mods/originInfo/place/placeTerm[@authority]", 0): "kormarccountry",
    ("/mods/originInfo/place/placeTerm[@type]", 1): "code",
    ("/mods/originInfo/issuance", 0): "단행자료",
    ("/mods/language/languageTerm", 0): "kor",
    ("/mods/language/languageTerm[@authority]", 0): "iso639-2b",
    ("/mods/language/languageTerm[@type]", 0): "code",
    ("/mods/physicalDescription/form", 0): "전자자료(Image)",
    ("/mods/physicalDescription/reformattingQuality", 0): "access",
    ("/mods/physicalDescription/digitalOrigin", 0): "BornDigital",
    ("/mods/note[@type]", 1): "target audience",
    ("/mods/note[@type]", 2): "acquisition",
    ("/mods/subject/topic", 0): "만화",
    ("/mods/classification", 0): 810,
    ("/mods/classification[@authority]", 0): "KDC",
    ("/mods/classification[@edition]", 0): "6",
    ("/mods/location/physicalLocation", 0): "국립중앙도서관",
    ("/mods/extension/regionOfPublishing", 0): "한국",
    ("/mods/typeOfResource", 0): "텍스트",
    ("/mods/genre", 0): "만화",
    ("currency_code", 0): "\\",
}

# 에이전트가 쓰는 항목 이름 → 양식의 열. 완료 사례(83열): 주기 1번째=기타, 2번째=이용대상, 3번째=수집처 / 원문주소 1번째=작품이 들어 있는 목록 페이지, 2번째=작품의 회차 목록 페이지
FIELDS: dict[str, Key] = {
    "title": ("/mods/titleInfo/title", 0), "title_parallel": ("/mods/titleInfo/title", 1), "subTitle": ("/mods/titleInfo/subTitle", 0),
    "partNumber": ("/mods/titleInfo/partNumber", 0), "partName": ("/mods/titleInfo/partName", 0),
    "place": ("/mods/originInfo/place/placeTerm", 0), "place_code": ("/mods/originInfo/place/placeTerm", 1),
    "publisher": ("/mods/originInfo/publisher", 0), "dateIssued": ("/mods/originInfo/dateIssued", 0),
    "targetAudience": ("/mods/targetAudience", 0), "other_note": ("/mods/note", 0), "audience_note": ("/mods/note", 1),
    "acquisition_note": ("/mods/note", 2), "identifier": ("/mods/identifier", 0), "identifier_type": ("/mods/identifier[@type]", 0),
    "url_main": ("/mods/location/url", 0), "url_work": ("/mods/location/url", 1), "licenseType": ("/mods/accessCondition/licenseType", 0),
    "price": ("contents_price", 0), "compensation": ("compensation", 0), "reward_yn": ("reward_yn", 0),
    "edition": ("/mods/originInfo/edition", 0),
}
# 반복·묶음 항목(값이 있으면 양식에 열을 그만큼 늘려 쓴다 — 완료 사례 88열: 저자 4묶음, 발행처 2칸)
#   names: [{name,type,role}] 상한 없음 / publishers: 첫 출처정보의 추가 발행처(제작처·유통사, 가이드 5 순서) / notes_extra: [{text,type}] 넷째 주기부터(수상 등)
#   origin2: {"type","issuance","publisher","dateIssued"} 둘째 출처정보(발행지가 다른 제작처가 있을 때만, 가이드 5.2)
REPEAT_FIELDS = ("names", "publishers", "notes_extra", "origin2")
ORIGIN2_KEYS = {"type": ("/mods/originInfo[@type]", 0), "issuance": ("/mods/originInfo/issuance", 1), "dateIssued": ("/mods/originInfo/dateIssued", 1)}   # publisher 는 첫 출처정보의 발행처 수에 따라 번째가 달라져 write 에서 계산
NAME_GROUP = ["/mods/name", "/mods/name/namePart", "/mods/name[@type]", "/mods/name/role", "/mods/name/role/roleTerm"]     # 둘째 이후 저자 묶음의 열(예시 88열)


def name_keys(i: int) -> dict:
    return {"name": ("/mods/name/namePart", i), "type": ("/mods/name[@type]", i), "role": ("/mods/name/role/roleTerm", i)}


NAME_KEYS = [name_keys(i) for i in range(3)]
COMPUTED = {"extent": ("/mods/physicalDescription/extent", 0), "format": ("/mods/physicalDescription/internetMediaType", 0), "thumb_file": ("thum_files", 0)}
REQUIRED = ["title", "publisher", "place", "place_code", "dateIssued", "targetAudience", "audience_note", "acquisition_note", "identifier",
            "url_main", "url_work", "licenseType", "price", "reward_yn"]
AUDIENCES = ("고등학생", "성인용", "아동용", "일반이용자", "중학생", "초등학생", "취학전아동", "취학전 아동", "특수계층", "미상")


NUMERIC = ("price", "compensation", "dateIssued", "identifier", "licenseType")      # 완료 사례는 이 칸들을 숫자로 저장한다(글자로 넣으면 엑셀이 "텍스트 형식 숫자"로 표시)
ISBN_FORMAT = r"0_);[Red]\(0\)"          # 완료 사례의 식별기호 칸 서식(13자리가 지수 표기로 바뀌지 않게)

KOREAN = {"title": "본표제", "title_parallel": "대등표제", "subTitle": "표제관련정보", "partNumber": "권차", "partName": "권차표제", "place": "발행지", "place_code": "발행국 부호",
          "publisher": "발행처", "dateIssued": "발행일", "targetAudience": "이용대상자", "other_note": "주기", "audience_note": "이용대상자 주기",
          "acquisition_note": "입수처 주기", "identifier": "식별기호", "identifier_type": "식별기호 유형", "url_main": "원문주소(작품이 들어 있는 목록 페이지)",
          "url_work": "원문주소(작품의 회차 목록 페이지)", "licenseType": "접근제한유형", "price": "정가", "compensation": "보상금", "reward_yn": "보상여부",
          "names": "저자명", "color": "수량(크기)의 색", "thumb_file": "썸네일 파일명", "extent": "수량(크기)", "format": "디지털 자료유형",
          "edition": "판사항", "publishers": "발행처(둘째 이후)", "notes_extra": "주기(넷째 이후)", "origin2": "둘째 출처정보"}
ROLE_BAD = ("지은", "그린", "엮은", "옮긴", "쓴", "펴낸")      # 관형형 역할어(가이드 2.3: 명사형으로)


def column_label(name: str) -> str:
    """사람에게 보여 줄 칸 이름: 반입용 엑셀의 열 이름(MODS) 기준, 한글은 괄호 안. 같은 열 이름이 여러 번 나오면 몇 번째인지 붙인다."""
    key = FIELDS.get(name) or COMPUTED.get(name) or {"names": NAME_KEYS[0]["name"], "color": COMPUTED["extent"], "publishers": ("/mods/originInfo/publisher", 1), "notes_extra": ("/mods/note", 3), "origin2": ("/mods/originInfo[@type]", 0)}.get(name)
    if not key:
        return name
    repeated = sum(1 for k in list(FIELDS.values()) + list(COMPUTED.values()) + [n["name"] for n in NAME_KEYS] if k[0] == key[0]) > 1 or name in ("publishers", "notes_extra")
    head = key[0] + (f" #{key[1] + 1}" if repeated else "")
    return f"{head} ({KOREAN.get(name, name)})"


def _columns(ws) -> dict[Key, int]:
    header = [(c.value or "").strip() if isinstance(c.value, str) else "" for c in ws[1]]
    out, seen = {}, {}
    for i, h in enumerate(header):
        out[(h, seen.get(h, 0))] = i + 1
        seen[h] = seen.get(h, 0) + 1
    return out


def field_names() -> list[str]:
    return [*FIELDS, *REPEAT_FIELDS, "color", "thumb_file"]


def check(data: dict, folder: Path) -> list[str]:
    """import.json 을 파일 증거·형식과 대조한다. 걸린 항목 목록(비어 있으면 통과)."""
    out = []
    rows = data.get("rows") or []
    if not rows:
        return ["rows 가 비어 있습니다"]
    # 정리 계획(arrange)이 있으면 "옮긴 뒤의 모습"으로 검사한다. 실제로 옮기는 것은 프로그램의 마무리 단계다
    plan_fails = arrange.check(folder, data)
    if plan_fails:
        return plan_fails
    vw = arrange.View(folder, data)
    root = str(data.get("manuscripts_root") or "")
    if not root or not vw.is_dir(root):
        return [f"원고 폴더(manuscripts_root)가 없습니다: {data.get('manuscripts_root')!r}. 회차 폴더만 들어 있는 상위 폴더여야 합니다"
                "(지금 그런 폴더가 없으면 arrange.moves 로 회차 폴더를 한 폴더 아래로 모으세요)"]
    kids = vw.children(root)
    subs = {n: p for n, p in kids if p.is_dir() and not n.startswith("_kolis")}
    if len(rows) != len(subs):
        out.append(f"자료 {len(rows)}행 ≠ 원고 하위 폴더 {len(subs)}개 {sorted(subs)[:15]}")
    stray = [n for n, p in kids if p.is_file()]
    if stray:
        out.append(f"원고 폴더 바로 아래에 파일이 있습니다(원문일괄등록이 멈춥니다): {stray[:3]}")
    th = str(data.get("thumbs_dir") or "")
    names = {n for n, p in vw.children(th) if p.is_file()} if th and vw.is_dir(th) else None
    if th and names is None:
        out.append(f"썸네일 폴더가 없습니다: {th}")
    used, parts, idents = set(), [], []
    for i, r in enumerate(rows, 1):
        v = r.get("values") or {}
        confirm = {str(c.get("field")) for c in r.get("confirm") or [] if str(c.get("reason") or "").strip()}
        for c in r.get("confirm") or []:
            if str(c.get("field")) not in field_names():
                out.append(f"{i}행 confirm: '{c.get('field')}' 는 없는 항목 이름입니다. 쓸 수 있는 이름: {', '.join(field_names())}")
            if not str(c.get("reason") or "").strip():
                out.append(f"{i}행 confirm '{c.get('field')}': 이유가 없습니다")
            if not str(c.get("ask") or "").strip():
                out.append(f"{i}행 confirm '{c.get('field')}': ask(직원이 무엇을 정해야 하는지)가 없습니다. 이 작품을 처음 보는 직원이 그 문장만 읽고 답할 수 있게 적으세요")
        unknown = [k for k in v if k not in field_names()]
        if unknown:
            out.append(f"{i}행 values: 없는 항목 이름 {unknown}. 쓸 수 있는 이름: {', '.join(field_names())}")
        f = str(r.get("folder") or "")
        if f not in subs:
            out.append(f"{i}행: folder '{f}' 에 해당하는 원고 하위 폴더가 없습니다")
        elif f in used:
            out.append(f"{i}행: 폴더 '{f}' 가 두 번 쓰였습니다")
        else:
            used.add(f)
            if not vw.images(f"{root}/{f}"):
                out.append(f"{i}행: 폴더 '{f}' 에 이미지가 없습니다")
            extra = vw.others(f"{root}/{f}")
            if extra:
                out.append(f"{i}행: 폴더 '{f}' 에 이미지가 아닌 것이 있습니다(원문일괄등록이 멈춥니다): {extra[:3]}. arrange.set_aside 로 빼 두세요")
        for k in REQUIRED:
            if str(v.get(k) if v.get(k) is not None else "").strip() == "" and k not in confirm:
                out.append(f"{i}행: '{k}' 가 비어 있습니다. 값을 넣거나, 정할 수 없으면 confirm 에 이유와 함께 넣으세요")
        ns = [n for n in v.get("names") or [] if str(n.get("name") or "").strip()]
        if not ns and "names" not in confirm:
            out.append(f"{i}행: 저자(names)가 없습니다")
        src = r.get("from") or {}
        for n in ns:
            if n.get("type") not in ("개인명", "단체명"):
                out.append(f"{i}행 저자 '{n.get('name')}': type 은 '개인명' 또는 '단체명'")
            if re.search(r"[\[\]]", str(n.get("name") or "")) and "플랫폼" not in str(src.get("names") or ""):
                out.append(f"{i}행 저자 '{n.get('name')}': 각괄호는 플랫폼 화면에서만 확인한 저자에만 씁니다(가이드 2 채택 순서). 원고·출판사 엑셀에 있는 이름이면 각괄호를 빼고, 플랫폼에서만 봤으면 from.names 에 '플랫폼'을 적으세요")
            role = str(n.get("role") or "").strip()
            if role and any(role == b or role.endswith(b) for b in ROLE_BAD):
                out.append(f"{i}행 저자 '{n.get('name')}': 역할어 '{role}' 는 관형형입니다. 명사형으로 적습니다(가이드 2.3, 예: 지은 → 지음)")
        for k, pubs in (("publishers", v.get("publishers") or []),):
            if not isinstance(pubs, list) or any(not isinstance(x, str) for x in pubs):
                out.append(f"{i}행: {k} 는 문자열 목록이어야 합니다")
            elif pubs and not str(v.get("other_note") or "").strip() and "other_note" not in confirm:
                out.append(f"{i}행: 발행처가 둘 이상({v.get('publisher')}, {', '.join(pubs)})인데 관계를 설명하는 일반 주기(other_note)가 없습니다. 가이드 9: 임프린트·브랜드 관계는 일반 주기에 적습니다(예: 'A은 B의 임프린트임'). 모르면 other_note 를 confirm 에 넣으세요")
        pub = str(v.get("publisher") or "").strip()
        if re.fullmatch(r"[A-Z][A-Z0-9&\s]{3,}", pub) and "publisher" not in confirm:
            out.append(f"{i}행: publisher '{pub}' 가 전부 대문자입니다. 가이드 5.2: 머리글자 약어가 아니면 각 단어 첫 글자만 대문자로 정규화합니다(예: KWBOOKS → Kwbooks). 약어가 맞으면 confirm 에 이유를 적으세요")
        um = str(v.get("url_main") or "").strip()
        if um and re.fullmatch(r"https?://[^/]+/?", um) and "url_main" not in confirm:
            out.append(f"{i}행: url_main '{um}' 은 사이트 첫 화면입니다. 가이드 14.1·3절: 그 작품이 들어 있는 목록 페이지(예 https://www.mrblue.com/comic, https://ridibooks.com/comics/ebook)를 적으세요")
        pn = str(v.get("partNumber") or "").strip()
        if pn.isdigit() and "partNumber" not in confirm and "원고" not in str(src.get("partNumber") or ""):
            out.append(f"{i}행: partNumber '{pn}' 는 단위 없는 숫자입니다. 가이드 1.3: 원고 표기 그대로('1화', '01회', '1권' …). 원고에 정말 숫자만 있으면 from.partNumber 에 '원고'를 적고, 아니면 표기를 맞추거나 confirm 에 넣으세요")
        o2 = v.get("origin2") or {}
        if o2 and not isinstance(o2, dict):
            out.append(f"{i}행: origin2 는 {{type, issuance, publisher, dateIssued}} 꼴이어야 합니다")
        elif o2 and str(o2.get("dateIssued") or "") and not re.fullmatch(r"\d{8}|\d{4}-{4}|\d{6}-{2}", str(o2.get("dateIssued"))):
            out.append(f"{i}행: origin2.dateIssued '{o2.get('dateIssued')}' 는 YYYYMMDD(모르는 자리는 -)")
        for x in v.get("notes_extra") or []:
            if not isinstance(x, dict) or not str(x.get("text") or "").strip():
                out.append(f"{i}행: notes_extra 항목은 {{text, type}} 꼴이고 text 가 있어야 합니다")
        d = str(v.get("dateIssued") or "")
        if d and not re.fullmatch(r"\d{8}", d):
            out.append(f"{i}행: dateIssued '{d}' 는 숫자 8자리(YYYYMMDD)여야 합니다")
        if v.get("color") not in ("천연색", "흑백"):
            out.append(f"{i}행: color 는 '천연색' 또는 '흑백'(매뉴얼 7.4)")
        if str(v.get("targetAudience") or "") and v.get("targetAudience") not in AUDIENCES:
            out.append(f"{i}행: targetAudience '{v.get('targetAudience')}' 는 양식의 선택값이 아닙니다({', '.join(AUDIENCES[:4])} …)")
        code = str(v.get("place_code") or "")
        if code and code not in REGION_CODE.values():
            out.append(f"{i}행: place_code '{code}' 는 발행국 부호 표에 없습니다({', '.join(f'{k} {c}' for k, c in list(REGION_CODE.items())[:4])} …)")
        if str(v.get("price") or "").strip() and not re.fullmatch(r"\d+", str(v.get("price")).strip()):
            out.append(f"{i}행: price '{v.get('price')}' 는 숫자만")
        if str(v.get("reward_yn") or "") not in ("", "Y", "N"):
            out.append(f"{i}행: reward_yn 은 Y 또는 N")
        raw_ident = str(v.get("identifier") or "").strip()
        ident = re.sub(r"[\s-]", "", raw_ident)
        is_isbn = str(v.get("identifier_type") or "").lower() == "isbn"
        if is_isbn and raw_ident and not raw_ident.isdigit():
            out.append(f"{i}행: ISBN '{raw_ident}' 는 붙임표(-)·공백 없이 숫자만 적습니다")
        if is_isbn and ident and not _isbn13_ok(ident) and "identifier" not in confirm:
            out.append(f"{i}행: ISBN '{v.get('identifier')}' 의 검증 숫자가 맞지 않습니다. 확인하거나 confirm 에 넣으세요")
        if is_isbn and ident:
            idents.append((ident, "identifier" in confirm))
        # 직원 규칙(2026-10-01)의 형식: 발행처 각괄호 금지, 보상금 = 정가, 유료는 Y / 무료(정가 0)는 N
        if re.search(r"[\[\]]", str(v.get("publisher") or "")):
            out.append(f"{i}행: publisher '{v.get('publisher')}' 에 각괄호가 있습니다. 발행처에는 각괄호를 넣지 않습니다")
        price, comp, yn = (str(v.get(k) if v.get(k) is not None else "").strip() for k in ("price", "compensation", "reward_yn"))
        if price.isdigit():
            if comp.isdigit() and int(comp) != int(price):
                out.append(f"{i}행: compensation {comp} ≠ price {price}. 보상금은 정가와 같습니다")
            if int(price) > 0 and comp.isdigit() and int(comp) > 0 and yn != "Y":
                out.append(f"{i}행: 정가와 보상금이 둘 다 있으면 reward_yn 은 Y")
            if int(price) > 0 and not comp and "compensation" not in confirm:
                out.append(f"{i}행: 유료(정가 {price})인데 compensation 이 비어 있습니다. 보상금 = 정가")
            if int(price) == 0 and yn != "N":
                out.append(f"{i}행: 무료 회차(정가 0)의 reward_yn 은 N")
        t = str(r.get("thumb_file") or "")
        if names is None and t:
            out.append(f"{i}행: 썸네일 폴더가 없는 납품인데 thumb_file 에 값이 있습니다(반입하면 0 Bytes 자리표시자 행이 생깁니다). 비우세요")
        src = str(r.get("thumb_source") or "")
        if names is not None and t and t not in names and src not in names:
            out.append(f"{i}행: 썸네일 파일이 썸네일 폴더에 없습니다(thumb_source '{src}', thumb_file '{t}')")
        if names is not None and not t and "thumb_file" not in confirm:
            out.append(f"{i}행: 썸네일 폴더가 있는데 thumb_file 이 비어 있습니다")
        parts.append((str(v.get("partNumber") or ""), str(v.get("partName") or "")))
    dup = sorted({p for p in parts if any(p) and parts.count(p) > 1})
    if dup:
        out.append(f"권차·권차표제가 같은 행이 있습니다: {dup[:5]}")
    shared = sorted({x for x, _ in idents if sum(1 for y, _ in idents if y == x) > 1})
    if shared and not all(c for x, c in idents if x in shared):
        out.append(f"같은 ISBN 이 여러 행에 쓰였습니다: {shared[:3]}. 권·회차별 ISBN 을 찾아 넣거나, 끝내 세트 ISBN 뿐이면 그 행들 전부의 confirm 에 identifier 로 이유(어디를 찾아봤는지)를 적으세요")
    if not out:
        rv = data.get("review") or {}
        if not rv.get("done"):
            out.append("검수를 아직 받지 않았습니다: reviewer 에이전트에게 맡기고, 지적을 처리한 뒤 review.done 을 true 로 하세요")
        elif len(rv.get("findings") or []) != len(rv.get("resolved") or []):
            out.append(f"검수 지적 {len(rv.get('findings') or [])}건 중 처리 기록이 {len(rv.get('resolved') or [])}건입니다")
    return out


def confirm_text(c: dict) -> str:
    """확인할 칸 하나를 사람이 읽을 글로(엑셀 메모에 넣는다)."""
    parts = [str(c.get("reason") or "").strip()]
    for label, k in (("출판사 엑셀", "publisher_says"), ("근거", "evidence"), ("직원이 정할 것", "ask")):
        if str(c.get(k) or "").strip():
            parts.append(f"{label}: {str(c[k]).strip()}")
    return "\n".join(p for p in parts if p)


def _isbn13_ok(s: str) -> bool:
    if not re.fullmatch(r"\d{13}", s):
        return False
    total = sum(int(d) * (1 if i % 2 == 0 else 3) for i, d in enumerate(s[:12]))
    return (10 - total % 10) % 10 == int(s[12])


def _header(ws) -> list[str]:
    return [(c.value or "").strip() if isinstance(c.value, str) else "" for c in ws[1]]


def _insert_after(ws, after_idx: int, names: list[str]) -> None:
    """1행 머리글 기준 after_idx(1부터) 열 뒤에 열을 끼워 넣고 머리글을 쓴다."""
    ws.insert_cols(after_idx + 1, len(names))
    for k, n in enumerate(names):
        ws.cell(row=1, column=after_idx + 1 + k).value = n


def _grow_columns(ws, data: dict) -> dict:
    """값이 요구하는 만큼 양식 열을 늘린다(완료 사례 88열의 자리와 순서). 돌려주는 것: 반복 수 {names, publishers, notes}.
    - 저자 묶음: 템플릿 3묶음 → 필요한 수만큼, 마지막 저자 묶음 뒤에 [/mods/name, namePart, [@type], role, roleTerm] 을 반복
    - 첫 출처정보 발행처: 첫 /mods/originInfo/publisher 바로 뒤에 반복
    - 주기: 셋째 쌍 뒤에 [/mods/note, /mods/note[@type]] 반복
    - 주제명: 템플릿의 둘째 /mods/subject/topic 열은 지운다(가이드 10 — 반입에는 일반주제명 '만화'만)"""
    rows = data.get("rows") or []
    need_names = max([len([n for n in (r.get("values") or {}).get("names") or [] if str(n.get("name") or "").strip()]) for r in rows] + [0])
    need_pubs = 1 + max([len((r.get("values") or {}).get("publishers") or []) for r in rows] + [0])
    need_notes = 3 + max([len((r.get("values") or {}).get("notes_extra") or []) for r in rows] + [0])
    h = _header(ws)
    topics = [i + 1 for i, x in enumerate(h) if x == "/mods/subject/topic"]
    if len(topics) > 1:
        ws.delete_cols(topics[1]); h = _header(ws)
    name_groups = [i + 1 for i, x in enumerate(h) if x == "/mods/name"]
    for _ in range(max(0, need_names - len(name_groups))):
        h = _header(ws)
        last_role = max(i + 1 for i, x in enumerate(h) if x == "/mods/name/role/roleTerm")
        _insert_after(ws, last_role, NAME_GROUP)
    h = _header(ws)
    pubs = [i + 1 for i, x in enumerate(h) if x == "/mods/originInfo/publisher"]
    first_pub = pubs[0]
    for k in range(max(0, need_pubs - 1)):
        _insert_after(ws, first_pub + k, ["/mods/originInfo/publisher"])
    h = _header(ws)
    note_types = [i + 1 for i, x in enumerate(h) if x == "/mods/note[@type]"]
    for _ in range(max(0, need_notes - len(note_types))):
        h = _header(ws)
        last = max(i + 1 for i, x in enumerate(h) if x == "/mods/note[@type]")
        _insert_after(ws, last, ["/mods/note", "/mods/note[@type]"])
    return {"names": max(need_names, 3), "publishers": need_pubs, "notes": need_notes}


def write(data: dict, folder: Path | None, out_xlsx: Path, template: Path = TEMPLATE_83) -> dict:
    """반입용 엑셀을 쓴다. 돌려주는 것: {'rows', 'confirm'(노란 칸 수), 'out', 'columns'}. folder 가 None 이면 수량·파일 형식 칸을 비운다(양식 시험용)."""
    out_xlsx = Path(out_xlsx)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template, out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx)
    ws = wb["Contents"]
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    rep_n = _grow_columns(ws, data)
    col = _columns(ws)
    vw = arrange.View(folder, data) if folder is not None else None
    nconfirm = 0
    for r in data["rows"]:
        v = r.get("values") or {}
        cells: dict[Key, object] = dict(CONSTANTS)
        for name, key in FIELDS.items():
            val = v.get(name)
            if val is not None and str(val).strip() != "":
                cells[key] = int(val) if name in NUMERIC and re.fullmatch(r"\d+", str(val).strip()) else str(val).strip()
        if str(v.get("title_parallel") or "").strip():
            cells[("/mods/titleInfo[@type]", 1)] = "parallel"
        for i, n in enumerate([n for n in v.get("names") or [] if str(n.get("name") or "").strip()]):
            nk = name_keys(i)
            cells[nk["name"]] = str(n["name"]).strip()
            cells[nk["type"]] = n.get("type") or "개인명"
            if str(n.get("role") or "").strip():
                cells[nk["role"]] = str(n["role"]).strip()
            if i == 0:
                cells[("/mods/name[@usage]", 0)] = "primary"
        for k, pub in enumerate([x for x in v.get("publishers") or [] if str(x).strip()], start=1):
            cells[("/mods/originInfo/publisher", k)] = str(pub).strip()
        o2 = v.get("origin2") or {}
        if isinstance(o2, dict) and any(str(o2.get(k) or "").strip() for k in ("publisher", "dateIssued", "type", "issuance")):
            for k, key in ORIGIN2_KEYS.items():
                if str(o2.get(k) or "").strip():
                    cells[key] = str(o2[k]).strip()
            if str(o2.get("publisher") or "").strip():
                cells[("/mods/originInfo/publisher", rep_n["publishers"])] = str(o2["publisher"]).strip()     # 둘째 출처정보의 발행처 = 첫 출처정보 발행처들 다음 번째
            if not str(o2.get("issuance") or "").strip():
                cells[ORIGIN2_KEYS["issuance"]] = "단행자료"
        for k, x in enumerate([x for x in v.get("notes_extra") or [] if isinstance(x, dict) and str(x.get("text") or "").strip()], start=3):
            cells[("/mods/note", k)] = str(x["text"]).strip()
            if str(x.get("type") or "").strip():
                cells[("/mods/note[@type]", k)] = str(x["type"]).strip()
        if vw is not None:
            imgs = vw.images(f"{data['manuscripts_root']}/{r['folder']}")
            cells[COMPUTED["extent"]] = extent_string(len(imgs), sum(p.stat().st_size for p in imgs), v.get("color") or "천연색")
            exts = {p.suffix.lower().lstrip(".") for p in imgs}
            cells[COMPUTED["format"]] = "JPG" if exts <= {"jpg"} else "JPEG" if exts <= {"jpeg", "jpg"} else sorted(exts)[0].upper()
        if str(r.get("thumb_file") or "").strip():
            cells[COMPUTED["thumb_file"]] = str(r["thumb_file"]).strip()
        ws.append([None] * ws.max_column)
        rn = ws.max_row
        for key, val in cells.items():
            if key in col:
                ws.cell(row=rn, column=col[key]).value = val
                if key == FIELDS["identifier"] and isinstance(val, int):
                    ws.cell(row=rn, column=col[key]).number_format = ISBN_FORMAT
                elif key == ("/mods/classification[@edition]", 0):
                    ws.cell(row=rn, column=col[key]).number_format = "@"
        for c in r.get("confirm") or []:
            name = str(c.get("field"))
            key = FIELDS.get(name) or COMPUTED.get(name) or {"names": NAME_KEYS[0]["name"], "color": COMPUTED["extent"], "publishers": ("/mods/originInfo/publisher", 1), "notes_extra": ("/mods/note", 3), "origin2": ("/mods/originInfo[@type]", 0)}.get(name)
            if key in col:
                cell = ws.cell(row=rn, column=col[key])
                cell.fill = YELLOW
                prev = cell.comment.text + "\n" if cell.comment else ""
                cell.comment = Comment((prev + confirm_text(c))[:2000], "kolis_tool")
                nconfirm += 1
    for name in wb.sheetnames:
        if name not in ("Sample", "Contents"):
            del wb[name]
    wb.save(out_xlsx)
    return {"rows": len(data["rows"]), "confirm": nconfirm, "out": str(out_xlsx), "columns": ws.max_column}


def marks(xlsx: Path) -> int:
    """반입용 엑셀에 남아 있는 확인 표시(노란색 칸·메모)의 수."""
    wb = openpyxl.load_workbook(xlsx)
    try:
        return sum(1 for row in wb["Contents"].iter_rows(min_row=2) for c in row if c.comment or (c.fill and c.fill.fill_type == "solid" and c.fill.fgColor.rgb in ("FFFFFF00", "00FFFF00")))
    finally:
        wb.close()


def clear_marks(xlsx: Path) -> int:
    """직원이 확인을 끝냈을 때: 프로그램이 넣은 확인 표시(노란색 칸·메모)를 지운다. 값은 건드리지 않는다. 완료 사례의 반입용 엑셀에는 색·메모가 없다."""
    wb = openpyxl.load_workbook(xlsx)
    n = 0
    for row in wb["Contents"].iter_rows(min_row=2):
        for c in row:
            yellow = c.fill and c.fill.fill_type == "solid" and c.fill.fgColor.rgb in ("FFFFFF00", "00FFFF00")
            if c.comment or yellow:
                n += 1
                c.comment = None
                if yellow:
                    c.fill = PatternFill(fill_type=None)
    wb.save(xlsx)
    return n


def main(import_json: str) -> int:
    """에이전트용 명령: 검사하고, 통과한 형식이면 반입용 엑셀을 쓴다. job.json 에서 납품 폴더와 출력 경로를 읽는다."""
    p = Path(import_json)
    job = json.loads(p.with_name("job.json").read_text(encoding="utf-8"))
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        print(f"걸린 항목 1건:\n- import.json 을 읽지 못했습니다: {e}")
        return 1
    fails = check(data, Path(job["folder"]))
    blocking = [f for f in fails if not f.startswith("검수")]
    if not blocking:
        r = write(data, Path(job["folder"]), Path(job["output_xlsx"]))
        print(f"반입용 엑셀을 썼습니다: {r['out']} ({r['rows']}행, 사람이 확인할 칸 {r['confirm']}개)")
    if not fails:
        print("통과")
        return 0
    print(f"걸린 항목 {len(fails)}건:")
    for f in fails:
        print(f"- {f}")
    return 1
