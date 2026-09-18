"""③-LLM 반입용 엑셀 보완 에이전트 (human-in-the-loop).

convert 가 만든 반입용 xlsx 에서 노란색(확인 필요) 셀만 골라, 근거 세 가지를 모아 Claude 에 묻고
'제안' 시트에 (행, 열, 현재값, 제안값, 근거, 확신도) 를 적는다. 본문 셀은 건드리지 않는다.
사람이 제안 시트를 보고 승인한 것만 `apply` 로 본문에 반영한다.

근거:
  ① 원문 이미지: 작품 폴더의 앞 N장 + 뒤 M장(표지·타이틀컷·크레딧면이 보통 여기 있음)
  ② 출판사용 엑셀 행(convert 가 source 로 남긴 값)
  ③ 플랫폼 페이지: 웹 검색·웹 페치 서버 도구로 조회(플랫폼 출처 값은 각괄호 규칙 적용)

규칙집: docs/rulebook/mods-input-guide.md 를 system 프롬프트로 넣고 캐시한다.
모델: claude-opus-5 (어댑티브 사고). API 키는 ANTHROPIC_API_KEY 또는 `ant auth login`.
외부 API 사용이 불허된 환경이면 --dry-run 으로 '질의 묶음(json)' 만 만들어 클로드코드 세션에서 직접 읽어 답하게 한다.
"""
from __future__ import annotations
import base64, json, re
from dataclasses import dataclass, asdict
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from .common import list_images, resolve_folder

HERE = Path(__file__).parent
RULEBOOK = HERE.parent / "docs/rulebook/mods-input-guide.md"
RULEBOOK2 = HERE.parent / "docs/rulebook/ebook-guideline-ch3.md"   # "전자책 정리 지침에 따른다"의 그 지침(3장)
MODEL = "claude-opus-5"
GREEN = PatternFill("solid", fgColor="CCFFCC")

SYSTEM = """당신은 국립중앙도서관 웹툰 납본 사업의 서지 구축 보조자입니다. 아래 규칙집(MODS 입력 가이드)에 따라
반입용 기초메타데이터의 '확인 필요' 항목에 값을 제안합니다.

원칙:
- 으뜸정보원 우선순위: 원문 이미지(표지·타이틀컷·크레딧면) > 출판사 기초메타데이터 > 플랫폼 페이지.
- 플랫폼 페이지에서만 확인한 값은 반드시 각괄호 [ ] 로 감쌉니다. 원문에 있는 각괄호는 원괄호 ( ) 로 바꿉니다.
- 근거가 없는 값은 지어내지 말고 제안값을 비워 두고 reason 에 '근거 없음' 이라고 씁니다.
- 각 제안에 evidence(어느 이미지 몇 번째 / 출판사 엑셀 열 / URL)와 confidence(high/medium/low)를 붙입니다.
- 권차는 원문 표기 그대로(예: '01회', '2부 3화'). 외전·프롤로그·에필로그는 권차를 비우고 권차표제(partName)에 씁니다.
- 영문 표제·저자명 대소문자 규칙, 발행지 각괄호·코드(kormarccountry), 역할어(글/그림/원작/각색) 규칙은 규칙집을 따릅니다.

규칙집:
"""

SCHEMA = {
    "type": "object",
    "properties": {
        "proposals": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "column": {"type": "string"}, "occurrence": {"type": "integer"},
                "value": {"type": "string"}, "reason": {"type": "string"},
                "evidence": {"type": "string"}, "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
            },
            "required": ["column", "occurrence", "value", "reason", "evidence", "confidence"], "additionalProperties": False}},
        "notes": {"type": "string"},
    },
    "required": ["proposals", "notes"], "additionalProperties": False,
}


# 회차마다 달라지는 열. 나머지 확인 항목은 작품 단위로 한 번만 묻는다.
RULE_OWNED = {"/mods/physicalDescription/extent", "thum_files", "reward_yn"}   # 노란색은 사람 확인용, LLM 에 묻지 않음
ROW_LEVEL = {"/mods/titleInfo/title", "/mods/titleInfo/subTitle", "/mods/titleInfo/partNumber", "/mods/titleInfo/partName",
             "/mods/identifier", "/mods/physicalDescription/extent", "/mods/originInfo/dateIssued", "thum_files", "contents_price", "reward_yn"}


@dataclass
class Query:
    row: int                   # 대표 행(작품 질의는 첫 회차 행)
    title: str
    publisher_row: dict
    flagged: list[dict]        # {column, occurrence, current, reason}
    images: list[str]          # 파일 경로
    platform_hint: str
    level: str = "row"         # "work"(작품 단위, rows 전체에 적용) | "row"(회차 단위)
    rows: list[int] = None     # level=work 일 때 적용 대상 행들


def _media(p: Path) -> str:
    return {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "gif": "image/gif"}.get(p.suffix.lower().lstrip("."), "image/jpeg")


def collect_queries(xlsx: Path, root: Path | None, head: int = 3, tail: int = 2) -> list[Query]:
    """노란 셀을 읽어 작품 단위 질의(저자·발행지·URL 등, 첫 회차 이미지 앞3·뒤2장)와
    회차 단위 질의(권차·권차표제 등, 타이틀컷 1장)로 나눈다."""
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Contents"]
    header = [(c.value or "") if isinstance(c.value, str) else "" for c in ws[1]]
    occ = []; seen = {}
    for h in header:
        occ.append(seen.get(h, 0)); seen[h] = seen.get(h, 0) + 1
    title_c = header.index("/mods/titleInfo/title") + 1
    start = 3 if str(ws.cell(row=2, column=title_c).value or "") == "제목" else 2
    src_ws = wb["_source"] if "_source" in wb.sheetnames else None
    total = ws.max_row - start + 1
    per_row: list[tuple[int, str, dict, list[dict], list[str]]] = []
    for r in range(start, ws.max_row + 1):
        flagged = []
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=r, column=c)
            if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb in ("00FFFF00", "FFFFFF00") and header[c - 1] not in RULE_OWNED:
                flagged.append({"column": header[c - 1], "occurrence": occ[c - 1], "current": cell.value,
                                "reason": cell.comment.text if cell.comment else ""})
        title = str(ws.cell(row=r, column=title_c).value or "")
        part = str(ws.cell(row=r, column=header.index("/mods/titleInfo/partNumber") + 1).value or "")
        pub = {}
        if src_ws is not None:
            pub = {k: v for k, v in zip([c.value for c in src_ws[1]], [c.value for c in src_ws[r - start + 2]]) if v is not None}
        images: list[str] = []
        if root is not None:
            fname = re.split(r"[\\/]", str(pub.get("파일명") or ""))[-1].rsplit(".", 1)[0]
            folder, _how = resolve_folder(root, title, part, fname, r - start, total)
            if folder:
                images = [str(p) for p in list_images(folder)]
        per_row.append((r, title, pub, flagged, images))
    out: list[Query] = []
    by_title: dict[str, list] = {}
    for item in per_row:
        by_title.setdefault(item[1], []).append(item)
    for title, items in by_title.items():
        r0, _, pub0, _, imgs0 = items[0]
        work_flags: dict[tuple, dict] = {}
        for r, _, pub, flagged, imgs in items:
            for f in flagged:
                if f["column"] not in ROW_LEVEL:
                    work_flags.setdefault((f["column"], f["occurrence"]), f)
        if work_flags:
            im = imgs0[:head] + imgs0[-tail:] if len(imgs0) > head + tail else imgs0
            out.append(Query(r0, title, pub0, list(work_flags.values()), im, str(pub0.get("최초연재플랫폼") or ""),
                             level="work", rows=[it[0] for it in items]))
        for r, _, pub, flagged, imgs in items:
            rf = [f for f in flagged if f["column"] in ROW_LEVEL]
            if rf:
                out.append(Query(r, title, pub, rf, imgs[:1], str(pub.get("최초연재플랫폼") or ""), level="row", rows=[r]))
    return out


TILE_H = 1600   # 웹툰 원고는 세로 1만 px 이상이라 통째로는 판독이 안 됨. 앞·뒤 타일만 잘라 준다.

def tile_images(paths: list[str], out_dir: Path, head_tiles: int = 2, tail_tiles: int = 2) -> list[str]:
    """각 이미지의 위쪽 head_tiles 장 + 아래쪽 tail_tiles 장(각 TILE_H px)을 PNG 로 저장해 경로 목록 반환.
    첫 이미지(표지·타이틀컷)는 위쪽, 마지막 이미지(크레딧면)는 아래쪽이 중요하다."""
    from PIL import Image
    out_dir.mkdir(parents=True, exist_ok=True)
    tiles: list[str] = []
    for p in paths:
        p = Path(p)
        with Image.open(p) as im:
            w, h = im.size
            if h <= TILE_H * (head_tiles + tail_tiles):
                tiles.append(str(p)); continue
            spans = [(i * TILE_H, (i + 1) * TILE_H) for i in range(head_tiles)]
            spans += [(h - (tail_tiles - i) * TILE_H, h - (tail_tiles - i - 1) * TILE_H) for i in range(tail_tiles)]
            for k, (y0, y1) in enumerate(spans):
                t = out_dir / f"{p.parent.name}_{p.stem}_t{k}_{y0}-{y1}.png"
                if not t.exists():
                    im.crop((0, y0, w, y1)).save(t)
                tiles.append(str(t))
    return tiles


def build_prompt_text(q: Query) -> str:
    ask = {"작품": q.title, "질의_단위": "작품(전 회차 공통 항목)" if q.level == "work" else f"회차({q.row}행)",
           "출판사_기초메타데이터_행": q.publisher_row, "최초연재플랫폼": q.platform_hint, "확인_필요_항목": q.flagged}
    web = ("플랫폼 페이지가 필요하면 웹 검색·페치로 작품 페이지(작품 홈 URL, 회차 URL)를 찾아 location/url 후보와 발행지·저자 역할어를 확인하세요. "
           if q.level == "work" else "이 질의는 회차 표기(권차·권차표제 등)만 다룹니다. 웹 조회는 필요 없습니다. ")
    return ("다음 작품의 확인 필요 항목에 값을 제안하세요. " + web + "\n\n" + json.dumps(ask, ensure_ascii=False, indent=1))


def build_user_content(q: Query) -> list[dict]:
    content: list[dict] = []
    for i, p in enumerate(tile_images(q.images, Path("work/tiles")) if q.images else [], 1):
        pth = Path(p)
        content.append({"type": "text", "text": f"[원문 이미지 {i}: {pth.name}]"})
        content.append({"type": "image", "source": {"type": "base64", "media_type": _media(pth),
                                                    "data": base64.standard_b64encode(pth.read_bytes()).decode()}})
    content.append({"type": "text", "text": build_prompt_text(q)})
    return content


# ---------- 헤드리스 실행(claude -p / codex exec) ----------
EXEC_SUFFIX = """

출력 형식: 마지막에 ```json 코드블록 하나로 다음 스키마의 JSON 만 출력하세요.
{"proposals":[{"column":"/mods/…","occurrence":0,"value":"…","reason":"…","evidence":"…","confidence":"high|medium|low"}],"notes":"…"}
근거 없는 항목은 value 를 빈 문자열로 두세요."""


def ask_exec(q: Query, runner: str = "claude", timeout: int = 600) -> dict:
    import subprocess
    system = SYSTEM + RULEBOOK.read_text(encoding="utf-8")
    tiles = tile_images(q.images, Path("work/tiles")) if q.images else []
    img_lines = "\n".join(f"- 원문 이미지 조각 {i}: {p}  (Read 도구로 열어 보세요. 파일명의 y 범위는 원본에서 잘라 낸 위치)" for i, p in enumerate(tiles, 1))
    ref = f"\n\n규칙집이 '전자책 정리 지침에 따른다'고 한 항목은 {RULEBOOK2} 를 Read 해서 해당 절을 확인하세요.\n" if RULEBOOK2.exists() else ""
    prompt = system + ref + "\n\n---\n" + (img_lines + "\n\n" if img_lines else "") + build_prompt_text(q) + EXEC_SUFFIX
    import shutil
    if runner == "claude":
        exe = shutil.which("claude") or shutil.which("claude.cmd")
        if not exe:
            return {"proposals": [], "notes": "claude 실행파일을 찾지 못함(PATH 확인)"}
        cmd = [exe, "-p", "--output-format", "json", "--allowedTools", "Read", "WebSearch", "WebFetch"]
        res = subprocess.run(cmd, input=prompt, capture_output=True, text=True, timeout=timeout, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        if res.returncode != 0:
            return {"proposals": [], "notes": f"claude -p 실패: {res.stderr[-500:]}"}
        try:
            text = json.loads(res.stdout).get("result", "")
        except json.JSONDecodeError:
            text = res.stdout
    elif runner == "codex":
        exe = shutil.which("codex") or shutil.which("codex.cmd")
        if not exe:
            return {"proposals": [], "notes": "codex 실행파일을 찾지 못함(PATH 확인)"}
        cmd = [exe, "exec", "--skip-git-repo-check", "-"]
        res = subprocess.run(cmd, input=prompt, capture_output=True, text=True, timeout=timeout, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        if res.returncode != 0:
            return {"proposals": [], "notes": f"codex exec 실패: {res.stderr[-500:]}"}
        text = res.stdout
    else:
        raise ValueError(runner)
    m = re.findall(r"```json\s*(\{.*?\})\s*```", text, re.S)
    if not m:
        return {"proposals": [], "notes": "JSON 블록 없음", "findings": text[-2000:]}
    try:
        data = json.loads(m[-1])
    except json.JSONDecodeError as e:
        return {"proposals": [], "notes": f"JSON 파싱 실패: {e}", "findings": text[-2000:]}
    data.setdefault("proposals", []); data.setdefault("notes", "")
    data["findings"] = text
    return data


def ask_claude(q: Query, client=None) -> dict:
    import anthropic
    client = client or anthropic.Anthropic()
    system = [{"type": "text", "text": SYSTEM + RULEBOOK.read_text(encoding="utf-8"), "cache_control": {"type": "ephemeral"}}]
    tools = [{"type": "web_search_20260209", "name": "web_search", "max_uses": 5},
             {"type": "web_fetch_20260209", "name": "web_fetch", "max_uses": 5}]
    # 1단계: 근거 수집 + 자유 서술 (서버 도구 사용)
    msgs = [{"role": "user", "content": build_user_content(q)}]
    with client.messages.stream(model=MODEL, max_tokens=16000, system=system, tools=tools, messages=msgs,
                                thinking={"type": "adaptive"}) as s:
        first = s.get_final_message()
    if first.stop_reason == "refusal":
        return {"proposals": [], "notes": f"refusal: {getattr(first.stop_details, 'explanation', '')}"}
    findings = "\n".join(b.text for b in first.content if b.type == "text")
    # 2단계: 구조화 (도구 없이 JSON 스키마 출력)
    second = client.messages.create(
        model=MODEL, max_tokens=8000, system=system,
        messages=[{"role": "user", "content": "아래 조사 결과를 확인_필요_항목별 제안 JSON 으로 정리하세요. 근거 없는 항목은 value 를 빈 문자열로.\n\n"
                   + "확인_필요_항목:\n" + json.dumps(q.flagged, ensure_ascii=False) + "\n\n조사 결과:\n" + findings}],
        output_config={"format": {"type": "json_schema", "schema": SCHEMA}},
    )
    text = next(b.text for b in second.content if b.type == "text")
    data = json.loads(text)
    data["findings"] = findings
    data["usage"] = {"cache_read": first.usage.cache_read_input_tokens, "in": first.usage.input_tokens, "out": first.usage.output_tokens}
    return data


def write_proposals(xlsx: Path, results: list[tuple[Query, dict]]) -> None:
    wb = openpyxl.load_workbook(xlsx)
    if "제안" in wb.sheetnames:
        del wb["제안"]
    ws = wb.create_sheet("제안")
    ws.append(["행(쉼표 구분)", "작품", "열", "n번째", "현재값", "제안값", "확신도", "근거", "이유", "승인(Y)"])
    for q, res in results:
        cur = {(f["column"], f["occurrence"]): f["current"] for f in q.flagged}
        rows = ",".join(str(r) for r in (q.rows or [q.row]))
        for p in res.get("proposals", []):
            ws.append([rows, q.title, p["column"], p.get("occurrence", 0), cur.get((p["column"], p.get("occurrence", 0))),
                       p.get("value", ""), p.get("confidence", ""), p.get("evidence", ""), p.get("reason", ""), ""])
        if res.get("notes"):
            ws.append([rows, q.title, "(비고)", "", "", "", "", "", res["notes"], ""])
    wb.save(xlsx)


def apply_approved(xlsx: Path) -> int:
    """'제안' 시트에서 승인(Y) 표시된 행만 Contents 본문에 반영하고 초록색으로 표시."""
    wb = openpyxl.load_workbook(xlsx)
    ws, ps = wb["Contents"], wb["제안"]
    header = [(c.value or "") if isinstance(c.value, str) else "" for c in ws[1]]
    col_of = {}; seen = {}
    for i, h in enumerate(header):
        col_of[(h, seen.get(h, 0))] = i + 1; seen[h] = seen.get(h, 0) + 1
    n = 0
    for r in ps.iter_rows(min_row=2, values_only=True):
        rows, _, col, occ, _, val, _, evid, _, ok = r[:10]
        if str(ok or "").strip().upper() != "Y" or (col, int(occ or 0)) not in col_of:
            continue
        for row in str(rows).split(","):
            cell = ws.cell(row=int(row), column=col_of[(col, int(occ or 0))])
            cell.value = val if val != "" else None
            cell.fill = GREEN
            cell.comment = Comment(f"승인 반영. 근거: {evid}", "kolis_tool")
            n += 1
    wb.save(xlsx)
    return n


def run(xlsx: Path, root: Path | None, dry_run: bool = False, limit: int | None = None, runner: str = "claude") -> None:
    qs = collect_queries(xlsx, root)
    if limit: qs = qs[:limit]
    nwork = sum(1 for q in qs if q.level == "work")
    print(f"질의 {len(qs)}건 (작품 단위 {nwork}, 회차 단위 {len(qs) - nwork})")
    if dry_run:
        pkt = xlsx.with_suffix(".queries.json")
        pkt.write_text(json.dumps([asdict(q) for q in qs], ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"질의 {len(qs)}건 → {pkt} (클로드코드 세션에서 읽어 '제안' 시트를 채우는 용도)")
        return
    results = []
    if runner in ("claude", "codex"):
        for i, q in enumerate(qs, 1):
            print(f"[{i}/{len(qs)}] {q.level} {q.row}행 {q.title} … 항목 {len(q.flagged)}, 이미지 {len(q.images)}")
            res = ask_exec(q, runner)
            print("   제안", len(res.get("proposals", [])), "|", (res.get("notes") or "")[:80])
            results.append((q, res))
        write_proposals(xlsx, results)
        print(f"'제안' 시트 기록 완료 → {xlsx}")
        return
    import anthropic
    client = anthropic.Anthropic()
    for i, q in enumerate(qs, 1):
        print(f"[{i}/{len(qs)}] {q.level} {q.row}행 {q.title} … 항목 {len(q.flagged)}, 이미지 {len(q.images)}")
        try:
            results.append((q, ask_claude(q, client)))
        except anthropic.RateLimitError as e:
            print("  rate limit:", e.message); break
        except anthropic.APIStatusError as e:
            print("  API 오류:", e.status_code, e.message); results.append((q, {"proposals": [], "notes": f"API 오류 {e.status_code}"}))
        except anthropic.APIConnectionError as e:
            print("  네트워크 오류:", e); break
    write_proposals(xlsx, results)
    print(f"'제안' 시트 기록 완료 → {xlsx}")
