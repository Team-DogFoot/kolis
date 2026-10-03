"""반입용 엑셀 — MODS 트리 ↔ 열 직렬화(형태가 정해진 일이라 코드가 한다. 2026-10-04 개선안 3-A).

열(머리글)을 고정하지 않는다. 에이전트가 행마다 MODS 트리를 쓰면, 이 모듈이 모든 행의 반복 수(최대치)로 열을 만들고 값을 옮긴다.
묶음 수에 상한이 없다(저자 n, 출처정보 n, 주기 n, 주제명 n, 식별기호 n, 원문주소 n, 표제정보 n, 다른이름 n …).

트리 표기(에이전트가 쓰는 import.json 의 rows[].mods):
  - 요소 이름 = MODS 태그 이름. 반복 요소는 목록. 속성은 "@이름". 요소 글자는 "_"(속성·자식이 없으면 글자만 써도 된다).
  - 예: {"titleInfo": [{"title": "로맨스 낫 로맨틱", "partNumber": "01회"}, {"@type": "parallel", "title": "Romance not romantic"}],
         "name": [{"@usage": "primary", "@type": "개인명", "@ID": "KAC…", "@authority": "국립중앙도서관전거데이터", "namePart": "홍길동",
                   "role": {"roleTerm": "글"}, "alternativeName": [{"@type": "nickname", "namePart": "길동이"}]}],
         "originInfo": [{"@eventType": "publication", "place": [{"placeTerm": {"_": "[서울]", "@type": "text"}}, {"placeTerm": {"_": "ulk", "@authority": "kormarccountry", "@type": "code"}}],
                         "issuance": "단행자료", "publisher": ["A", "B"], "dateIssued": "20240101", "edition": "개정판"}],
         "note": [{"_": "A은 B의 임프린트임"}, {"_": "15세 이용가", "@type": "target audience"}], "subject": [{"@ID": "KSH1998022212", "@authority": "…", "topic": "만화[漫畵]"}],
         "identifier": [{"_": "9791109764511", "@type": "isbn"}], "location": {"url": ["…", "…"], "physicalLocation": "국립중앙도서관"}, …}
  - 양식 고유 칸은 rows[].extra: {"currency_code", "contents_price", "compensation", "reward_yn", "thum_files"}.

열 순서(직렬화 규칙 — 예시 88열·템플릿 83열·2026-10-03 반입 헤더 시험 엑셀 그대로. Sample 시트 순서가 아니다):
  1. 요소 순서는 ORDER_TOP. 알 수 없는 요소는 extension 뒤(시험 전 항목으로 검사가 알린다).
  2. 반복 요소는 묶음마다 부모 열(PARENT_COLS, 값 없음) → 자식·속성을 GROUP_ORDER 의 순서로. 묶음 수 = 모든 행의 최대치(최소치는 MIN_COUNTS, 템플릿과 같게).
  3. 속성 열 위치는 예시 그대로(name: 부모·@usage·namePart·@type·@ID·@authority / titleInfo: 부모·title·@type / subject: 부모·topic·@ID·@authority / originInfo: 부모·@eventType …).
  4. 발행지: 첫 출처정보는 부모 열 없이 placeTerm 쌍(템플릿·예시·끝까지 등록된 사례와 같음). 둘째 이후 출처정보는 placeTerm 마다 부모 열
     `/mods/originInfo/place`(1차 시험 꼴 = <place> 둘 = 구축 완료 XML 그림 34 와 같음). 부모 열이 없으면 앞 묶음에 붙는다(RECON 9).
  5. 다른이름 유형 열은 `[@type]`(Sample·템플릿 표기. 반입되면 altType 이 된다. `[@altType]` 도 읽을 때는 받는다).
  6. 양식 고유 칸 5개는 맨 뒤.
  7. 원문주소는 예시대로 url·url·physicalLocation·url… (같은 묶음 안 형제 순서는 도서관 77열 예시가 physicalLocation 을 먼저 두므로 파서가 가리지 않는다).
"""
from __future__ import annotations
import copy, json, re, shutil
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

HERE = Path(__file__).parent
TEMPLATE = HERE / "templates" / "import_template_83.xlsx"      # Sample 시트(도서관 양식 설명)만 쓴다. Contents 머리글은 여기서 만든다
YELLOW = PatternFill("solid", fgColor="FFFF00")
EXTRA = ["currency_code", "contents_price", "compensation", "reward_yn", "thum_files"]
NUMERIC_PATHS = {"/mods/originInfo/dateIssued", "/mods/identifier", "/mods/classification", "/mods/accessCondition/licenseType", "contents_price", "compensation"}
ISBN_FORMAT = r"0_);[Red]\(0\)"
TEXT_FORMAT_PATHS = {"/mods/classification[@edition]"}

ORDER_TOP = ["titleInfo", "name", "originInfo", "language", "physicalDescription", "targetAudience", "note", "subject", "classification",
             "identifier", "location", "accessCondition", "extension", "typeOfResource", "genre"]
KNOWN_AFTER = "extension"           # 알 수 없는 요소는 이 뒤에
# 요소별 자식·속성 순서. "@x" 는 속성, 그 밖은 자식 요소. "_" 는 글자 열(자기 경로). 없는 요소는 generic(글자 → 속성 → 자식).
GROUP_ORDER: dict[str, list[str]] = {
    "/mods/titleInfo": ["title", "@type", "subTitle", "partNumber", "partName", "nonSort"],
    "/mods/name": ["@usage", "namePart", "@type", "@ID", "@authority", "role", "displayForm", "alternativeName"],
    "/mods/name/role": ["roleTerm"],
    "/mods/name/alternativeName": ["@type", "namePart"],
    "/mods/originInfo": ["@eventType", "@type", "place", "issuance", "publisher", "dateIssued", "edition"],
    "/mods/originInfo/place": ["placeTerm"],
    "/mods/originInfo/place/placeTerm": ["_", "@authority", "@type"],
    "/mods/language": ["languageTerm"],
    "/mods/language/languageTerm": ["_", "@authority", "@type"],
    "/mods/physicalDescription": ["form", "extent", "reformattingQuality", "internetMediaType", "digitalOrigin"],
    "/mods/note": ["_", "@type"],
    "/mods/subject": ["topic", "genre", "@ID", "@authority"],
    "/mods/classification": ["_", "@authority", "@edition"],
    "/mods/identifier": ["_", "@type"],
    "/mods/location": ["url", "physicalLocation"],
    "/mods/accessCondition": ["licenseType"],
    "/mods/extension": ["regionOfPublishing"],
}
PARENT_COLS = {"/mods/titleInfo": "/mods/titleInfo/", "/mods/name": "/mods/name", "/mods/name/role": "/mods/name/role",
               "/mods/name/alternativeName": "/mods/name/alternativeName", "/mods/originInfo": "/mods/originInfo",
               "/mods/originInfo/place": "/mods/originInfo/place", "/mods/subject": "/mods/subject/"}
REPEAT = {"/mods/titleInfo", "/mods/name", "/mods/name/alternativeName", "/mods/originInfo", "/mods/originInfo/place", "/mods/originInfo/publisher",
          "/mods/note", "/mods/subject", "/mods/identifier", "/mods/location/url", "/mods/classification"}
# 템플릿 83열과 같은 최소 묶음 수(값이 없어도 열은 둔다 — 승인 양식과 같은 모양을 유지)
MIN_COUNTS = {"/mods/titleInfo": 2, "/mods/name": 3, "/mods/originInfo": 2, "/mods/note": 3, "/mods/subject": 1, "/mods/identifier": 1, "/mods/location/url": 3}
FIRST_ONLY = {("/mods/name", "@usage"), ("/mods/name", "alternativeName")}     # 템플릿·예시에서 첫 묶음에만 있는 열(값이 있으면 다른 묶음에도 낸다)
# 템플릿 첫 출처정보의 발행지 모양(부모 열 없음, placeTerm 둘). 둘째 이후는 place 마다 부모 열(1차 시험 꼴)
FIRST_ORIGIN_PLACES = [{"placeTerm": {"@type": "text"}}, {"placeTerm": {"@authority": "kormarccountry", "@type": "code"}}]

# 사람에게 보여 줄 한글 이름: Sample 시트의 한글명 + 보충(Sample 에 없는 열. 2026-10-04 개선안 리뷰 2-6)
LABELS_EXTRA = {
    "/mods/titleInfo/partName": "권차표제", "/mods/titleInfo[@type]": "표제 유형", "/mods/name[@usage]": "대표저자", "/mods/name[@ID]": "저자 전거 번호",
    "/mods/name[@authority]": "저자 전거 기관", "/mods/name[@type]": "저자 유형", "/mods/name/displayForm": "디스플레이형식",
    "/mods/name/alternativeName[@type]": "다른이름 유형", "/mods/name/alternativeName[@altType]": "다른이름 유형", "/mods/name/alternativeName/namePart": "다른이름",
    "/mods/originInfo": "출처정보", "/mods/originInfo[@eventType]": "출처정보 유형", "/mods/originInfo[@type]": "출처정보 유형", "/mods/originInfo/edition": "판사항",
    "/mods/originInfo/place/placeTerm[@type]": "발행지 표기 유형", "/mods/originInfo/place/placeTerm[@authority]": "발행국 부호표",
    "/mods/language/languageTerm[@authority]": "언어 부호표", "/mods/note[@type]": "주기 유형", "/mods/subject/genre": "장르주제명", "/mods/subject[@ID]": "주제명 번호",
    "/mods/classification[@edition]": "분류표 판", "currency_code": "통화", "contents_price": "정가", "compensation": "보상금", "reward_yn": "보상여부", "thum_files": "썸네일 파일명",
}


# ---------------------------------------------------------------- 트리 다루기
def _is_scalar(v) -> bool:
    return isinstance(v, (str, int, float)) or v is None


def _text(node) -> str:
    if _is_scalar(node):
        return "" if node is None else str(node)
    return "" if node.get("_") is None else str(node["_"])


def _node(node) -> dict:
    """스칼라면 {"_": v} 로."""
    return {"_": node} if _is_scalar(node) else node


def _lst(v) -> list:
    return v if isinstance(v, list) else [v]


def parse_path(path: str) -> list:
    """'name[1].alternativeName[0].namePart' / 'originInfo[0].publisher[1]' / 'identifier[1].@type' / 'extra.contents_price' → 키 목록."""
    keys: list = []
    for part in path.strip().split("."):
        m = re.fullmatch(r"(@?[A-Za-z_]+)((?:\[\d+\])*)", part)
        if not m:
            raise ValueError(f"경로 표기가 틀렸습니다: {path!r}")
        keys.append(m.group(1))
        for idx in re.findall(r"\[(\d+)\]", m.group(2)):
            keys.append(int(idx))
    return keys


def get(tree: dict, path: str):
    """트리 경로의 값(글자). 없으면 None."""
    cur = tree
    for k in parse_path(path):
        if isinstance(k, int):
            cur = _lst(cur)
            if k >= len(cur):
                return None
            cur = cur[k]
        else:
            if _is_scalar(cur):
                return None
            if isinstance(cur, list):
                cur = cur[0] if cur else None
                if cur is None or _is_scalar(cur):
                    return None
            cur = cur.get(k)
        if cur is None:
            return None
    return _text(cur) if not isinstance(cur, (dict, list)) or (isinstance(cur, dict) and "_" in cur) else cur


def set_value(tree: dict, path: str, value) -> None:
    keys = parse_path(path)
    cur = tree
    for i, k in enumerate(keys):
        last = i == len(keys) - 1
        if isinstance(k, int):
            if not isinstance(cur, list):
                raise ValueError(f"목록이 아닌 곳에 번호가 왔습니다: {path}")
            while len(cur) <= k:
                cur.append({})
            if last:
                cur[k] = value
            else:
                if _is_scalar(cur[k]):
                    cur[k] = {"_": cur[k]}
                cur = cur[k]
        else:
            if isinstance(cur, list):
                cur = cur[0]
            if _is_scalar(cur):
                raise ValueError(f"글자 값 아래에 키를 넣을 수 없습니다: {path}")
            if last:
                cur[k] = value
            else:
                nxt = cur.get(k)
                if nxt is None:
                    nxt = [] if (i + 1 < len(keys) and isinstance(keys[i + 1], int)) else {}
                    cur[k] = nxt
                elif _is_scalar(nxt):
                    nxt = {"_": nxt}; cur[k] = nxt
                cur = nxt


# ---------------------------------------------------------------- 뼈대(모든 행의 반복 수 최대치)
def _skeleton(node, path: str) -> dict:
    """노드의 모양: {"attrs": {…}, "text": bool, "children": {name: [skeleton, …]}} (목록 길이 = 묶음 수)."""
    sk = {"attrs": {}, "text": False, "children": {}}
    if _is_scalar(node):
        sk["text"] = node not in (None, "")
        return sk
    for k, v in node.items():
        if k == "_":
            sk["text"] = v not in (None, "")
        elif k.startswith("@"):
            sk["attrs"][k] = v not in (None, "")
        else:
            cp = f"{path}/{k}"
            sk["children"][k] = [_skeleton(x, cp) for x in _lst(v)]
    return sk


def _merge(a: dict, b: dict) -> dict:
    out = {"attrs": dict(a["attrs"]), "text": a["text"] or b["text"], "children": {}}
    for k, v in b["attrs"].items():
        out["attrs"][k] = out["attrs"].get(k, False) or v
    names = list(a["children"]) + [k for k in b["children"] if k not in a["children"]]
    for k in names:
        la, lb = a["children"].get(k, []), b["children"].get(k, [])
        n = max(len(la), len(lb))
        merged = []
        for i in range(n):
            if i < len(la) and i < len(lb):
                merged.append(_merge(la[i], lb[i]))
            else:
                merged.append(copy.deepcopy(la[i] if i < len(la) else lb[i]))
        out["children"][k] = merged
    return out


def _empty_sk() -> dict:
    return {"attrs": {}, "text": False, "children": {}}


def _ensure_min(sk: dict, path: str) -> None:
    """최소 묶음 수(템플릿 모양)와 첫 출처정보의 발행지 모양을 보장한다."""
    for child, lst in list(sk["children"].items()):
        cp = f"{path}/{child}"
        n = MIN_COUNTS.get(cp, 0)
        while len(lst) < n:
            lst.append(_empty_sk())
        for i, s in enumerate(lst):
            _ensure_min(s, cp)
    for cp, n in MIN_COUNTS.items():
        parent, child = cp.rsplit("/", 1)
        if parent == path and child not in sk["children"]:
            sk["children"][child] = [_empty_sk() for _ in range(n)]
    if path == "/mods" and "originInfo" in sk["children"]:
        first = sk["children"]["originInfo"][0]
        places = first["children"].setdefault("place", [])
        want = [_skeleton(p, "/mods/originInfo/place") for p in FIRST_ORIGIN_PLACES]
        for i, w in enumerate(want):
            if i < len(places):
                places[i] = _merge(places[i], w)
            else:
                places.append(w)
        for o in sk["children"]["originInfo"]:
            o["children"].setdefault("publisher", [_empty_sk()])
            if not o["children"]["publisher"]:
                o["children"]["publisher"] = [_empty_sk()]
            for key in ("issuance", "dateIssued"):
                o["children"].setdefault(key, [_empty_sk()])
            if o is first:
                o["attrs"].setdefault("@eventType", True)
                o["children"].setdefault("edition", [_empty_sk()])
            else:
                o["attrs"].setdefault("@type", True)
        # 첫 저자 묶음의 템플릿 열(usage, 다른이름 1묶음)
        if "name" in sk["children"]:
            n0 = sk["children"]["name"][0]
            n0["attrs"].setdefault("@usage", True)
            n0["children"].setdefault("alternativeName", [_empty_sk()])
            if not n0["children"]["alternativeName"]:
                n0["children"]["alternativeName"] = [_empty_sk()]
            for nm in sk["children"]["name"]:
                nm["children"].setdefault("namePart", [_empty_sk()]); nm["attrs"].setdefault("@type", True)
                nm["children"].setdefault("role", [_empty_sk()])
                nm["children"]["role"][0]["children"].setdefault("roleTerm", [_empty_sk()])
            for alt in n0["children"]["alternativeName"]:
                alt["attrs"].setdefault("@type", True); alt["children"].setdefault("namePart", [_empty_sk()])
        for t in sk["children"].get("titleInfo", []):
            t["attrs"].setdefault("@type", True)
            for key in ("title", "subTitle", "partNumber", "partName"):
                t["children"].setdefault(key, [_empty_sk()])
        for nt in sk["children"].get("note", []):
            nt["attrs"].setdefault("@type", True)
        for s in sk["children"].get("subject", []):
            if "topic" not in s["children"] and "genre" not in s["children"]:
                s["children"]["topic"] = [_empty_sk()]
        for idf in sk["children"].get("identifier", []):
            idf["attrs"].setdefault("@type", True)
        for key, attrs in (("language", {}), ("physicalDescription", {}), ("classification", {"@authority": True, "@edition": True}), ("accessCondition", {}), ("extension", {})):
            sk["children"].setdefault(key, [_empty_sk()])
        lang = sk["children"]["language"][0]["children"].setdefault("languageTerm", [_empty_sk()])[0]
        lang["attrs"].setdefault("@authority", True); lang["attrs"].setdefault("@type", True)
        pd = sk["children"]["physicalDescription"][0]["children"]
        for key in ("form", "extent", "reformattingQuality", "internetMediaType", "digitalOrigin"):
            pd.setdefault(key, [_empty_sk()])
        cl = sk["children"]["classification"][0]
        cl["attrs"].setdefault("@authority", True); cl["attrs"].setdefault("@edition", True)
        sk["children"]["location"] = sk["children"].get("location") or [_empty_sk()]
        loc = sk["children"]["location"][0]["children"]
        loc.setdefault("physicalLocation", [_empty_sk()])
        urls = loc.setdefault("url", [])
        while len(urls) < MIN_COUNTS["/mods/location/url"]:
            urls.append(_empty_sk())
        sk["children"]["accessCondition"][0]["children"].setdefault("licenseType", [_empty_sk()])
        sk["children"]["extension"][0]["children"].setdefault("regionOfPublishing", [_empty_sk()])
        for key in ("targetAudience", "typeOfResource", "genre"):
            sk["children"].setdefault(key, [_empty_sk()])


# ---------------------------------------------------------------- 열 배치(layout)
class Col:
    __slots__ = ("header", "keys", "unknown")

    def __init__(self, header: str, keys: tuple | None, unknown: bool = False):
        self.header, self.keys, self.unknown = header, keys, unknown     # keys=None 이면 부모 열(값 없음)

    def __repr__(self):
        return f"Col({self.header}, {self.keys})"


def _walk(sk: dict, path: str, keys: tuple, out: list[Col], group_index: int = 0, parent_path: str = "", known: bool = True) -> None:
    """뼈대를 열 목록으로. path 는 '/mods/name' 꼴, keys 는 트리 키 경로(name, 1, alternativeName, 0 …)."""
    order = GROUP_ORDER.get(path)
    if path in PARENT_COLS:
        out.append(Col(PARENT_COLS[path], None))
    tokens = list(order) if order else (["_"] + sorted(sk["attrs"]) + list(sk["children"]))
    # 순서에 없는 속성·자식은 뒤에(시험 전 항목)
    for a in sk["attrs"]:
        if a not in tokens:
            tokens.append(a)
    for c in sk["children"]:
        if c not in tokens:
            tokens.append(c)
    if order and "_" not in order and sk["text"]:
        tokens.insert(0, "_")
    for t in tokens:
        if t == "_":
            if order is None or sk["text"] or "_" in (order or []):
                if path != "/mods":
                    out.append(Col(path, keys + ("_",)))
        elif t.startswith("@"):
            if t not in sk["attrs"]:
                continue
            if (path, t) in FIRST_ONLY and group_index > 0 and not sk["attrs"].get(t):
                continue
            out.append(Col(f"{path}[{t}]", keys + (t,), unknown=not known or (order is not None and t not in order)))
        else:
            if t not in sk["children"]:
                continue
            if (path, t) in FIRST_ONLY and group_index > 0 and not any(_has_value(s) for s in sk["children"][t]):
                continue
            cp = f"{path}/{t}"
            lst = sk["children"][t]
            is_known = known and (order is None or t in order) and (path != "/mods" or t in ORDER_TOP)
            for i, s in enumerate(lst):
                k2 = keys + (t, i) if (cp in REPEAT or len(lst) > 1) else keys + (t,)
                if cp == "/mods/originInfo/place" and group_index == 0 and path == "/mods/originInfo" and keys[-1] == 0:
                    # 첫 출처정보: 부모 열 없이 placeTerm 쌍(템플릿 모양)
                    _walk_no_parent(s, cp, k2, out, is_known)
                else:
                    _walk(s, cp, k2, out, group_index=i, parent_path=path, known=is_known)


def _walk_no_parent(sk: dict, path: str, keys: tuple, out: list[Col], known: bool) -> None:
    saved = PARENT_COLS.pop(path, None)
    try:
        _walk(sk, path, keys, out, known=known)
    finally:
        if saved is not None:
            PARENT_COLS[path] = saved


def _has_value(sk: dict) -> bool:
    return sk["text"] or any(sk["attrs"].values()) or any(_has_value(s) for lst in sk["children"].values() for s in lst)


def _top_sorted(sk: dict) -> dict:
    """최상위 요소를 ORDER_TOP 순으로, 모르는 요소는 extension 뒤에."""
    kids = sk["children"]
    ordered = {}
    for k in ORDER_TOP:
        if k in kids:
            ordered[k] = kids[k]
        if k == KNOWN_AFTER:
            for u in kids:
                if u not in ORDER_TOP:
                    ordered[u] = kids[u]
    for u in kids:
        if u not in ordered:
            ordered[u] = kids[u]
    return {"attrs": sk["attrs"], "text": sk["text"], "children": ordered}


def layout(rows: list[dict]) -> list[Col]:
    """행 목록(각 행 {"mods": 트리, "extra": {…}}) → 열 배치. 모든 행의 최대 반복 수, 템플릿 최소 모양."""
    sk = _empty_sk()
    for r in rows:
        sk = _merge(sk, _skeleton(r.get("mods") or {}, "/mods"))
    _ensure_min(sk, "/mods")
    sk = _top_sorted(sk)
    out: list[Col] = []
    _walk(sk, "/mods", (), out)
    # 원문주소: 예시(url·url·physicalLocation·url…) 순서로 맞춘다
    out = _reorder_location(out)
    for e in EXTRA:
        out.append(Col(e, ("extra", e)))
    return out


def _reorder_location(cols: list[Col]) -> list[Col]:
    idx = [i for i, c in enumerate(cols) if c.header.startswith("/mods/location/")]
    if not idx:
        return cols
    block = [cols[i] for i in idx]
    urls = [c for c in block if c.header == "/mods/location/url"]
    phys = [c for c in block if c.header == "/mods/location/physicalLocation"]
    rest = [c for c in block if c not in urls and c not in phys]
    new = urls[:2] + phys + urls[2:] + rest
    return cols[:idx[0]] + new + cols[idx[-1] + 1:]


# ---------------------------------------------------------------- 값 옮기기
def _resolve(row: dict, keys: tuple):
    cur = (row.get("extra") or {}) if keys and keys[0] == "extra" else (row.get("mods") or {})
    keys = keys[1:] if keys and keys[0] == "extra" else keys
    for i, k in enumerate(keys):
        if isinstance(k, int):
            cur = _lst(cur)
            if k >= len(cur):
                return None
            cur = cur[k]
            continue
        if isinstance(cur, list):
            cur = cur[0] if cur else None
        if k == "_":
            return _text(cur) if (not isinstance(cur, dict) or "_" in cur) else None
        if cur is None or _is_scalar(cur):
            return None
        cur = cur.get(k)
        if cur is None:
            return None
    if _is_scalar(cur):
        return None if cur is None else str(cur)
    if isinstance(cur, dict):
        return _text(cur) or None
    return None


def cell_value(row: dict, col: Col):
    if col.keys is None:
        return None
    v = _resolve(row, col.keys)
    if v in (None, ""):
        return None
    s = str(v).strip()
    if col.header in NUMERIC_PATHS and re.fullmatch(r"\d+", s):
        return int(s)
    return s


def column_index(cols: list[Col], path: str) -> int | None:
    """트리 경로('name[1].alternativeName[0].namePart', 'extra.contents_price') → 열 번호(1부터)."""
    keys = tuple(parse_path(path))
    if keys and keys[0] != "extra":
        keys = _normalize_keys(cols, keys)
    for i, c in enumerate(cols):
        if c.keys == keys:
            return i + 1
    # 글자 열은 '_' 로 끝난다
    for i, c in enumerate(cols):
        if c.keys == keys + ("_",):
            return i + 1
    return None


def _normalize_keys(cols: list[Col], keys: tuple) -> tuple:
    """'titleInfo.title' 처럼 번호를 뺀 표기를 열의 키 표기(titleInfo,0,title)로 맞춘다."""
    out: list = []
    path = "/mods"
    i = 0
    while i < len(keys):
        k = keys[i]
        out.append(k)
        if isinstance(k, str) and not k.startswith("@") and k != "_":
            path = f"{path}/{k}"
            nxt = keys[i + 1] if i + 1 < len(keys) else None
            if not isinstance(nxt, int) and any(c.keys and len(c.keys) > len(out) and c.keys[:len(out)] == tuple(out) and isinstance(c.keys[len(out)], int) for c in cols):
                out.append(0)
        i += 1
    return tuple(out)


def label(header: str, sample_labels: dict[str, str] | None = None) -> str:
    base = header.split("[")[0] if "[@" in header else header
    name = (sample_labels or {}).get(header) or LABELS_EXTRA.get(header) or (sample_labels or {}).get(base) or ""
    name = re.sub(r"\(사용안함\)", "", name).strip()
    return name


def sample_labels(template: Path = TEMPLATE) -> dict[str, str]:
    wb = openpyxl.load_workbook(template, read_only=True)
    try:
        ws = wb["Sample"]
        out = {}
        for row in ws.iter_rows(min_row=7, max_col=2, values_only=True):
            if row[0] and isinstance(row[0], str) and row[0].startswith("/mods") and row[1]:
                out[row[0].strip()] = str(row[1]).strip()
        return out
    finally:
        wb.close()


def column_label(cols: list[Col], index: int, labels: dict[str, str] | None = None) -> str:
    """사람에게 보여 줄 칸 이름: 경로 #묶음번호 (한글)."""
    c = cols[index - 1]
    same = [i for i, x in enumerate(cols) if x.header == c.header]
    n = same.index(index - 1) + 1
    head = c.header + (f" #{n}" if len(same) > 1 else "")
    kor = label(c.header, labels)
    return f"{head} ({kor})" if kor else head


# ---------------------------------------------------------------- 쓰기
def write(data: dict, out_xlsx: Path, confirm_text=None, template: Path = TEMPLATE) -> dict:
    """반입용 엑셀을 쓴다. data = {"rows": [{"mods": 트리, "extra": {…}, "confirm": [{"path", "reason", …}]}]}.
    돌려주는 것: {"rows", "confirm"(노란 칸 수), "out", "columns", "headers", "unknown"(시험 전 열)}."""
    rows = data.get("rows") or []
    cols = layout(rows)
    out_xlsx = Path(out_xlsx)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template, out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx)
    ws = wb["Contents"]
    ws.delete_rows(1, ws.max_row)
    ws.append([c.header for c in cols])
    nconfirm = 0
    for r in rows:
        values = [cell_value(r, c) for c in cols]
        ws.append(values)
        rn = ws.max_row
        for i, c in enumerate(cols, 1):
            if c.header == "/mods/identifier" and isinstance(values[i - 1], int):
                ws.cell(row=rn, column=i).number_format = ISBN_FORMAT
            elif c.header in TEXT_FORMAT_PATHS and values[i - 1] is not None:
                ws.cell(row=rn, column=i).number_format = "@"
        for cf in r.get("confirm") or []:
            try:
                ci = column_index(cols, str(cf.get("path") or ""))
            except ValueError:
                ci = None
            if ci:
                cell = ws.cell(row=rn, column=ci)
                cell.fill = YELLOW
                text = confirm_text(cf) if confirm_text else str(cf.get("reason") or "")
                prev = cell.comment.text + "\n" if cell.comment else ""
                cell.comment = Comment((prev + text)[:2000], "kolis_tool")
                nconfirm += 1
    for name in wb.sheetnames:
        if name not in ("Sample", "Contents"):
            del wb[name]
    wb.save(out_xlsx)
    return {"rows": len(rows), "confirm": nconfirm, "out": str(out_xlsx), "columns": len(cols), "headers": [c.header for c in cols],
            "unknown": sorted({c.header for c in cols if c.unknown})}


# ---------------------------------------------------------------- 읽기(역직렬화): 반입용 엑셀 → 행 트리
def read_sheet(xlsx: Path) -> list[dict]:
    """Contents 시트를 행 트리로. 부모 열로 묶음 경계를 잡고, 부모 열이 없는 반복 요소(발행처·주기·식별기호·원문주소)는 같은 묶음 안에서 반복으로 본다."""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    try:
        ws = wb["Contents"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [(h or "").strip() if isinstance(h, str) else "" for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if not any(v not in (None, "") for v in raw):
            continue
        out.append(_row_from(headers, list(raw)))
    return out


def _row_from(headers: list[str], raw: list) -> dict:
    mods: dict = {}
    extra: dict = {}
    # 현재 열린 묶음(경로 → 노드) 스택
    open_nodes: dict[str, dict] = {}
    for h, v in zip(headers, raw):
        if not h:
            continue
        val = None if v in (None, "") else (str(v).strip() if not isinstance(v, (int, float)) else (str(int(v)) if float(v).is_integer() else str(v)))
        if h in EXTRA:
            if val is not None:
                extra[h] = val
            continue
        if not h.startswith("/mods"):
            continue
        is_parent = h in PARENT_COLS.values()
        path = h.rstrip("/") if is_parent else h
        attr = None
        m = re.fullmatch(r"(.*)\[@(\w+)\]", path)
        if m:
            path, attr = m.group(1), "@" + m.group(2)
            if attr == "@altType":
                attr = "@type"
        parts = path[len("/mods/"):].split("/") if path != "/mods" else []
        if is_parent:
            node = _open_group(mods, parts, open_nodes, path, new=True)
            continue
        if val is None:
            continue
        node = _open_group(mods, parts[:-1], open_nodes, "/mods/" + "/".join(parts[:-1]) if parts[:-1] else "/mods", new=False) if parts[:-1] else mods
        leaf = parts[-1]
        leaf_path = path
        parent_path = "/mods/" + "/".join(parts[:-1]) if parts[:-1] else "/mods"
        if not attr and leaf_path not in REPEAT and leaf in node and parent_path in REPEAT:
            # 부모 열 없이 같은 요소가 또 나오면(첫 출처정보의 placeTerm 둘) 새 묶음
            node = _open_group(mods, parts[:-1], open_nodes, parent_path, new=True)
        if attr:
            # 속성: 그 요소(마지막 같은 이름)에 붙인다. 요소 열이 아직 없으면(부모 속성) 부모 노드에
            target = _last_leaf(node, leaf, leaf_path) if not (leaf_path in PARENT_COLS or leaf_path in REPEAT and leaf not in node) else None
            if leaf_path in PARENT_COLS or (leaf_path in REPEAT and leaf not in node and leaf_path not in ("/mods/originInfo/publisher", "/mods/note", "/mods/identifier", "/mods/location/url", "/mods/classification")):
                target = _open_group(mods, parts, open_nodes, leaf_path, new=False)
            if target is None:
                target = _last_leaf(node, leaf, leaf_path, create=True)
            target[attr] = val
        else:
            if leaf_path in REPEAT and leaf_path not in PARENT_COLS:
                node.setdefault(leaf, [])
                node[leaf].append({"_": val})
            elif leaf_path in PARENT_COLS:
                g = _open_group(mods, parts, open_nodes, leaf_path, new=False)
                g["_"] = val
            else:
                node[leaf] = {"_": val}
    return {"mods": _simplify(mods), "extra": extra}


def _open_group(mods: dict, parts: list[str], open_nodes: dict, path: str, new: bool) -> dict:
    """경로의 묶음 노드를 연다. new=True 면 새 묶음(부모 열)을 만든다."""
    if not parts:
        return mods
    parent = _open_group(mods, parts[:-1], open_nodes, "/mods/" + "/".join(parts[:-1]) if parts[:-1] else "/mods", new=False)
    name = parts[-1]
    full = "/mods/" + "/".join(parts)
    if full in REPEAT:
        lst = parent.setdefault(name, [])
        if new or not lst:
            lst.append({})
            open_nodes[full] = lst[-1]
        return lst[-1]
    return parent.setdefault(name, {})


def _last_leaf(node: dict, leaf: str, leaf_path: str, create: bool = False):
    v = node.get(leaf)
    if isinstance(v, list):
        return v[-1] if v else (v.append({}) or v[-1]) if create else None
    if isinstance(v, dict):
        return v
    if v is None and create:
        node[leaf] = {}
        return node[leaf]
    return None


def _simplify(node):
    """{"_": "x"} 만 있는 노드는 "x" 로, 빈 묶음은 그대로."""
    if isinstance(node, list):
        return [_simplify(x) for x in node]
    if isinstance(node, dict):
        if set(node) == {"_"}:
            return node["_"]
        return {k: _simplify(v) for k, v in node.items()}
    return node


# ---------------------------------------------------------------- 비교(시험·대조용)
def value_cells(rows: list[dict]) -> set[tuple]:
    """행마다 값이 있는 셀의 (행 번호, 열 경로, 같은 경로 중 n번째, 값) 집합. 머리글 모양이 달라도 값 비교가 되게."""
    out = set()
    cols = layout(rows)
    for rn, r in enumerate(rows):
        seen: dict[str, int] = {}
        for c in cols:
            seen[c.header] = seen.get(c.header, 0) + 1
            v = cell_value(r, c)
            if v is not None:
                out.add((rn, c.header, seen[c.header], str(v)))
    return out


def sheet_value_cells(xlsx: Path, skip: set[str] = frozenset()) -> set[tuple]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    try:
        ws = wb["Contents"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    headers = [(h or "").strip() if isinstance(h, str) else "" for h in rows[0]]
    out = set()
    for rn, raw in enumerate(rows[1:]):
        if not any(v not in (None, "") for v in raw):
            continue
        seen: dict[str, int] = {}
        for h, v in zip(headers, raw):
            if not h:
                continue
            seen[h] = seen.get(h, 0) + 1
            if v in (None, "") or h in skip:
                continue
            s = str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip()
            out.add((rn, h, seen[h], s))
    return out


def dump(rows: list[dict]) -> str:
    return json.dumps(rows, ensure_ascii=False, indent=1)


# ---------------------------------------------------------------- 확인 표시(노란 칸·메모)
def marks(xlsx: Path) -> int:
    """반입용 엑셀에 남아 있는 확인 표시(노란색 칸·메모)의 수."""
    wb = openpyxl.load_workbook(xlsx)
    try:
        return sum(1 for row in wb["Contents"].iter_rows(min_row=2) for c in row if c.comment or (c.fill and c.fill.fill_type == "solid" and c.fill.fgColor.rgb in ("FFFFFF00", "00FFFF00")))
    finally:
        wb.close()


def clear_marks(xlsx: Path) -> int:
    """직원이 확인을 끝냈을 때: 프로그램이 넣은 확인 표시(노란색 칸·메모)를 지운다. 값은 건드리지 않는다. 완료 사례의 반입용 엑셀에는 색·메모가 없다."""
    wb = openpyxl.load_workbook(xlsx)
    n = 0
    for row in wb["Contents"].iter_rows(min_row=2):
        for c in row:
            yellow = c.fill and c.fill.fill_type == "solid" and c.fill.fgColor.rgb in ("FFFFFF00", "00FFFF00")
            if c.comment or yellow:
                n += 1
                c.comment = None
                if yellow:
                    c.fill = PatternFill(fill_type=None)
    wb.save(xlsx)
    return n
