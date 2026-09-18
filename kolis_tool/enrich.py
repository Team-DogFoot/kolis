"""② 기초메타데이터(출판사용) 보완 — 웹 리서치 에이전트.

출판사가 보낸 기초메타데이터는 보통 '최초 연재 플랫폼·주제 구분·저자·발행연월일·정가·소개·썸네일' 이 비어 온다.
  research : 작품 단위로 claude -p(WebSearch/WebFetch) 를 돌려 작품 정보 JSON 을 만든다.  (KOLIS 접근 없음)
  apply    : 그 JSON 으로 빈 칸을 채워 새 xlsx 를 쓴다. 채운 셀은 노란색 + 메모(근거). 원본 시트는 '_원본' 으로 보존.

규칙(작업자 완료 사례 5차-190/201/206, 유저 확정 2026-09-18):
  - 발행연월일 = **회차별로**, 여러 플랫폼 중 그 회차가 가장 이르게 공개된 날짜(유저 확정 2026-09-18).
    회차 날짜를 못 찾으면 비워 두고 노란색(사람 확인). 참고: 작업자 완료 사례 3건은 전 회차를 첫 회차 날짜 하나로 통일했었음.
  - 저자 표기 = '글: 이름; 그림: 이름' (스토리숲 사례 형식). 역할어는 글/그림/원작/작화 등 매뉴얼 2장 범위만.
  - 썸네일 = '<제목 공백제거><회차 2자리>.jpg' (thumbs 명령이 같은 규칙으로 파일명을 바꾼다).
  - 플랫폼에서만 확인한 값은 convert 단계에서 각괄호 규칙을 적용하므로 여기서는 값 그대로 적고 근거만 남긴다.
"""
from __future__ import annotations
import json, re, shutil, subprocess, threading, time
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

YELLOW = PatternFill("solid", fgColor="FFFF00")

RESEARCH_PROMPT = """당신은 국립중앙도서관 웹툰 납본 사업의 서지 조사 보조자입니다. 아래 작품의 정보를 웹에서 찾아 JSON 으로만 답하세요.
KOLIS(도서관 시스템)에는 절대 접근하지 마세요.

작품: {title}
출판사: {publisher}
ISBN: {isbn}
회차 수(납본 파일 기준): {n}
출판사 기재 최초 연재 플랫폼: {platform_hint}
출판사 기재 이용대상: {audience_hint}

{editable}
{extra}
출력 형식: 마지막에 ```json 코드블록 하나로 아래 스키마의 JSON 만.
{{"title":"", "platform_first":"", "platforms":[{{"name":"", "url_main":"", "url_work":"", "start_date":"YYYYMMDD 또는 빈칸"}}],
 "authors":[{{"name":"", "role":"글|그림|원작|작화|각색"}}], "genre":"", "rating":"전체|15세|19세", "adult":false,
 "first_publish_date":"YYYYMMDD", "episodes":[{{"no":1, "date":"YYYYMMDD", "title":"", "price":"", "source":"URL"}}],
 "completed":true, "total_episodes":0, "price_per_episode":"", "free_episodes":0,
 "summary":"", "publisher_place":"출판사 소재 시/도(예: 서울) 또는 빈칸", "evidence":["URL — 무엇을 확인했는지"], "notes":"확신이 낮은 항목과 이유"}}
모르는 값은 빈 문자열/0/false 로 두고 notes 에 적으세요. 지어내지 마세요."""

# 편집 가능한 부분: 어디를 보고 무엇을 찾을지. 역할·KOLIS 금지·출력 형식·지어내지 말 것은 고정(위 RESEARCH_PROMPT).
EDITABLE_DEFAULT = """볼 곳: 플랫폼(카카오페이지·네이버시리즈·리디·봄툰·코미카·레진·카카오웹툰 등)·출판사 홈페이지·웹툰가이드·나무위키.

찾을 것:
1. 최초 연재 플랫폼 이름과, 확인되는 모든 서비스 플랫폼(이름, 플랫폼 메인 URL, 이 작품 페이지 URL).
2. 저자: 이름과 역할(글/그림/원작/작화/각색 중 하나). 어시스트·콘티·채색·편집은 제외.
3. 연재 시작일: 플랫폼별로 확인되는 첫 회차 공개일(YYYYMMDD). 그중 **가장 이른** 날짜를 first_publish_date 에.
   그리고 **회차별 공개일**: 회차 목록이 보이는 플랫폼(카카오페이지·네이버시리즈·리디·코미카 등)에서 회차 번호별 공개일을 최대한 모아 episodes 에.
   같은 회차가 여러 플랫폼에 있으면 가장 이른 날짜. 못 찾은 회차는 넣지 마세요(지어내지 말 것).
4. 장르(주제 구분: 로맨스/BL/판타지/소년/드라마 등 플랫폼 표기), 이용등급(전체/15세/19세·성인), 완결 여부, 총 회차.
5. 회차 가격(원 또는 코인/캐시 단위, 무료 회차 수).
6. 작품 소개 문구(플랫폼 원문 그대로, 200자 이내).
7. 각 값의 근거 URL."""


def norm(s) -> str:
    return re.sub(r"\s+", "", str(s or ""))


def read_sheet(path: Path):
    wb = openpyxl.load_workbook(path)
    ws = wb["작성"] if "작성" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and any(norm(c) == "No" for c in r))
    header = [norm(c) for c in rows[hi]]
    col = {h: i + 1 for i, h in enumerate(header) if h}
    data_rows = [i + 1 for i in range(hi + 1, len(rows)) if rows[i] and rows[i][col["제목(도서명)"] - 1]]
    return wb, ws, col, data_rows


NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)   # 창 없는 실행(pythonw)에서 콘솔이 튀어나오지 않게


class Cancelled(Exception):
    pass


def claude_exe() -> str | None:
    return shutil.which("claude") or shutil.which("claude.cmd") or shutil.which("claude.exe")


def friendly_error(stderr: str, returncode: int) -> str:
    s = (stderr or "").lower()
    if "not logged in" in s or "login" in s or "authentication" in s or "unauthorized" in s or "401" in s:
        return "클로드코드에 로그인이 되어 있지 않습니다. 터미널에서 `claude` 를 한 번 실행해 로그인한 뒤 다시 시도하세요."
    if "enotfound" in s or "econnrefused" in s or "fetch failed" in s or "network" in s or "certificate" in s:
        return "네트워크 또는 인증서 문제로 클로드가 외부에 접속하지 못했습니다(보안 에이전트 확인). 원문: " + stderr[-300:]
    if "rate limit" in s or "overloaded" in s or "429" in s:
        return "클로드 사용량 한도에 걸렸습니다. 잠시 뒤 다시 시도하세요."
    return f"클로드 실행 실패(코드 {returncode}): {stderr[-400:]}"


PROMPT_OVERRIDE = Path("work") / "research_prompt.txt"   # 프로그램에서 편집한 '볼 곳·찾을 것' 부분(있으면 이걸 씀)


def current_editable() -> str:
    if PROMPT_OVERRIDE.exists():
        return PROMPT_OVERRIDE.read_text(encoding="utf-8")
    return EDITABLE_DEFAULT


def save_editable(text: str | None) -> str:
    """편집 가능한 부분만 저장. None/빈 값이면 기본값으로 되돌림."""
    if not text or not text.strip():
        if PROMPT_OVERRIDE.exists():
            PROMPT_OVERRIDE.unlink()
        return EDITABLE_DEFAULT
    if "{" in text or "}" in text:
        raise ValueError("중괄호 { } 는 쓸 수 없습니다(자리표시자 충돌)")
    PROMPT_OVERRIDE.parent.mkdir(parents=True, exist_ok=True)
    PROMPT_OVERRIDE.write_text(text, encoding="utf-8")
    return text


def build_prompt(title, publisher, isbn, n, platform_hint, audience_hint, extra: str = "") -> str:
    ex = f"\n이번 작품 추가 지시(우선 적용):\n{extra.strip()}\n" if extra and extra.strip() else ""
    return RESEARCH_PROMPT.format(title=title, publisher=publisher, isbn=isbn, n=n, platform_hint=platform_hint,
                                  audience_hint=audience_hint, editable=current_editable().strip(), extra=ex)


def research(pub_xlsx: Path, out_json: Path, runner: str = "claude", timeout: int = 900, log=None, handle: dict | None = None,
             extra: str = "") -> dict:
    """헤드리스 클로드로 작품 정보 JSON 을 만든다. log(msg) 로 진행(검색어·읽는 URL)을 실시간 보고. handle['proc'] 에 프로세스를 두어 취소 가능.
    extra: 이번 작품에만 붙이는 추가 지시(저장하지 않음)."""
    log = log or (lambda m: None)
    wb, ws, col, rows = read_sheet(pub_xlsx)
    g = lambda r, k: ws.cell(row=r, column=col[k]).value if k in col else None
    r0 = rows[0]
    if PROMPT_OVERRIDE.exists():
        log("편집된 조사 지시 사용(work/research_prompt.txt)")
    if extra and extra.strip():
        log("이번 작품 추가 지시 포함")
    prompt = build_prompt(g(r0, "제목(도서명)"), g(r0, "출판사"), g(r0, "ISBN/UCI"), len(rows),
                          g(r0, "최초연제플랫폼") or g(r0, "최초연재플랫폼") or "(비어 있음)", g(r0, "이용대상") or "(비어 있음)", extra)
    exe = claude_exe() if runner == "claude" else (shutil.which(runner) or shutil.which(runner + ".cmd"))
    if not exe:
        raise SystemExit("클로드코드(claude)가 설치되어 있지 않거나 PATH 에 없습니다." if runner == "claude" else f"{runner} 실행파일을 찾지 못함")
    if runner == "claude":
        # 모델: 검색·정리 작업이라 소넷으로 충분(빠르고 저렴). 환경변수 KOLIS_RESEARCH_MODEL 로 바꿀 수 있다
        import os
        cmd = [exe, "-p", "--verbose", "--output-format", "stream-json", "--model", os.environ.get("KOLIS_RESEARCH_MODEL", "sonnet"),
               "--allowedTools", "WebSearch", "WebFetch"]
    else:
        cmd = [exe, "exec", "--skip-git-repo-check", "-"]
    proc = subprocess.Popen(cmd, stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8",
                            errors="replace", creationflags=NO_WINDOW)
    if handle is not None:
        handle["proc"] = proc
    proc.stdin.write(prompt); proc.stdin.close()
    text, lines = "", []
    t0 = time.time()
    err_buf: list[str] = []
    threading.Thread(target=lambda: err_buf.append(proc.stderr.read()), daemon=True).start()
    for line in proc.stdout:
        lines.append(line)
        if runner != "claude":
            continue
        try:
            ev = json.loads(line)
        except json.JSONDecodeError:
            continue
        if ev.get("type") == "assistant":
            for b in (ev.get("message") or {}).get("content", []):
                if b.get("type") == "tool_use":
                    inp = b.get("input") or {}
                    what = inp.get("query") or inp.get("url") or ""
                    log(f"  [{int(time.time() - t0)}s] {'검색' if b.get('name') == 'WebSearch' else '페이지 읽기'}: {str(what)[:90]}")
                elif b.get("type") == "text" and b.get("text", "").strip():
                    log(f"  [{int(time.time() - t0)}s] 클로드: {b['text'].strip()[:120]}")
        elif ev.get("type") == "result":
            text = ev.get("result") or ""
    rc = proc.wait(timeout=timeout)
    if handle is not None and handle.get("cancel"):
        raise Cancelled("사용자가 중단함")
    if rc != 0:
        raise SystemExit(friendly_error("".join(err_buf), rc))
    if runner != "claude":
        text = "".join(lines)
    m = re.findall(r"```json\s*(\{.*?\})\s*```", text, re.S)
    if not m:
        (out_json.with_suffix(".raw.txt")).write_text(text, encoding="utf-8")
        raise SystemExit(f"클로드 답변에서 JSON 을 찾지 못했습니다 → {out_json.with_suffix('.raw.txt')} 확인")
    data = json.loads(m[-1])
    data["_raw"] = text
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    log(f"  리서치 완료 ({int(time.time() - t0)}초)")
    return data


ROLE_MAP = {"글": "글", "그림": "그림", "원작": "원작", "작화": "작화", "각색": "각색", "스토리": "글", "작가": ""}


def _domain_main(url: str) -> str:
    m = re.match(r"(https?://[^/]+)", url or "")
    return (m.group(1) + "/") if m else ""


def from_cliptoon(paths: list[Path]) -> dict:
    """ClipToon V8(직원 제작 도구) '엑셀 다운로드' 파일들 → research JSON 과 같은 구조.
    시트: 작품정보(제목·플랫폼·저자·저자별 역할어·1회 발행일·출판사·이용등급·링크), 회차정보(회차번호·회차명·발행일·대여/소장 가격·무료 여부·회차 링크).
    파일이 여러 개(플랫폼별)면 회차별 발행일은 가장 이른 것, 최초 연재 플랫폼은 1회 발행일이 가장 이른 플랫폼."""
    plats, episodes, authors, evidence = [], {}, [], []
    rating, publisher_name, title = "", "", ""
    for p in paths:
        wb = openpyxl.load_workbook(p, data_only=True)
        w = dict(zip([str(c or "") for c in next(wb["작품정보"].iter_rows(min_row=1, max_row=1, values_only=True))],
                     next(wb["작품정보"].iter_rows(min_row=2, max_row=2, values_only=True))))
        title = title or str(w.get("제목") or "")
        d = re.sub(r"\D", "", str(w.get("1회 발행일") or ""))
        plats.append({"name": str(w.get("플랫폼") or ""), "url_main": _domain_main(str(w.get("링크") or "")),
                      "url_work": str(w.get("링크") or ""), "start_date": d, "file": p.name})
        evidence.append(f"{w.get('링크')} — ClipToon {p.name}")
        if not authors:
            for part in re.split(r"[,;]\s*", str(w.get("저자별 역할어") or "")):
                m = re.match(r"(.+?)\((.+?)\)\s*$", part.strip())
                if m:
                    name, role = m.group(1).strip(), m.group(2).strip()
                    authors.append({"name": name, "role": ROLE_MAP.get(role, "" if "확인" in role else role)})
                elif part.strip():
                    authors.append({"name": part.strip(), "role": ""})
        r = str(w.get("이용등급") or "")
        if r and "확인" not in r and (not rating or "19" in r or "성인" in r):
            rating = r
        pub = str(w.get("출판사") or "")
        if pub and "확인" not in pub:
            publisher_name = publisher_name or pub
        ws = wb["회차정보"]
        hdr = [str(c or "") for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            e = dict(zip(hdr, row))
            name = str(e.get("회차명") or "")
            m = re.search(r"(?<![\d])(\d{1,3})\s*화", name)
            if not m or re.search(r"외전|프롤로그|트레일러|예고|후기|특별", name):
                continue                      # 본편 'N화' 만 회차 번호로 인정. 외전·트레일러는 제외(사람 확인)
            no = int(m.group(1))
            d = re.sub(r"\D", "", str(e.get("발행일") or ""))
            price = ""
            for k in ("소장 원화 가격", "대여 원화 가격", "소장 원본 가격", "대여 원본 가격"):
                v = str(e.get(k) or "")
                if v and "확인" not in v:
                    price = re.sub(r"\D", "", v) or v; break
            cur = episodes.get(no)
            if cur is None or (len(d) == 8 and (not cur["date"] or d < cur["date"])):
                episodes[no] = {"no": no, "date": d, "title": name, "price": price or (cur or {}).get("price", ""),
                                "free": str(e.get("무료 여부") or ""), "source": f"{w.get('플랫폼')} {e.get('회차 링크') or ''}"}
            elif cur is not None and not cur.get("price") and price:
                cur["price"] = price
    dated = [p for p in plats if p["start_date"]]
    first = min(dated, key=lambda p: p["start_date"]) if dated else (plats[0] if plats else {})
    paid = [e["price"] for e in episodes.values() if e.get("price") and e.get("free") != "무료"]
    price_per = max(set(paid), key=paid.count) if paid else ""
    adult = "19" in rating or "성인" in rating
    return {"title": title, "platform_first": first.get("name", ""), "platforms": plats, "authors": authors,
            "genre": "", "rating": rating, "adult": adult, "first_publish_date": first.get("start_date", ""),
            "episodes": [episodes[k] for k in sorted(episodes)], "completed": any("완결" in e["title"] for e in episodes.values()),
            "total_episodes": len(episodes), "price_per_episode": price_per,
            "free_episodes": sum(1 for e in episodes.values() if e.get("free") == "무료"), "summary": "",
            "publisher_place": "", "publisher_name": publisher_name, "evidence": evidence,
            "notes": "ClipToon 내보내기 기반. 장르·소개·발행지는 ClipToon 에 없어 비움(사람 또는 research 로 보완)"}


def authors_text(authors: list[dict]) -> str:
    return "; ".join(f"{a.get('role', '')}: {a.get('name', '')}".strip(": ") for a in authors if a.get("name"))


def price_number(raw) -> str:
    """'리디 200원(대여, 4~50화)' → '200', '500원' → '500', '무료' → '', 숫자만이면 그대로. 문장에서 '원' 앞 숫자를 우선."""
    s = str(raw or "").strip()
    if not s:
        return ""
    m = re.search(r"(\d[\d,]*)\s*원", s) or re.search(r"(\d[\d,]*)", s)
    return m.group(1).replace(",", "") if m else ""


def thumb_name(title: str, i: int, ext: str = ".jpg") -> str:
    return f"{norm(title)}{i:02d}{ext}"


def apply(pub_xlsx: Path, info: dict, out_xlsx: Path, thumbs_dir: Path | None = None) -> dict:
    wb, ws, col, rows = read_sheet(pub_xlsx)
    # 원본 보존
    if "_원본" not in wb.sheetnames:
        src = wb.copy_worksheet(ws); src.title = "_원본"
    filled: dict[str, int] = {}
    ev = "; ".join(info.get("evidence", [])[:5])

    def put(r, key, value, why):
        if key not in col or value in (None, ""):
            return
        c = ws.cell(row=r, column=col[key])
        if c.value not in (None, ""):
            return
        c.value = value; c.fill = YELLOW
        c.comment = Comment(f"enrich(웹 리서치): {why}\n근거: {ev}"[:2000], "kolis_tool")
        filled[key] = filled.get(key, 0) + 1

    first = re.sub(r"\D", "", str(info.get("first_publish_date") or ""))
    ep_date: dict[int, str] = {}
    ep_price: dict[int, str] = {}
    for e in info.get("episodes", []) or []:
        try:
            no = int(e.get("no"))
        except (TypeError, ValueError):
            continue
        d = re.sub(r"\D", "", str(e.get("date") or ""))
        if len(d) == 8 and (no not in ep_date or d < ep_date[no]):
            ep_date[no] = d
        if e.get("price") not in (None, ""):
            ep_price[no] = str(e["price"])
    authors = authors_text(info.get("authors", []))
    price = price_number(info.get("price_per_episode"))
    ep_price = {k: price_number(v) for k, v in ep_price.items() if price_number(v)}
    for i, r in enumerate(rows, 1):
        put(r, "최초연제플랫폼", info.get("platform_first"), "플랫폼 조사")
        put(r, "최초연재플랫폼", info.get("platform_first"), "플랫폼 조사")
        put(r, "주제구분", info.get("genre"), "플랫폼 장르 표기")
        put(r, "저자", authors, "플랫폼·출판사 크레딧. 역할어는 매뉴얼 2장 범위")
        try:
            no = int(re.sub(r"\D", "", str(ws.cell(row=r, column=col["권차"]).value or "")) or i)
        except ValueError:
            no = i
        if no in ep_date:
            d = min(ep_date[no], first) if (no == 1 and len(first) == 8) else ep_date[no]   # 1회차는 '가장 이른 연재 시작일' 도 후보
            put(r, "발행연월일", d, f"{no}회차 공개일, 여러 플랫폼 중 가장 이른 날짜(유저 규칙 2026-09-18)")
        elif no == 1 and len(first) == 8:
            put(r, "발행연월일", first, "1회차 = 가장 이른 연재 시작일")
        else:
            c = ws.cell(row=r, column=col["발행연월일"])
            if c.value in (None, ""):
                c.fill = YELLOW; c.comment = Comment(f"enrich: {no}회차 공개일을 웹에서 못 찾음. 원문·플랫폼에서 확인", "kolis_tool")
        put(r, "정가", ep_price.get(no, price), "플랫폼 회차 가격(무료 회차도 유료 회차 가격으로 적는 것이 완료 사례 관행)")
        put(r, "도서정보및소개", info.get("summary"), "플랫폼 작품 소개")
        put(r, "썸네일", thumb_name(str(ws.cell(row=r, column=col["제목(도서명)"]).value), i), "thumbs 명령 파일명 규칙")
        if info.get("adult") or str(info.get("rating", "")).startswith("19"):
            c = ws.cell(row=r, column=col["이용대상"])
            c.fill = YELLOW
            c.comment = Comment(f"enrich: 플랫폼 이용등급 '{info.get('rating')}' — 출판사 기재 '{c.value}' 와 다름. 성인용이면 convert 가 성인용/19세 미만 구독불가/공개 1 로 씀", "kolis_tool")
    # 근거 시트
    if "근거" in wb.sheetnames:
        del wb["근거"]
    g = wb.create_sheet("근거")
    g.append(["항목", "값"])
    for k, v in info.items():
        if k == "_raw":
            continue
        g.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v])
    wb.save(out_xlsx)
    return filled


def rename_thumbs(folder: Path, title: str, dry_run: bool = False) -> list[tuple[str, str]]:
    """회차 썸네일을 '<제목 공백제거><NN>.jpg' 로. 원래 순서는 자연 정렬. manifest 로 undo 가능."""
    from .common import natural_key, list_images
    files = list_images(folder)
    pairs = [(p.name, thumb_name(title, i, p.suffix.lower())) for i, p in enumerate(files, 1)]
    if dry_run:
        return pairs
    mf = folder / "thumbs_manifest.json"
    if mf.exists():
        raise SystemExit(f"이미 바꾼 폴더(manifest 있음): {mf}")
    for a, b in pairs:
        if a != b:
            (folder / a).rename(folder / b)
    mf.write_text(json.dumps(pairs, ensure_ascii=False, indent=1), encoding="utf-8")
    return pairs


def undo_thumbs(folder: Path) -> int:
    mf = folder / "thumbs_manifest.json"
    pairs = json.loads(mf.read_text(encoding="utf-8"))
    n = 0
    for a, b in reversed(pairs):
        if a != b and (folder / b).exists():
            (folder / b).rename(folder / a); n += 1
    mf.unlink()
    return n
