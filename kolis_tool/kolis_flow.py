"""반입용 엑셀 확인(사람) 뒤의 KOLIS 등록을 처음부터 끝까지 한 번에, **브라우저 없이**:
로그인 → 일괄반입 → 콘텐츠ID 받기 → 원고 폴더명 → 원문일괄등록(전송·정보입력·원문등록) → (썸네일) → 등록대상처리 → 가원부번호 → 가원부 파일.

- 매번 처음부터 한다(이어서 하기 없음, 2026-09-30 유저 확정). 같은 작품을 다시 실행하면 새 접수번호가 생긴다.
- 실행한 접수번호는 전부 `work/취소요청_목록.csv` 에 남는다. 작품마다 끝까지 성공한 마지막 접수번호만 '유지', 나머지는 '취소 요청'이다.
  이 목록을 주무관에게 보내 취소를 요청한다(사람이 함).
- 단계마다 실행 전 조건을 확인하고, 끝난 뒤 KOLIS 의 값을 다시 읽어 맞춰 본다. 하나라도 다르면 그 자리에서 멈춘다.
- 썸네일: 납품에 동봉돼 있으면 등록하고(화면 방식, Edge 필요), 없으면 하지 않는다.
- 전부 기록한다(journal): 조건과 값, 단계의 시작·끝·시간, 화면 로그, 보낸 요청과 받은 응답, 멈춘 이유.
"""
from __future__ import annotations
import csv, datetime, re, threading, time
from pathlib import Path

STEPS = [("login", "로그인"), ("import", "일괄반입"), ("export", "콘텐츠ID 받기"), ("cnts", "원고 폴더명 → 콘텐츠ID"), ("upload", "원문일괄등록"),
         ("thumbs", "썸네일 등록"), ("register", "등록대상처리 → 가원부번호 → 가원부 파일")]
LEDGER = "취소요청_목록.csv"
COLUMNS = ["처리", "작품", "접수번호", "가원부번호", "건수", "콘텐츠ID", "비고", "결과", "반입 시각", "실행 기록"]
KEEP, CANCEL = "유지", "취소 요청"
_ledger_lock = threading.Lock()


class Stop(RuntimeError):
    pass


def ledger_read(work_dir: Path) -> list[dict]:
    p = Path(work_dir) / LEDGER
    if not p.exists():
        return []
    with open(p, encoding="utf-8-sig", newline="") as f:
        return [{c: r.get(c, "") for c in COLUMNS} for r in csv.DictReader(f)]


def ledger_put(work_dir: Path, row: dict) -> None:
    """접수번호 한 건을 적거나 고친다. 끝까지 성공한 건이 생기면 같은 작품의 앞선 건은 전부 '취소 요청'이 된다."""
    with _ledger_lock:
        rows = ledger_read(work_dir)
        old = next((r for r in rows if r["접수번호"] == str(row["접수번호"])), None)
        if old:
            old.update({k: str(v) for k, v in row.items()})
        else:
            rows.append({c: str(row.get(c, "")) for c in COLUMNS})
        if row.get("처리") == KEEP:
            for r in rows:
                if r["작품"] == row["작품"] and r["접수번호"] != str(row["접수번호"]):
                    r["처리"] = CANCEL
        p = Path(work_dir) / LEDGER
        with open(p, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=COLUMNS)
            w.writeheader(); w.writerows(rows)


def _folders(root: Path) -> list[Path]:
    return [p for p in Path(root).iterdir() if p.is_dir() and not p.name.startswith("_kolis")]


def _thumb_names(xlsx: Path) -> list[str]:
    """반입용 엑셀의 thum_files 칸을 행 순서대로."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        rows = [r for r in wb["Contents"].iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
    finally:
        wb.close()
    ti = [str(v or "") for v in rows[0]].index("thum_files")
    return [str(r[ti] or "").strip() for r in rows[1:]]


def _excel_rows(xlsx: Path) -> dict:
    """반입용 엑셀의 데이터 행 수와 본표제(반입 뒤 KOLIS 목록과 맞춰 볼 값)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        names = list(wb.sheetnames)
        ws = wb["Contents"] if "Contents" in names else wb.worksheets[-1]
        rows = [r for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
    finally:
        wb.close()
    head = [str(v or "") for v in rows[0]]
    ti = next((i for i, h in enumerate(head) if h.startswith("/mods/titleInfo/title")), None)
    return {"rows": len(rows) - 1, "sheets": names, "titles": sorted({str(r[ti]).strip() for r in rows[1:] if ti is not None and r[ti]})}


def run(work: dict, note: str, yes: str, work_dir: Path, log, handle: dict | None = None, progress=None, thumb_progress=None, save=None) -> dict:
    """work: 1단계 결과(title, output_xlsx, rows, manuscripts, thumbs). save(patch): 작품 상태 기록(화면 표시용)."""
    from . import kolis_http, kolis_request as req, cnts_folders
    from .ids_from_export import receipt_map
    from .journal import Journal
    handle = handle or {}
    progress = progress or (lambda key, status, text="": None)
    save = save or (lambda patch: None)
    work_dir = Path(work_dir)
    title = str(work.get("title") or Path(str(work.get("folder") or "작품")).name)
    jr = Journal(work_dir, title)
    log = jr.wrap(log)
    year = str(datetime.date.today().year)
    client = kolis_http.Client(log, jr)
    t_all = time.time()
    checks: list = []
    receipt = ""
    row: dict = {}

    def check(name: str, ok: bool, value="", why: str = ""):
        jr.write("check", name=name, ok=bool(ok), value=value)
        checks.append((name, bool(ok)))
        log(f"  확인 {'✓' if ok else '✗'} {name}: {value}")
        if not ok:
            raise Stop(f"{name} — {why or value}")

    def step(key: str, fn):
        label = dict(STEPS)[key]
        if handle.get("cancel"):
            raise Stop(f"사용자가 중단함({label} 전)")
        jr.step = key
        jr.write("step", status="start")
        log(f"■ {label}: 시작")
        progress(key, "run")
        t0 = time.time()
        try:
            text = fn() or ""
        except BaseException as e:
            jr.error(e, seconds=round(time.time() - t0, 1))
            jr.write("step", status="fail", text=str(e), seconds=round(time.time() - t0, 1))
            log(f"■ {label}: 멈춤 — {e}")
            progress(key, "fail", str(e)[:300])
            raise
        sec = round(time.time() - t0, 1)
        jr.write("step", status="done", text=text, seconds=sec)
        progress(key, "done", f"{text} ({sec:g}초)")
        log(f"■ {label}: 완료 {text} ({sec:g}초)")
        if receipt:
            ledger_put(work_dir, {**row, "결과": f"{label}까지 끝남"})

    try:
        # ---------- 시작 전 확인 ----------
        jr.step = "start"
        jr.write("input", note=note, year=year, work={k: work.get(k) for k in ("title", "folder", "output_xlsx", "rows", "manuscripts", "thumbs", "confirm_cells")})
        log(f"=== KOLIS 등록 시작: {title} · 브라우저 없음 · 기록 {jr.path}")
        check("직원 동의(YES)", yes == "YES", "입력됨" if yes == "YES" else "없음", "직원과 반입용 엑셀을 확인한 뒤 YES 를 넣어야 합니다")
        check("비고(차수·번호)", bool(note) and not re.search(r"\(차\)|\(\)", note), note, "비고에 차수와 번호를 넣어야 합니다(예: 2026-납본-웹툰대행(5차)(394))")
        xlsx = Path(work.get("output_xlsx") or "")
        root = Path(work.get("manuscripts") or "")
        check("반입용 엑셀 파일", xlsx.is_file(), str(xlsx), f"반입용 엑셀이 없습니다: {xlsx}")
        check("원고 폴더", root.is_dir(), str(root), f"원고 폴더가 없습니다: {root}")
        info = _excel_rows(xlsx)
        rows = info["rows"]
        check("반입용 엑셀의 시트", info["sheets"] == ["Sample", "Contents"], info["sheets"], "도서관 양식(Sample, Contents) 밖의 시트가 있습니다")
        check("반입용 엑셀 행 수", 0 < rows <= 500, rows, "행이 없거나 500행을 넘습니다(반입 엑셀 1개당 500건)")
        if work.get("rows"):
            check("엑셀 행 수 = 1단계가 만든 행 수", rows == int(work["rows"]), f"{rows} / {work['rows']}", "1단계 뒤에 엑셀의 행 수가 바뀌었습니다 → 직원 확인")
        # 지난 실행이 원고 폴더 이름을 콘텐츠ID 로 바꿔 놓았으면 원래 이름으로 되돌린다(새 접수번호의 콘텐츠ID 로 다시 바꿔야 한다)
        if any(p.name.startswith("CNTS-") for p in _folders(root)):
            n = cnts_folders.undo(root)
            log(f"  지난 실행이 바꾼 원고 폴더 이름 {n}개를 원래대로 되돌림")
            left = [p.name for p in _folders(root) if p.name.startswith("CNTS-")]
            check("원고 폴더 이름을 되돌림", not left, left[:3] or f"{n}개", "되돌리기 기록이 없어 폴더 이름을 원래대로 돌리지 못했습니다 → 폴더 이름을 직접 고치세요")
        check("원고 폴더 수 = 엑셀 행 수", len(_folders(root)) == rows, f"{len(_folders(root))} / {rows}")
        thumbs = Path(work["thumbs"]) if work.get("thumbs") and Path(work["thumbs"]).is_dir() and any(Path(work["thumbs"]).iterdir()) else None
        log(f"  썸네일: {'동봉됨 → 등록함 (' + str(thumbs) + ')' if thumbs else '동봉되지 않음 → 등록하지 않음'}")
        jr.write("branch", thumbs=bool(thumbs), path=str(thumbs or ""))
        for k in ("kolis_prepare", "kolis_submit", "export", "cnts", "upload", "thumbs_register", "register"):
            save({k: None})

        # ---------- 0. 로그인 ----------
        def do_login():
            try:
                client.login()
            except req.Stop as e:
                raise Stop(str(e)) from e
            return ""
        step("login", do_login)

        # ---------- 1. 일괄반입 ----------
        def do_import():
            nonlocal receipt, row
            try:
                r = req.import_excel(client, year, note, xlsx, rows, yes, log)
            except req.Stop as e:
                raise Stop(str(e)) from e
            receipt = r["receipt"]
            row = {"처리": CANCEL, "작품": title, "접수번호": receipt, "가원부번호": "", "건수": r["count"], "콘텐츠ID": f"{r['ids'][0]} ~ {r['ids'][-1]}",
                   "비고": note, "결과": "일괄반입", "반입 시각": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "실행 기록": jr.path.name}
            ledger_put(work_dir, row)
            save({"kolis_submit": {"ok": True, "message": r["message"], "receipt": receipt, "count": r["count"], "note": note}})
            check("반입된 본표제 = 엑셀의 본표제", not info["titles"] or set(r["titles"]) == set(info["titles"]), f"{r['titles']} / {info['titles']}",
                  "반입된 자료의 제목이 엑셀과 다릅니다 → 직원 확인")
            return f"접수번호 {receipt}, {r['count']}건"
        step("import", do_import)

        # ---------- 2. 콘텐츠ID 받기 ----------
        ex_file = Path()
        ids: list[str] = []

        def do_export():
            nonlocal ex_file, ids
            try:
                r = req.export_receipt(client, year, receipt, work_dir, log)
            except req.Stop as e:
                raise Stop(str(e)) from e
            check("받은 목록의 접수번호", str(r["receipt"]) == str(receipt), f"{r['receipt']} / {receipt}")
            check("KOLIS 목록 건수 = 엑셀 행 수", int(r["count"]) == rows, f"{r['count']} / {rows}")
            m = receipt_map(Path(r["file"]))
            ids = [x["cnts"] for x in m]
            check("콘텐츠ID 가 건마다 다름", len(set(ids)) == rows, len(set(ids)))
            check("접수번호 뒷번호가 1부터 이어짐", [x["seq"] for x in m] == list(range(1, rows + 1)), [x["seq"] for x in m][:5])
            ex_file = Path(r["file"])
            save({"export": {"file": r["file"], "receipt": r["receipt"], "count": r["count"]}})
            return f"{r['count']}건 ({r['first']} … {r['last']})"
        step("export", do_export)

        # ---------- 3. 폴더명 ----------
        def do_cnts():
            try:
                p = cnts_folders.plan(ex_file, root)
                jr.write("plan", pairs=p["pairs"])
                for a in p["pairs"]:
                    log(f"  {a[0]} → {a[1]} ({a[2]})")
                r = cnts_folders.apply(ex_file, root)
            except SystemExit as e:
                raise Stop(str(e)) from e
            after = sorted(p.name for p in _folders(root))
            check("바꾼 뒤 폴더 이름 = 콘텐츠ID", after == sorted(ids), f"{len(after)}개")
            save({"cnts": {"count": r["count"], "receipt": r["receipt"]}})
            return f"{r['count']}개 폴더"
        step("cnts", do_cnts)

        # ---------- 4. 원문일괄등록 ----------
        def do_upload():
            before = {i["id"]: i["files"] for i in req.contents(client, year, receipt)["items"]}
            check("원문이 아직 등록되지 않음", not any(before.values()), {k: v for k, v in before.items() if v} or "전 건 0")
            try:
                r = req.upload_folders(client, year, receipt, root, ids, log, handle, check)
            except req.Stop as e:
                raise Stop(str(e)) from e
            save({"upload": {"folders": r["folders"], "files": r["files"], "message": r["message"]}})
            return f"폴더 {r['folders']}개, 파일 {r['files']}개, {r['mb']}MB"
        step("upload", do_upload)

        # ---------- 5. 썸네일(동봉된 납품만) ----------
        def do_thumbs():
            # 짝은 반입용 엑셀의 행 순서로 짓는다: n 번째 행의 thum_files ↔ n 번째 콘텐츠ID(접수번호 뒷번호 순). 편/권차 글자로 맞추지 않는다(본편·외전의 숫자가 겹친다)
            names = _thumb_names(xlsx)
            check("엑셀의 thum_files 가 행마다 있음", len(names) == rows and all(names), [n for n in names if n][:3] or "없음",
                  "썸네일이 동봉된 납품인데 반입용 엑셀의 thum_files 가 비어 있는 행이 있습니다")
            check("thum_files 가 행마다 다름", len(set(names)) == rows, len(set(names)))
            pairs = [(cid, thumbs / n) for cid, n in zip(ids, names)]
            jr.write("thumb_pairs", pairs=[(c, f.name) for c, f in pairs])
            try:
                r = req.upload_thumbs(client, year, receipt, pairs, log, handle, check)
            except req.Stop as e:
                raise Stop(str(e)) from e
            jr.write("thumbs", results=r["items"])
            save({"thumbs_register": {"done": r["count"], "skipped": 0, "complete": True}})
            check("썸네일 건수 = 엑셀 행 수", r["count"] == rows, f"{r['count']} / {rows}")
            return f"등록 {r['count']}건"
        jr.step = "thumbs"
        if thumbs is None:
            jr.write("step", status="skip", text="썸네일 없음")
            log("■ 썸네일 등록: 썸네일이 동봉되지 않은 납품 — 하지 않음")
            progress("thumbs", "skip", "썸네일이 동봉되지 않음(등록하지 않음)")
        else:
            step("thumbs", do_thumbs)

        # ---------- 6. 등록대상처리 → 가원부번호 → 가원부 파일 ----------
        out: dict = {}

        def do_register():
            nonlocal out, row
            try:
                r = req.run(receipt, year, work_dir, log, handle, yes, client)
            except req.Stop as e:
                raise Stop(str(e)) from e
            out = {"record_no": r["record_no"], "file": r["file"], "count": r["count"], "receipt": receipt}
            row = {**row, "가원부번호": r["record_no"]}
            ledger_put(work_dir, row)
            save({"register": out})
            check("가원부 파일 건수 = 엑셀 행 수", int(r["count"]) == rows, f"{r['count']} / {rows}")
            check("가원부 파일의 콘텐츠ID = 반입한 콘텐츠ID", sorted(r["export"].get("ids") or []) == sorted(ids), f"{len(r['export'].get('ids') or [])}건")
            return f"가원부번호 {r['record_no']}, {r['count']}건"
        step("register", do_register)

        jr.step = "end"
        sec = round(time.time() - t_all, 1)
        row = {**row, "처리": KEEP, "결과": "끝까지 성공"}
        ledger_put(work_dir, row)
        jr.write("done", receipt=receipt, record_no=out.get("record_no"), file=out.get("file"), seconds=sec, requests=client.count, checks=len(checks))
        log(f"=== 전체 완료: 접수번호 {receipt}, 가원부번호 {out.get('record_no')} · {sec:g}초 · 확인 {len(checks)}개 통과 · 요청 {client.count}건 · 기록 {jr.path.name}")
        return {"receipt": receipt, "record_no": out.get("record_no"), "file": out.get("file"), "count": out.get("count"), "export_file": str(ex_file),
                "seconds": sec, "journal": str(jr.path), "checks": len(checks), "requests": client.count, "ledger": str(work_dir / LEDGER)}
    except BaseException as e:
        if jr.step not in dict(STEPS):
            jr.error(e)
        jr.write("stopped", reason=str(e), receipt=receipt, seconds=round(time.time() - t_all, 1), requests=client.count)
        if receipt:
            ledger_put(work_dir, {**row, "처리": CANCEL, "결과": f"{dict(STEPS).get(jr.step, jr.step)}에서 멈춤: {str(e)[:120]}"})
        log(f"=== 멈춤: {e}" + (f" · 접수번호 {receipt} 는 취소 요청 목록에 적음" if receipt else "") + f" · 기록 {jr.path}")
        if isinstance(e, Stop) or not isinstance(e, Exception):
            raise
        raise Stop(f"{type(e).__name__}: {e}") from e
    finally:
        client.close()
