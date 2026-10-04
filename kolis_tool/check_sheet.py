"""C 단계 점검·납품(프로세스 가이드 5.4~5.9) — 2026-10-04 새 설계. MODStoXL·정리 매크로를 대신한다.

  C-1 MODS 받기 : 원부의 콘텐츠 전부 → getHarContentsXml.do(요청) → work/wonbu/<원부>/xml/*.xml (받은 그대로)
  C-2 점검 시트 : XML → 트리 → 매크로 꼴 열(경로 '/mods/…', 같은 경로 n+1 번째는 '[n] /mods/…') → 가이드 5.7 세팅
                  (1행 오류 종류, 2행 데이터 개수 COUNTA, 3행 한글, 4행 경로, A열 콘텐츠ID, 필터, 틀 고정 A~C, 장르주제명 열은 분류기호 앞)
  C-3 점검      : 형식 규칙(코드) + 판단 규칙(에이전트 `check-mods`) → 오류 = 노란 배경·빨간 글자 + 1행에 오류 종류 / 확인 = 파란 배경(틀린 게 없으면 직원이 채우기 없음으로)
  C-4 확정·납품 : 직원이 확인(파란 칸 처리, 오류 답 기재) → 파일명 "2026-<원부> n차점검(오류없음|오류종류)" → 웹하드 업로드는 사람.
                  오류가 있으면 B-3 보정으로 돌아가 고치고 n+1 차 점검.
기록: work/wonbu/<원부>/check/<n차>/ (findings.json, sheet.xlsx)
"""
from __future__ import annotations
import datetime, json, re
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill, Font

from . import mods_xml, wonbu as wb
from .common import CODE_REGION, region_code_from_place

YELLOW = PatternFill("solid", fgColor="FFFF00"); BLUE = PatternFill("solid", fgColor="CCE5FF"); RED = Font(color="FF0000")
CONST = {"/mods/genre": "만화", "/mods/typeOfResource": "텍스트", "/mods/extension/regionOfPublishing": "한국", "/mods/originInfo/issuance": "단행자료",
         "/mods/language/languageTerm": "kor", "/mods/physicalDescription/form": "전자자료(Image)", "/mods/location/physicalLocation": "국립중앙도서관"}
NOTE_TYPES = ("target audience", "acquisition", "awards", "funding")


def cdir(wonbu: str, nth: int) -> Path:
    d = wb.wdir(wonbu) / "check" / f"{nth}차"; d.mkdir(parents=True, exist_ok=True); return d


def next_nth(wonbu: str) -> int:
    base = wb.wdir(wonbu) / "check"
    ns = [int(p.name[:-1]) for p in base.glob("*차") if p.name[:-1].isdigit()] if base.exists() else []
    return (max(ns) + 1) if ns else 1


# ---------------------------------------------------------------- C-2 시트 자료
def load_rows(wonbu: str) -> tuple[list[str], list[dict], list[dict]]:
    """xml 폴더의 파일을 콘텐츠 목록 순서로. 돌려주는 것: (열, 행(dict), 항목[{contents_id, vol, title}])."""
    cj = wb._rd(wb.wdir(wonbu) / "contents.json"); items = cj.get("contents") or []
    xd = wb.wdir(wonbu) / "xml"
    rows, meta = [], []
    for it in items:
        p = xd / f"{it['contents_id']}.xml"
        if not p.exists():
            continue
        tree = mods_xml.tree_from_xml(p.read_text(encoding="utf-8"))
        rows.append(mods_xml.sheet_row(tree)); meta.append({"contents_id": it["contents_id"], "vol": it.get("vol"), "title": it.get("title")})
    cols = mods_xml.sheet_columns(rows)
    return cols, rows, meta


# ---------------------------------------------------------------- C-3 형식 규칙(코드. 판단하지 않는다)
def _isbn_ok(s: str) -> bool:
    s = re.sub(r"[\s-]", "", s)
    if not re.fullmatch(r"\d{13}", s):
        return False
    total = sum(int(d) * (1 if k % 2 == 0 else 3) for k, d in enumerate(s[:12]))
    return (10 - total % 10) % 10 == int(s[12])


def rule_findings(cols: list[str], rows: list[dict]) -> list[dict]:
    """[{row, col, msg, level: error|check, by: rule}] — 가이드 5.8 과 정리 매크로의 형식 검사."""
    F: list[dict] = []
    bp = mods_xml.base_path
    like = lambda path: [c for c in cols if bp(c) == path]
    for i, r in enumerate(rows):
        v = lambda c: r.get(c, "")
        for c in cols:
            p, val = bp(c), v(c)
            if not val:
                continue
            if "  " in val:
                F.append({"row": i, "col": c, "msg": "공백 두 개 이상", "level": "error"})
            if p in ("/mods/titleInfo/title",) and re.search(r"[A-Za-z]", val) and val[0].isalpha() and val[0].islower():
                F.append({"row": i, "col": c, "msg": "외국어 표제 첫 글자가 소문자(가이드 1.1)", "level": "check"})
            if p in ("/mods/name/namePart", "/mods/name/displayForm", "/mods/name/alternativeName/namePart") and re.search(r"[A-Za-z]", val):
                if re.search(r",(?! )", val[:-1]):
                    F.append({"row": i, "col": c, "msg": "쉼표 뒤 공백 없음", "level": "error"})
                if re.search(r"[A-Z]{3,}", val):
                    F.append({"row": i, "col": c, "msg": "영문 대문자 연속(필명이면 정상)", "level": "check"})
            if p == "/mods/name/alternativeName[@altType]" and val not in ("nickname", "formal name", "no specific type"):
                F.append({"row": i, "col": c, "msg": "다른이름 유형값이 셋 중 하나가 아님(가이드 2.4)", "level": "error"})
            if p == "/mods/originInfo/place/placeTerm[@authority]" and val != "kormarccountry":
                F.append({"row": i, "col": c, "msg": "kormarccountry 가 아님", "level": "error"})
            if p == "/mods/originInfo/dateIssued" and not re.fullmatch(r"[\d-]{8}", val):
                F.append({"row": i, "col": c, "msg": "발행일 8자리 아님(가이드 5.3)", "level": "error"})
            if p == "/mods/identifier":
                t = v(c.replace("/mods/identifier", "/mods/identifier[@type]"))
                if t == "isbn" and not _isbn_ok(val):
                    F.append({"row": i, "col": c, "msg": "ISBN 형식·검증 숫자 오류", "level": "error"})
                if " " in val or "-" in val:
                    F.append({"row": i, "col": c, "msg": "식별기호에 공백·붙임표", "level": "error"})
            if p == "/mods/location/url" and " " in val:
                F.append({"row": i, "col": c, "msg": "원문주소에 공백", "level": "error"})
            if p == "/mods/classification" and not re.fullmatch(r"\d{3}(\.\d+)?", val):
                F.append({"row": i, "col": c, "msg": "분류기호 꼴 아님", "level": "error"})
            if p == "/mods/physicalDescription/extent" and not re.fullmatch(r"이미지 파일 \d+개 \([\d.]+ (MB|KB|GB|bytes)\) : (천연색|흑백)", val):
                F.append({"row": i, "col": c, "msg": "수량 문구 꼴 확인(가이드 7.4)", "level": "check"})
            if p in CONST and val != CONST[p]:
                F.append({"row": i, "col": c, "msg": f"'{CONST[p]}' 이어야 함", "level": "error"})
            if p == "/mods/name[@type]" and val not in ("personal", "corporate", "conference", "개인명", "단체명"):
                F.append({"row": i, "col": c, "msg": "저자 유형값 아님", "level": "error"})
            if p == "/mods/name[@ID]" and not re.fullmatch(r"KAC\d+", val):
                F.append({"row": i, "col": c, "msg": "전거 번호 꼴 아님", "level": "error"})
            if p == "/mods/subject[@ID]" and not re.fullmatch(r"KSH\d{10}", val):
                F.append({"row": i, "col": c, "msg": "주제명 번호 꼴 아님", "level": "error"})
        # 행 단위 대조
        for nc in like("/mods/name/namePart"):
            tc = nc.replace("/mods/name/namePart", "/mods/name[@type]")
            if v(nc) and not v(tc):
                F.append({"row": i, "col": nc, "msg": "저자 유형값 없음(가이드 2)", "level": "error"})
            ic = nc.replace("/mods/name/namePart", "/mods/name[@ID]"); ac = nc.replace("/mods/name/namePart", "/mods/name[@authority]")
            if v(ic) and not v(ac):
                F.append({"row": i, "col": ic, "msg": "전거 번호는 있는데 전거 기관이 없음", "level": "error"})
            dc = nc.replace("/mods/name/namePart", "/mods/name/displayForm")
            if v(dc) and not v(ic):
                F.append({"row": i, "col": dc, "msg": "전거 미연결인데 디스플레이형식 있음(가이드 2.2)", "level": "error"})
        pt, ptype = like("/mods/originInfo/place/placeTerm"), like("/mods/originInfo/place/placeTerm[@type]")
        texts = [v(c) for c, t in zip(pt, ptype) if v(t) == "text"]; codes = [v(c) for c, t in zip(pt, ptype) if v(t) == "code"]
        for text, code in zip(texts, codes):
            exp = region_code_from_place(text) if text else None
            if text and code and exp and exp != code:
                F.append({"row": i, "col": pt[0], "msg": f"발행지 '{text}' 와 부호 '{code}' 불일치(예상 {exp} {CODE_REGION.get(exp, '')})", "level": "error"})
            if code and code not in CODE_REGION:
                F.append({"row": i, "col": pt[0], "msg": f"발행국 부호 '{code}' 가 국내 부호표에 없음", "level": "check"})
        if len(texts) != len(codes):
            F.append({"row": i, "col": pt[0] if pt else cols[0], "msg": "발행지 글자/부호 개수 불일치", "level": "error"})
        for nc in like("/mods/note"):
            tc = nc.replace("/mods/note", "/mods/note[@type]"); note, typ = v(nc), v(tc)
            if not note:
                continue
            if typ and typ not in NOTE_TYPES:
                F.append({"row": i, "col": tc, "msg": "주기 유형값이 쓰는 유형이 아님(가이드 9)", "level": "error"})
            if re.search(r"이용가|세 미만|청소년", note) and typ != "target audience":
                F.append({"row": i, "col": tc, "msg": "이용대상 주기인데 유형이 target audience 아님", "level": "error"})
            if "수집한 자료" in note and typ != "acquisition":
                F.append({"row": i, "col": tc, "msg": "입수처 주기인데 유형이 acquisition 아님", "level": "error"})
        topics = [v(c) for c in like("/mods/subject/topic")]; genres = [v(c) for c in like("/mods/subject/genre")]
        if not any("만화" in t for t in topics):
            F.append({"row": i, "col": cols[0], "msg": "일반주제명 만화 없음(가이드 10.1)", "level": "error"})
        if not any("웹툰" in g for g in genres):
            F.append({"row": i, "col": cols[0], "msg": "장르주제명 웹툰 없음(가이드 10.6)", "level": "error"})
        for sc in like("/mods/subject[@ID]"):
            if not v(sc):
                F.append({"row": i, "col": sc, "msg": "주제명 전거 번호 없음(연결 안 됨)", "level": "error"})
        for path, nm in (("/mods/titleInfo/title", "표제"), ("/mods/originInfo/dateIssued", "발행일"), ("/mods/originInfo/publisher", "발행처"), ("/mods/physicalDescription/extent", "수량")):
            if not any(v(c) for c in like(path)):
                F.append({"row": i, "col": cols[0], "msg": f"{nm} 없음", "level": "error"})
    return [{**f, "by": "rule"} for f in F]


# ---------------------------------------------------------------- C-3 판단 규칙(에이전트)
def agent_job(wonbu: str, nth: int, cols: list[str], rows: list[dict], meta: list[dict], work_dir: Path = Path("work")) -> dict:
    imp = wb.import_for(wonbu, work_dir)
    return {"wonbu": wonbu, "nth": nth, "columns": cols, "rows": [{"contents_id": m["contents_id"], "vol": m.get("vol"), "values": r} for m, r in zip(meta, rows)],
            "observations_dir": (imp or {}).get("observations_dir"), "import_available": bool(imp), "authority_cache_dir": str(Path("work/authority").resolve()),
            "labels": {c: mods_xml.label_for(re.sub(r"^\[\d+\] /mods/", "", c).replace("/", ".").replace("[@", ".@").rstrip("]")) for c in cols}}


def check_agent_findings(result: Path, job: Path | None = None) -> list[str]:
    """에이전트 결과(findings.json)의 형식 검사(판단하지 않는다)."""
    out = []
    if not result.exists():
        return ["결과 파일이 없습니다"]
    try:
        d = json.loads(result.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        return [f"JSON 으로 읽히지 않습니다: {e}"]
    j = json.loads(Path(job).read_text(encoding="utf-8")) if job and Path(job).exists() else {}
    cols = set(j.get("columns") or []); n = len(j.get("rows") or [])
    for k, f in enumerate(d.get("findings") or []):
        if not isinstance(f.get("row"), int) or not (0 <= f["row"] < n):
            out.append(f"findings[{k}]: row 가 0~{n - 1} 사이 정수가 아닙니다")
        if f.get("col") not in cols:
            out.append(f"findings[{k}]: col '{f.get('col')}' 이 시트 열이 아닙니다")
        if f.get("level") not in ("error", "check"):
            out.append(f"findings[{k}]: level 은 error 또는 check")
        if not str(f.get("msg") or "").strip():
            out.append(f"findings[{k}]: msg 가 없습니다")
        if f.get("level") == "error" and not str(f.get("basis") or "").strip():
            out.append(f"findings[{k}]: error 는 basis(지침 조항·직원 답변 번호·원문 파일)가 있어야 합니다")
    if not out:
        rv = d.get("review") or {}
        if not rv.get("done"):
            out.append("검수를 아직 받지 않았습니다")
    return out


def run_agent(wonbu: str, nth: int, log=print, handle: dict | None = None, work_dir: Path = Path("work")) -> tuple[list[dict], list[str]]:
    from . import agent
    cols, rows, meta = load_rows(wonbu)
    job = agent_job(wonbu, nth, cols, rows, meta, work_dir)
    result, fails, jd = agent.run_job("check-mods", f"check_{wonbu}_{nth}", job, "findings.json", lambda p: check_agent_findings(p, p.with_name("job.json")), log, handle, timeout=2400)
    fs = [{**f, "by": "agent"} for f in result.get("findings") or []]
    (cdir(wonbu, nth) / "findings_agent.json").write_text(json.dumps({"findings": fs, "notes": result.get("notes_for_staff") or [], "usage": agent.usage_line(handle or {})}, ensure_ascii=False, indent=1), encoding="utf-8")
    return fs, fails


# ---------------------------------------------------------------- C-2/C-3 시트 쓰기
def write_sheet(wonbu: str, nth: int, findings: list[dict], out_path: Path | None = None, confirmed: dict | None = None) -> Path:
    """가이드 5.7 세팅의 점검 시트. confirmed = {"<row>|<col>": "ok"|"error"} 직원 처리(파란 칸을 채우기 없음으로, 오류 답 기재)."""
    cols, rows, meta = load_rows(wonbu)
    cols = list(cols)
    g = [c for c in cols if mods_xml.base_path(c) == "/mods/subject/genre"]
    for c in g:
        cols.remove(c)
    k = next((i for i, c in enumerate(cols) if mods_xml.base_path(c) == "/mods/classification"), len(cols))
    cols[k:k] = g
    wb_ = openpyxl.Workbook(); ws = wb_.active; ws.title = "점검"
    ncol = len(cols) + 1
    ws.append([None] * ncol); ws.append([None] * ncol)
    ws.append(["컨텐츠 아이디"] + [mods_xml.LABELS.get(mods_xml.base_path(c)) or mods_xml.base_path(c).rsplit("/", 1)[-1] for c in cols])
    ws.append(["컨텐츠 아이디"] + cols)
    for m, r in zip(meta, rows):
        ws.append([m["contents_id"]] + [r.get(c, "") for c in cols])
    last = ws.max_row
    for ci in range(1, ncol + 1):
        L = ws.cell(row=5, column=ci).column_letter
        ws.cell(row=2, column=ci).value = f"=COUNTA({L}5:{L}{last})"
    col_idx = {c: i + 2 for i, c in enumerate(cols)}
    hdr: dict[int, set[str]] = {}
    confirmed = confirmed or {}
    for f in findings:
        ci = col_idx.get(f["col"], 1); rn = f["row"] + 5
        cell = ws.cell(row=rn, column=ci)
        state = confirmed.get(f"{f['row']}|{f['col']}")
        if state == "ok":
            continue        # 직원이 '틀린 게 없다'고 한 확인 칸: 채우기 없음(가이드 5.8 3)
        level = "error" if (f["level"] == "error" or state == "error") else "check"
        if level == "error":
            cell.fill = YELLOW; cell.font = RED
            hdr.setdefault(ci, set()).add(f["msg"])
        elif cell.fill != YELLOW:
            cell.fill = BLUE
        prev = cell.comment.text + "\n" if cell.comment else ""
        cell.comment = Comment((prev + f["msg"] + (f" [{f.get('basis')}]" if f.get("basis") else ""))[:2000], "kolis_tool")
    for ci, msgs in hdr.items():
        h = ws.cell(row=1, column=ci); h.fill = YELLOW; h.font = RED; h.value = " / ".join(sorted(msgs))
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:{ws.cell(row=4, column=ncol).column_letter}{last}"
    out_path = out_path or (cdir(wonbu, nth) / "점검.xlsx")
    wb_.save(out_path)
    return out_path


def file_name(wonbu: str, nth: int, findings: list[dict], confirmed: dict | None, year: str | None = None) -> str:
    """가이드 5.9: 2026-(원부) n차점검(오류없음 | 오류종류 | 전체적 오류 많음)."""
    year = year or str(datetime.date.today().year)
    confirmed = confirmed or {}
    errs = [f for f in findings if (f["level"] == "error" or confirmed.get(f"{f['row']}|{f['col']}") == "error") and confirmed.get(f"{f['row']}|{f['col']}") != "ok"]
    if not errs:
        tag = "오류없음"
    else:
        by: dict[str, int] = {}
        for f in errs:
            by[f["msg"]] = by.get(f["msg"], 0) + 1
        top, cnt = max(by.items(), key=lambda kv: kv[1])
        tag = top if cnt / len(errs) > 0.6 else "전체적 오류 많음"
    return f"{year}-{wonbu} {nth}차점검({tag}).xlsx"


# ---------------------------------------------------------------- 한 바퀴(프로그램용)
def run_check(wonbu: str, log=print, with_agent: bool = True, handle: dict | None = None, work_dir: Path = Path("work")) -> dict:
    """C-1 받기 → C-2·C-3 시트. 직원 확정 전 상태(파일명은 임시 '점검.xlsx')."""
    nth = next_nth(wonbu)
    wb.fetch_all(wonbu, log)
    cols, rows, meta = load_rows(wonbu)
    findings = rule_findings(cols, rows)
    fails: list[str] = []
    if with_agent:
        fs, fails = run_agent(wonbu, nth, log, handle, work_dir)
        findings += fs
    d = cdir(wonbu, nth)
    (d / "findings.json").write_text(json.dumps({"wonbu": wonbu, "nth": nth, "findings": findings, "confirmed": {}, "at": wb._now()}, ensure_ascii=False, indent=1), encoding="utf-8")
    p = write_sheet(wonbu, nth, findings)
    errs = sum(1 for f in findings if f["level"] == "error"); checks = len(findings) - errs
    wb._patch(wonbu, check_nth=nth, check_at=wb._now(), check_errors=errs, check_checks=checks, check_file=str(p), check_confirmed=False)
    log(f"{nth}차 점검: {len(rows)}건 · 오류 {errs} · 확인 {checks} → {p}")
    return {"wonbu": wonbu, "nth": nth, "count": len(rows), "errors": errs, "checks": checks, "file": str(p), "remaining": fails, "findings": findings}


def confirm(wonbu: str, nth: int, decisions: dict, log=print) -> dict:
    """직원 확정: decisions = {"<row>|<col>": "ok"|"error"}. 파일명을 가이드 꼴로 바꾼 납품 파일을 만든다."""
    d = cdir(wonbu, nth)
    fj = json.loads((d / "findings.json").read_text(encoding="utf-8"))
    fj["confirmed"] = {**(fj.get("confirmed") or {}), **decisions}; fj["confirmed_at"] = wb._now()
    (d / "findings.json").write_text(json.dumps(fj, ensure_ascii=False, indent=1), encoding="utf-8")
    name = file_name(wonbu, nth, fj["findings"], fj["confirmed"])
    p = write_sheet(wonbu, nth, fj["findings"], d / name, fj["confirmed"])
    errs = sum(1 for f in fj["findings"] if (f["level"] == "error" or fj["confirmed"].get(f"{f['row']}|{f['col']}") == "error") and fj["confirmed"].get(f"{f['row']}|{f['col']}") != "ok")
    wb._patch(wonbu, check_confirmed=True, check_file=str(p), check_errors=errs, delivered_name=name)
    log(f"{nth}차 점검 확정: 오류 {errs} → {p}")
    return {"file": str(p), "name": name, "errors": errs, "nth": nth}


def status(wonbu: str) -> dict:
    st = wb.state(wonbu)
    nth = st.get("check_nth")
    fj = json.loads((cdir(wonbu, nth) / "findings.json").read_text(encoding="utf-8")) if nth and (cdir(wonbu, nth) / "findings.json").exists() else {}
    return {"nth": nth, "at": st.get("check_at"), "errors": st.get("check_errors"), "checks": st.get("check_checks"), "file": st.get("check_file"), "confirmed": st.get("check_confirmed"),
            "delivered_name": st.get("delivered_name"), "findings": fj.get("findings") or [], "decisions": fj.get("confirmed") or {}}
