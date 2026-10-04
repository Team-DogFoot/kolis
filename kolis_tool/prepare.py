"""납품 폴더 → 반입용 엑셀. 버튼 하나로 끝까지(유저 확정 2026-09-29). 2026-10-04 개선안: 반입에서 모든 것을 끝낸다.

프로그램이 하는 일은 셋뿐이다.
  1. 납품 폴더의 zip 을 풀고, 폴더 구성과 엑셀의 모든 칸을 있는 그대로 글로 옮긴다(snapshot). 양식을 가정하지 않는다.
  2. 에이전트를 한 번 실행한다(스킬 `prepare-import`). 에이전트가 원문 전체 관찰(회차마다 하위 에이전트 observe-episode) → 조사 → 전거·UCI·발행지 조회
     → MODS 트리로 값 결정 → 반입용 엑셀 쓰기(열은 값에 맞춰) → 검사 → 검수까지 스스로 돈다. 끝난 뒤 같은 검사를 돌려, 걸리면 같은 대화를 이어서 알려 준다.
     제한 시간은 이미지 수에 비례한다(agent.timeout_for). 실행기가 에이전트의 Read/look 호출을 기록해 "관찰 파일에 봤다고 적은 장을 실제로 열었는지" 검사한다.
  3. 마무리: 반입용 엑셀을 최종 값으로 다시 쓰고, 원고 파일명을 8자리 일련번호로, 썸네일 파일명을 반입용 엑셀의 값으로 바꾼다(되돌리기 가능). look 조각을 지운다.
중간 엑셀은 만들지 않는다. 직원이 보는 파일은 출판사 엑셀과 반입용 엑셀 둘뿐이다. 결과 파일(import.json)의 꼴은 import_check 모듈 머리 참조.
"""
from __future__ import annotations
import json, os, re, time, zipfile
from pathlib import Path
import openpyxl
import shutil
from . import agent, checks, import_check
from .common import IMAGE_EXT_ACCEPTED, natural_key, manifest_path, find_manifest, migrate_legacy

SKILL = "prepare-import"
MAX_CELLS = 6000


def stem(folder: Path) -> str:
    return re.sub(r"[^\w가-힣]+", "", Path(folder).name)


def result_path(folder: Path, work_dir: Path) -> Path:
    """작업 기록(에이전트 결과 묶음). 프로그램이 상태를 되살릴 때 읽는다. 직원이 볼 파일이 아니다."""
    return Path(work_dir) / f"{stem(folder)}.작업.json"


def load(folder: Path, work_dir: Path) -> dict | None:
    p = result_path(folder, work_dir)
    if p.exists():
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001
            return None
    return None


def _zip_names(path: Path) -> list[tuple[str, bool]]:
    out = []
    with zipfile.ZipFile(path) as z:
        for i in z.infolist():
            n = i.filename
            if not (i.flag_bits & 0x800):
                try:
                    n = n.encode("cp437").decode("cp949")
                except (UnicodeEncodeError, UnicodeDecodeError):
                    pass
            out.append((n, i.is_dir()))
    return out


def unpack_zips(folder: Path, log=None) -> list[str]:
    """아직 풀리지 않은 zip 을 `<폴더>/<zip 이름>/` 에 푼다(한글 파일명 cp949). 이미 그 폴더가 있으면 건드리지 않는다."""
    from .unzip_kr import extract
    log = log or (lambda m: None)
    done = []
    for z in sorted(Path(folder).glob("*.zip")):
        dest = Path(folder) / z.stem
        if dest.exists():
            continue
        log(f"압축 풀기: {z.name} → {dest.name}\\")
        log(f"  {extract(z, dest)}개 파일")
        done.append(str(dest))
    return done


def snapshot(folder: Path) -> str:
    """폴더 구성과 엑셀 내용을 있는 그대로 글로. 어떤 열이 무엇인지는 여기서 판단하지 않는다."""
    folder = Path(folder)
    lines = [f"# 납품 폴더: {folder}", "", "## 폴더 구성"]

    def walk(d: Path, level: int):
        entries = sorted(d.iterdir(), key=lambda p: (p.is_file(), natural_key(p.name)))
        imgs = [p for p in entries if p.is_file() and p.suffix.lower() in IMAGE_EXT_ACCEPTED]
        pad = "  " * level
        if imgs:
            size = sum(p.stat().st_size for p in imgs)
            lines.append(f"{pad}[이미지 {len(imgs)}개, {size / 1048576:.2f} MB, 확장자 {','.join(sorted({p.suffix.lower() for p in imgs}))}] "
                         f"처음 {imgs[0].name} … 끝 {imgs[-1].name}")
        for p in entries:
            if p.is_file() and p not in imgs:
                lines.append(f"{pad}{p.name}  ({p.stat().st_size:,} B)")
        for p in entries:
            if p.is_dir() and not p.name.startswith("_kolis"):
                lines.append(f"{pad}{p.name}\\")
                if level < 4:
                    walk(p, level + 1)
    walk(folder, 0)
    for z in sorted(folder.rglob("*.zip")):
        dirs: dict[str, int] = {}
        for n, is_dir in _zip_names(z):
            if not is_dir:
                d = n.rsplit("/", 1)[0] if "/" in n else "(최상위)"
                dirs[d] = dirs.get(d, 0) + 1
        lines += ["", f"## zip 내용: {z.relative_to(folder)}"] + [f"  {d}: 파일 {c}개" for d, c in sorted(dirs.items(), key=lambda kv: natural_key(kv[0]))][:200]
    for x in sorted(folder.rglob("*.xls*")):
        if x.name.startswith("~$") or x.suffix.lower() not in (".xlsx", ".xlsm"):
            continue
        for ws in openpyxl.load_workbook(x, data_only=True).worksheets:
            lines += ["", f"## 엑셀: {x.relative_to(folder)} / 시트 '{ws.title}' ({ws.max_row}행 × {ws.max_column}열)", "칸 주소 = 값 (빈 칸은 생략)"]
            n = 0
            for row in ws.iter_rows():
                cells = [f"{c.coordinate}={str(c.value).strip()!r}" for c in row if c.value not in (None, "")]
                if cells:
                    lines.append("  " + " | ".join(cells)); n += len(cells)
                if n > MAX_CELLS:
                    lines.append(f"  … (칸 {MAX_CELLS}개를 넘어 생략)"); break
    return "\n".join(lines)


def _check_all(import_json: Path) -> list[str]:
    """프로그램이 끝에 돌리는 검사. 에이전트가 작업 중에 돌리는 두 명령과 같은 검사다(+ 실행기의 Read/look 기록 대조)."""
    jd = import_json.parent
    job = json.loads((jd / "job.json").read_text(encoding="utf-8"))
    out = [f"[조사] {f}" for f in checks.check_research(jd / "research.json", jd / "job.json")]
    if not import_json.exists():
        return out + ["[반입용] import.json 이 없습니다"]
    try:
        data = json.loads(import_json.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        return out + [f"[반입용] import.json 이 JSON 으로 읽히지 않습니다: {e}"]
    data["_job_dir"] = str(jd)
    reads = None
    rf = jd / "_reads.json"
    if rf.exists():
        try:
            reads = set(json.loads(rf.read_text(encoding="utf-8")))
        except json.JSONDecodeError:
            reads = None
    return out + [f"[반입용] {f}" for f in import_check.check(data, Path(job["folder"]), reads)]


def count_images(folder: Path) -> int:
    return sum(1 for p in Path(folder).rglob("*") if p.is_file() and p.suffix.lower() in IMAGE_EXT_ACCEPTED and "_kolis" not in p.parts)


def output_name(title: str, batch_note: str) -> str:
    """반입용 엑셀 파일명(작업자 관행): 기초메타데이터(26웹툰대행5차-203)_코리스 반입용_제목.xlsx. 차수·번호는 작업 번호(비고)에서 읽는다."""
    m = re.search(r"(\d{4})?-?납본-웹툰대행\((\d*)차\)\((\d*)\)", batch_note or "")
    yy = (m.group(1) or "2026")[2:] if m else "26"
    nth, no = (m.group(2), m.group(3)) if m else ("", "")
    return f"기초메타데이터({yy}웹툰대행{nth}차-{no})_코리스 반입용_{re.sub(r'[\\/:*?\"<>|\s]+', '', title or '')}.xlsx"


def run(folder: Path, work_dir: Path, output_xlsx: Path | None = None, instructions: str = "", batch_note: str = "", log=None,
        handle: dict | None = None, timeout: int = 3600) -> dict:
    """output_xlsx 를 비우면 제목을 읽은 뒤 관행대로 이름을 붙여 work 폴더에 저장한다."""
    log = log or (lambda m: None)
    handle = handle if handle is not None else {}
    folder, work_dir = Path(folder), Path(work_dir)
    auto = output_xlsx is None
    output_xlsx = work_dir / f"반입용_{stem(folder)}.xlsx" if auto else Path(output_xlsx)
    if not folder.is_dir():
        raise SystemExit(f"폴더 없음: {folder}")
    t0 = time.time()
    unpack_zips(folder, log)
    home = agent.deploy(log)
    name = stem(folder)
    jd = home / "jobs" / name
    jd.mkdir(parents=True, exist_ok=True)
    for old in ("research.json", "import.json"):
        if (jd / old).exists():
            (jd / old).unlink()
    (jd / "snapshot.txt").write_text(snapshot(folder), encoding="utf-8")
    common = Path(work_dir) / "research_prompt.txt"      # 직원이 프로그램에서 적은 '모든 작품에 적용할 지시'
    settings = project_settings(work_dir)
    n_images = count_images(folder)
    job = {"task": "납품 폴더 → 반입용 엑셀", "folder": str(folder), "snapshot": f"jobs/{name}/snapshot.txt", "output_xlsx": str(output_xlsx),
           "observations_dir": str(jd / "observations"), "image_count": n_images,
           "instructions": (instructions or "").strip(), "instructions_for_all_works": common.read_text(encoding="utf-8").strip() if common.exists() else "",
           "batch_note": (batch_note or "").strip(), "project": settings}
    (jd / "observations").mkdir(exist_ok=True)
    t_out = max(int(timeout), agent.timeout_for(n_images))
    log(f"에이전트에게 맡김: {folder.name} (이미지 {n_images}장, 제한 {t_out}초)" + (f" / 지시: {job['instructions'][:80]}" if job["instructions"] else ""))
    data, fails, _ = agent.run_job(SKILL, name, job, "import.json", _check_all, log, handle, add_dirs=[folder],
                                   model=os.environ.get("KOLIS_AGENT_MODEL") or None, timeout=t_out)
    blocking = [f for f in fails if "검수" not in f]
    if any(f.startswith("[반입용]") for f in blocking):
        raise SystemExit("반입용 값이 검사를 통과하지 못해 엑셀을 만들지 않았습니다: " + "; ".join(blocking[:5]))
    if auto:
        draft, output_xlsx = output_xlsx, work_dir / output_name(str(data.get("title") or stem(folder)), batch_note)
        if draft.exists() and draft != output_xlsx:
            draft.unlink()
    data["_job_dir"] = str(jd)
    written = import_check.write_xlsx(data, folder, output_xlsx)
    research = json.loads((jd / "research.json").read_text(encoding="utf-8")) if (jd / "research.json").exists() else {}
    done = finalize(folder, data, log)
    shutil.rmtree(jd / "_look", ignore_errors=True)
    result = {"folder": str(folder), "title": data.get("title"), "unit": data.get("unit"), "output_xlsx": str(output_xlsx), "rows": written["rows"],
              "columns": written["columns"], "unknown_columns": written.get("unknown") or [],
              "confirm_cells": written["confirm"], "manuscripts": str(folder / data["manuscripts_root"]),
              "thumbs": str(folder / data["thumbs_dir"]) if data.get("thumbs_dir") else "", "import": data, "research": research,
              "observations_dir": str(jd / "observations"), "adult": bool(data.get("adult")), "adult_reason": data.get("adult_reason") or "",
              "remaining": fails, "finalize": done,
              "run": {"seconds": int(time.time() - t0), "turns": handle.get("turns"), "session_id": handle.get("session_id"), "job_dir": str(jd), "images": n_images,
                      "usage": handle.get("usage") or {}, "usage_text": agent.usage_line(handle)}}
    result_path(folder, work_dir).parent.mkdir(parents=True, exist_ok=True)
    result_path(folder, work_dir).write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    log(f"반입용 엑셀: {output_xlsx.name} — {written['rows']}행, {written['columns']}열, 사람이 확인할 칸 {written['confirm']}개" + (f", 검사에 남은 항목 {len(fails)}건" if fails else ""))
    if handle.get("usage"):
        log(f"토큰 사용량(작품 전체, 하위 에이전트 포함): {agent.usage_line(handle)}")
    return result


def project_settings(work_dir: Path) -> dict:
    """사업 고정값(에이전트에게 주는 기본값. 자료마다 다른 값은 에이전트가 정한다). 파일이 없으면 기본값."""
    p = Path(work_dir) / "build_settings.json"
    base = {"acquisition_note": "한국웹툰산업협회를 통해 수집한 자료임", "classification": "810", "classification_authority": "KDC", "classification_edition": "6",
            "physical_location": "국립중앙도서관", "region": "한국", "type_of_resource": "텍스트", "genre": "만화", "form": "전자자료(Image)",
            "reformatting_quality": "access", "digital_origin": "BornDigital", "language_default": "kor", "currency_code": "\\",
            "subjects": [{"kind": "topic", "term": "만화[漫畵]", "id": "KSH1998022212", "authority": "국립중앙도서관주제명표목표"},
                         {"kind": "genre", "term": "웹툰[webtoon]", "id": "KSH2016000049", "authority": "국립중앙도서관주제명표목표"}],
            "author_authority": "국립중앙도서관전거데이터"}
    if p.exists():
        try:
            base.update(json.loads(p.read_text(encoding="utf-8")))
        except json.JSONDecodeError:
            pass
    return base


def rewrite_xlsx(folder: Path, work_dir: Path) -> dict:
    """직원 결정(전거 번호 등)으로 import.json 이 바뀐 뒤 반입용 엑셀을 다시 쓴다(확인 완료 전에만)."""
    w = load(folder, work_dir)
    if not w:
        raise SystemExit("작업 기록이 없습니다")
    data = w["import"]
    data["_job_dir"] = w.get("run", {}).get("job_dir") or data.get("_job_dir")
    written = import_check.write_xlsx(data, Path(folder), Path(w["output_xlsx"]))
    w["confirm_cells"] = written["confirm"]; w["columns"] = written["columns"]; w["confirmed"] = False
    result_path(Path(folder), work_dir).write_text(json.dumps(w, ensure_ascii=False, indent=1), encoding="utf-8")
    return written


def finalize(folder: Path, data: dict, log=None) -> dict:
    """마무리(기계적인 일): 원고 파일명을 8자리 일련번호로, 썸네일을 반입용 엑셀에 적힌 이름으로. 되돌리기 기록은 폴더 밖 `_kolis_manifests/`."""
    from . import rename_files, arrange
    log = log or (lambda m: None)
    arranged = arrange.apply(folder, data, log)      # 에이전트가 정한 정리 계획대로 폴더·파일을 옮긴다(되돌리기 가능)
    ms = folder / data["manuscripts_root"]
    migrate_legacy(ms)
    nfiles = nfolders = 0
    for r in data["rows"]:
        d = ms / str(r["folder"])
        if d.is_dir() and not rename_files.is_done(d):
            nfiles += len(rename_files.apply(d)); nfolders += 1
    if nfolders:
        log(f"원고 파일명 정리: 폴더 {nfolders}개, {nfiles}장 → 8자리 일련번호(되돌리기 가능)")
    nthumb = 0
    if data.get("thumbs_dir"):
        td = folder / data["thumbs_dir"]
        if not find_manifest(td, "thumbs", "thumbs_manifest.json"):
            pairs = [(str(r.get("thumb_source") or ""), str(r.get("thumb_file") or "")) for r in data["rows"]]
            pairs = [(a, b) for a, b in pairs if a and b and a != b and (td / a).exists()]
            for a, b in pairs:
                (td / a).rename(td / f"__tmp__{b}")
            for a, b in pairs:
                (td / f"__tmp__{b}").rename(td / b)
            if pairs:
                manifest_path(td, "thumbs").write_text(json.dumps(pairs, ensure_ascii=False, indent=1), encoding="utf-8")
                nthumb = len(pairs)
                log(f"썸네일 파일명 정리: {nthumb}장(되돌리기 가능)")
    return {"manuscript_folders": nfolders, "manuscript_files": nfiles, "thumbs": nthumb, "arranged": arranged}


def undo(folder: Path, data: dict) -> dict:
    """마무리에서 바꾼 파일명을 되돌린다."""
    from . import rename_files
    ms = Path(folder) / data["manuscripts_root"]
    n = sum(rename_files.undo(ms / str(r["folder"])) for r in data["rows"] if rename_files.is_done(ms / str(r["folder"])))
    t = 0
    if data.get("thumbs_dir"):
        td = Path(folder) / data["thumbs_dir"]
        mf = find_manifest(td, "thumbs", "thumbs_manifest.json")
        if mf:
            for a, b in reversed(json.loads(mf.read_text(encoding="utf-8"))):
                if (td / b).exists():
                    (td / b).rename(td / a); t += 1
            mf.unlink()
    from . import arrange
    return {"manuscript_files": n, "thumbs": t, "arranged": arrange.undo(Path(folder))}      # 파일명을 먼저 되돌린 뒤 폴더 정리를 되돌린다
