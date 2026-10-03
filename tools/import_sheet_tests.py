"""개발용 시험(KOLIS 없이 돈다): 반입용 엑셀 직렬화기(kolis_tool/mods_sheet.py)의 2026-10-04 개선안 ① 시험.
  1. 예시 88열(완료사례 로맨스낫로맨틱)을 트리로 역변환해 다시 쓰면 머리글이 같고, 값이 있는 셀이 전부 같다.
  2. 2026-10-03 반입 헤더 시험(접수 1027~1029, tools/import_probe.py 로 끼운 열)의 묶음 배치를 트리로 만들어 쓰면 그 배치가 나온다
     (전거 열은 저자 유형 뒤, 주제명 2묶음, 둘째 식별기호 UCI, 둘째 출처정보는 발행지 부모 열 + placeTerm 쌍, 디스플레이형식·발행처 2칸·판사항·수상 주기).
  3. 반복 확장: 저자 5·다른이름 2·출처정보 3·식별기호 2·주기 5·주제명 2·원문주소 4 + 모르는 요소(relatedItem) → 열이 생기고, 역변환하면 값이 같고, 모르는 요소는 '시험 전'으로 표시된다.
  4. 확인할 칸(path) → 노란 칸 위치.
실행: python3 -X utf8 tools/import_sheet_tests.py   (출력 폴더: work/eval/ 또는 --out)"""
import sys, json, argparse
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from kolis_tool import mods_sheet as ms

EX = Path("docs/source/3_샘플양식/완료사례_코리스반입용_로맨스낫로맨틱(88열).xlsx")


def headers_of(xlsx: Path) -> list[str]:
    ws = openpyxl.load_workbook(xlsx, read_only=True)["Contents"]
    return [(c or "").strip() if isinstance(c, str) else "" for c in next(ws.iter_rows(values_only=True))]


def test_roundtrip_88(out: Path) -> list[str]:
    rows = ms.read_sheet(EX)
    r = ms.write({"rows": rows}, out / "roundtrip_88.xlsx")
    fails = []
    if r["headers"] != headers_of(EX):
        fails.append("예시 88열과 머리글이 다름: " + json.dumps([h for h in r["headers"] if h not in headers_of(EX)] + [h for h in headers_of(EX) if h not in r["headers"]], ensure_ascii=False))
    a, b = ms.sheet_value_cells(EX), ms.sheet_value_cells(out / "roundtrip_88.xlsx")
    if a != b:
        fails.append(f"값 셀 다름: 예시에만 {sorted(a - b)[:5]} / 산출에만 {sorted(b - a)[:5]}")
    return fails


def test_probe_layout(out: Path) -> list[str]:
    """10-03 시험 엑셀의 묶음 배치를 트리로 재현(값은 시험 때 넣은 것과 같은 꼴)."""
    rows = ms.read_sheet(EX)[:2]
    for r in rows:
        m = r["mods"]
        n0 = ms._node(m["name"][0]); n0["@ID"] = "KAC201418251"; n0["@authority"] = "국립중앙도서관전거데이터"
        n0["displayForm"] = "황재(원문 표기 시험)"; n0["alternativeName"] = [{"@type": "formal name", "namePart": "황재(시험 다른이름)"}]
        m["name"][0] = n0
        o = ms._node(m["originInfo"][0]); o["publisher"] = list(ms._lst(o.get("publisher"))) + ["시험유통사"]; o["edition"] = "개정판"; m["originInfo"][0] = o
        m["originInfo"] = [o, {"@type": "production", "place": [{"placeTerm": {"_": "[부산]", "@type": "text"}}, {"placeTerm": {"_": "bnk", "@authority": "kormarccountry", "@type": "code"}}],
                               "issuance": "단행자료", "publisher": ["시험제작사 [제작]"], "dateIssued": "20120924"}]
        m["note"] = list(ms._lst(m.get("note"))) + [{"_": "문화체육관광부장관상, 2025(시험)", "@type": "awards"}]
        m["subject"] = [{"@ID": "KSH1998022212", "@authority": "국립중앙도서관주제명표목표", "topic": "만화[漫畵]"}, {"@ID": "KSH2016000049", "@authority": "국립중앙도서관주제명표목표", "genre": "웹툰[webtoon]"}]
        m["identifier"] = list(ms._lst(m.get("identifier"))) + [{"_": "G903+NL-TEST-0001", "@type": "uci"}]
    r = ms.write({"rows": rows}, out / "probe_layout.xlsx")
    h = r["headers"]
    fails = []
    def after(a, b, msg):
        ia = [i for i, x in enumerate(h) if x == a]; ib = [i for i, x in enumerate(h) if x == b]
        if not ia or not ib or ib[0] != ia[0] + 1:
            fails.append(msg + f" (실제: {h[max(0, (ia or [0])[0] - 1):(ia or [0])[0] + 4]})")
    after("/mods/name[@type]", "/mods/name[@ID]", "전거 번호 열은 저자 유형 열 바로 뒤여야 함(10-03 시험 배치)")
    after("/mods/name[@ID]", "/mods/name[@authority]", "전거 기관 열은 전거 번호 뒤")
    after("/mods/name/role/roleTerm", "/mods/name/displayForm", "디스플레이형식은 역할어 뒤(3차 시험)")
    seq = ["/mods/subject/", "/mods/subject/topic", "/mods/subject[@ID]", "/mods/subject[@authority]", "/mods/subject/", "/mods/subject/genre", "/mods/subject[@ID]", "/mods/subject[@authority]"]
    i = h.index("/mods/subject/") if "/mods/subject/" in h else -1
    if i < 0 or h[i:i + len(seq)] != seq:
        fails.append(f"주제명 2묶음 배치가 다름: {h[i:i + len(seq)] if i >= 0 else '없음'}")
    ids = [i for i, x in enumerate(h) if x == "/mods/identifier"]
    if len(ids) != 2 or h[ids[1] + 1] != "/mods/identifier[@type]":
        fails.append("식별기호 2쌍(ISBN, UCI)이 아님")
    oi = [i for i, x in enumerate(h) if x == "/mods/originInfo"]
    if len(oi) != 2:
        fails.append("출처정보 2묶음이 아님")
    else:
        second = h[oi[1]:oi[1] + 12]
        want = ["/mods/originInfo", "/mods/originInfo[@type]", "/mods/originInfo/place", "/mods/originInfo/place/placeTerm", "/mods/originInfo/place/placeTerm[@type]",
                "/mods/originInfo/place", "/mods/originInfo/place/placeTerm", "/mods/originInfo/place/placeTerm[@authority]", "/mods/originInfo/place/placeTerm[@type]",
                "/mods/originInfo/issuance", "/mods/originInfo/publisher", "/mods/originInfo/dateIssued"]
        if second != want:
            fails.append(f"둘째 출처정보 배치(발행지 부모 열 + placeTerm 쌍, 1차 시험 꼴)가 다름: {second}")
        first = h[oi[0]:oi[1]]
        if "/mods/originInfo/place" in first:
            fails.append("첫 출처정보에는 발행지 부모 열을 두지 않는다(템플릿·예시 모양)")
        if first.count("/mods/originInfo/publisher") != 3:
            fails.append(f"첫 출처정보 발행처 3칸이어야 함(예시 2 + 시험 1): {first.count('/mods/originInfo/publisher')}")
    if "/mods/originInfo/edition" not in h or h.count("/mods/note[@type]") < 4:
        fails.append("판사항·넷째 주기 열 없음")
    back = ms.read_sheet(out / "probe_layout.xlsx")
    if ms.value_cells(rows) != ms.value_cells(back):
        fails.append("시험 배치 역변환 값 다름")
    return fails


def test_expand(out: Path) -> list[str]:
    row = {"mods": {
        "titleInfo": [{"title": "시험 작품", "partNumber": "1화", "partName": "시작"}, {"@type": "parallel", "title": "Test work"}],
        "name": [{"@usage": "primary", "@type": "개인명", "@ID": "KAC000000001", "@authority": "국립중앙도서관전거데이터", "namePart": "가", "role": {"roleTerm": "글"}, "displayForm": "가(원문)", "alternativeName": [{"@type": "no specific type", "namePart": "가(원문)"}]},
                 {"@type": "개인명", "namePart": "나", "role": {"roleTerm": "그림"}, "alternativeName": [{"@type": "nickname", "namePart": "나나"}, {"@type": "formal name", "namePart": "나본명"}]},
                 {"@type": "개인명", "namePart": "다", "role": {"roleTerm": "각색"}}, {"@type": "개인명", "namePart": "라", "role": {"roleTerm": "원작"}}, {"@type": "단체명", "namePart": "마스튜디오", "role": {"roleTerm": "제작"}}],
        "originInfo": [{"@eventType": "publication", "place": [{"placeTerm": {"_": "[서울]", "@type": "text"}}, {"placeTerm": {"_": "ulk", "@authority": "kormarccountry", "@type": "code"}}], "issuance": "단행자료", "publisher": ["A", "B", "C"], "dateIssued": "20240101", "edition": "개정판"},
                       {"@type": "production", "place": [{"placeTerm": {"_": "[부산]", "@type": "text"}}, {"placeTerm": {"_": "bnk", "@authority": "kormarccountry", "@type": "code"}}], "issuance": "단행자료", "publisher": ["D [제작]"]},
                       {"@type": "production", "place": [{"placeTerm": {"_": "[대구]", "@type": "text"}}, {"placeTerm": {"_": "tgk", "@authority": "kormarccountry", "@type": "code"}}], "issuance": "단행자료", "publisher": ["E [제작]"]}],
        "language": {"languageTerm": {"_": "kor", "@authority": "iso639-2b", "@type": "code"}},
        "physicalDescription": {"form": "전자자료(Image)", "reformattingQuality": "access", "digitalOrigin": "BornDigital"},
        "targetAudience": "일반이용자",
        "note": [{"_": "A은 B의 임프린트임"}, {"_": "15세 이용가", "@type": "target audience"}, {"_": "한국웹툰산업협회를 통해 수집한 자료임", "@type": "acquisition"}, {"_": "문화체육관광부장관상, 2025", "@type": "awards"}, {"_": "한국콘텐츠진흥원 2024", "@type": "funding"}],
        "subject": [{"@ID": "KSH1998022212", "@authority": "국립중앙도서관주제명표목표", "topic": "만화[漫畵]"}, {"@ID": "KSH2016000049", "@authority": "국립중앙도서관주제명표목표", "genre": "웹툰[webtoon]"}],
        "classification": {"_": "810", "@authority": "KDC", "@edition": "6"},
        "identifier": [{"_": "9791109764511", "@type": "isbn"}, {"_": "G903+NL-TEST", "@type": "uci"}],
        "location": {"url": ["https://a", "https://b", "https://c", "https://d"], "physicalLocation": "국립중앙도서관"},
        "accessCondition": {"licenseType": "2"}, "extension": {"regionOfPublishing": "한국"}, "typeOfResource": "텍스트", "genre": "만화",
        "relatedItem": [{"@type": "otherFormat", "recordInfo": {"recordIdentifier": "KMO000000000"}}]},
        "extra": {"currency_code": "\\", "contents_price": 500, "compensation": 500, "reward_yn": "Y", "thum_files": "x.jpg"},
        "confirm": [{"path": "name[1].alternativeName[1].namePart", "reason": "본명 근거 확인"}, {"path": "identifier[1]", "reason": "UCI 확인"}, {"path": "extra.contents_price", "reason": "가격"}, {"path": "titleInfo[0].title", "reason": "제목"}]}
    row2 = {"mods": {"titleInfo": [{"title": "시험 작품", "partNumber": "2화"}], "name": [{"@type": "개인명", "namePart": "가", "role": {"roleTerm": "글"}}]}, "extra": {}}
    r = ms.write({"rows": [row, row2]}, out / "expand.xlsx", confirm_text=lambda c: c["reason"])
    fails = []
    h = r["headers"]
    if h.count("/mods/name") != 5 or h.count("/mods/name/alternativeName") != 3 or h.count("/mods/originInfo") != 3 or h.count("/mods/identifier") != 2 or h.count("/mods/note") != 5 or h.count("/mods/subject/") != 2 or h.count("/mods/location/url") != 4:
        fails.append(f"반복 수가 다름: name {h.count('/mods/name')} alt {h.count('/mods/name/alternativeName')} origin {h.count('/mods/originInfo')} ident {h.count('/mods/identifier')} note {h.count('/mods/note')} subject {h.count('/mods/subject/')} url {h.count('/mods/location/url')}")
    if not any(x.startswith("/mods/relatedItem") for x in r["unknown"]):
        fails.append("모르는 요소(relatedItem)가 '시험 전'으로 표시되지 않음")
    back = ms.read_sheet(out / "expand.xlsx")
    if ms.value_cells([row, row2]) != ms.value_cells(back):
        fails.append("확장 역변환 값 다름")
    ws = openpyxl.load_workbook(out / "expand.xlsx")["Contents"]
    yellow = [ws.cell(1, c.column).value for c in ws[2] if c.fill and c.fill.fgColor.rgb in ("FFFFFF00", "00FFFF00")]
    if sorted(yellow) != sorted(["/mods/titleInfo/title", "/mods/name/alternativeName/namePart", "/mods/identifier", "contents_price"]):
        fails.append(f"노란 칸 위치가 다름: {yellow}")
    cols = ms.layout([row, row2])
    ci = ms.column_index(cols, "name[1].alternativeName[1].namePart")
    if not ci or ws.cell(2, ci).value != "나본명":
        fails.append("confirm 경로가 엉뚱한 열을 가리킴")
    return fails


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default="work/eval")
    ns = ap.parse_args()
    out = Path(ns.out); out.mkdir(parents=True, exist_ok=True)
    total = 0
    for name, fn in (("1 예시 88열 역변환", test_roundtrip_88), ("2 10-03 시험 배치 재현", test_probe_layout), ("3 반복 확장·모르는 요소·노란 칸", test_expand)):
        fails = fn(out)
        total += len(fails)
        print(f"{name}: {'통과' if not fails else '걸림 ' + str(len(fails))}")
        for f in fails:
            print("  -", f)
    print("전부 통과" if not total else f"걸린 항목 {total}건")
    return 0 if not total else 1


if __name__ == "__main__":
    sys.exit(main())
