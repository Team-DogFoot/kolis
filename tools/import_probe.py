"""헤더 시험 반입(2026-10-03 유저 지시 "다 테스트 해봐, 취소하면 되잖아"): 원시인 삼촌 2행의 반입용 엑셀에 시험 열을 끼워 일괄반입(요청)하고,
접수 목록이 돌려주는 콘텐츠 MODS 로 어떤 열이 들어갔는지 본다. 접수번호는 취소 요청 목록에 '취소 요청'으로 적는다. 원문 등록·가원부번호는 하지 않는다.
시험 열: 저자 전거 name[@ID]·[@authority] / 다른이름 alternativeName + [@type] / 주제명 subject[@ID]·[@authority] + 둘째 주제명 묶음 subject/genre
        / 식별기호 둘째 identifier(uci) / 출처정보 2묶음 × 발행지 2쌍(발행지 부모 열 /mods/originInfo/place 를 넣어서)
실행: .venv/Scripts/python.exe -X utf8 tools/import_probe.py
"""
import sys, json, re, datetime, shutil
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from kolis_tool import kolis_http, kolis_request as req, kolis_flow

SRC = Path("work/기초메타데이터(26웹툰대행5차-901)_코리스 반입용_원시인삼촌.xlsx")
OUT = Path("work/probe"); OUT.mkdir(parents=True, exist_ok=True)
XLSX = OUT / ("기초메타데이터(26웹툰대행5차-901)_코리스 반입용_원시인삼촌_헤더시험" + ("3" if "--variant3" in sys.argv else "2" if "--variant2" in sys.argv else "") + ".xlsx")
NOTE = "2026-납본-웹툰대행(5차)(901)"
VARIANT2 = "--variant2" in sys.argv or "--variant3" in sys.argv
VARIANT3 = "--variant3" in sys.argv     # 3차: 디스플레이형식, 다른이름 [@altType] 표기, 첫 출처정보 발행처 2칸, 판사항, 넷째 주기(수상)
AUTH = "국립중앙도서관전거데이터"; SUBJ_AUTH = "국립중앙도서관주제명표목표"


def load(xlsx):
    wb = openpyxl.load_workbook(xlsx); ws = wb["Contents"]
    rows = list(ws.iter_rows(values_only=True))
    h = [(c or "").strip() if isinstance(c, str) else "" for c in rows[0]]
    data = [list(r) for r in rows[1:] if any(c not in (None, "") for c in r)]
    return h, data


def build(h, data, kac):
    """새 머리글 목록과 행 값을 만든다. (머리글, [행1값, 행2값]) 목록."""
    cols = [(x, [r[i] for r in data]) for i, x in enumerate(h)]
    out = []
    n_name_type = n_sub_topic = n_ident_type = n_origin = n_place = 0
    skip_second_topic = False
    for i, (x, vals) in enumerate(cols):
        if x == "/mods/subject/topic":
            n_sub_topic += 1
            if n_sub_topic == 2:
                continue        # 템플릿의 둘째 topic 열은 뺀다
        if x == "/mods/originInfo":
            n_origin += 1
        if x == "/mods/originInfo/place/placeTerm" and n_origin == 1:
            n_place += 1
            if n_place == 1 or (n_place == 2 and not VARIANT2):
                out.append(("/mods/originInfo/place", [None, None]))      # 발행지 부모 열(Sample: 반복에 따른 부모정의, 데이터 없음)
        if x == "/mods/originInfo[@type]" and n_origin == 2:
            # 둘째 출처정보: 유형 뒤에 발행지 2쌍을 끼운다
            out.append((x, ["production", "production"]))
            out.append(("/mods/originInfo/place", [None, None])); out.append(("/mods/originInfo/place/placeTerm", ["[부산]", "[부산]"])); out.append(("/mods/originInfo/place/placeTerm[@type]", ["text", "text"]))
            (None if VARIANT2 else out.append(("/mods/originInfo/place", [None, None]))); out.append(("/mods/originInfo/place/placeTerm", ["bnk", "bnk"])); out.append(("/mods/originInfo/place/placeTerm[@authority]", ["kormarccountry", "kormarccountry"])); out.append(("/mods/originInfo/place/placeTerm[@type]", ["code", "code"]))
            continue
        if x == "/mods/originInfo/issuance" and n_origin == 2:
            out.append((x, ["단행자료", "단행자료"])); continue
        if x == "/mods/originInfo/publisher" and n_origin == 2:
            out.append((x, ["시험제작사[제작]", "시험제작사[제작]"])); continue
        if x == "/mods/originInfo/dateIssued" and n_origin == 2:
            out.append((x, ["20120924", "20120924"])); continue
        out.append((x, vals))
        if x == "/mods/name[@type]":
            n_name_type += 1
            if n_name_type == 1:
                out.append(("/mods/name[@ID]", [kac, kac])); out.append(("/mods/name[@authority]", [AUTH, AUTH]))
        if x == "/mods/name/alternativeName[@type]" and VARIANT3:
            out[-1] = ("/mods/name/alternativeName[@altType]", ["formal name", "formal name"])      # 표기 시험: [@altType]
        if x == "/mods/name/role/roleTerm" and VARIANT3 and n_name_type == 1:
            out.append(("/mods/name/displayForm", ["황재(원문 표기 시험)", "황재(원문 표기 시험)"]))
        if x == "/mods/originInfo/publisher" and n_origin == 1 and VARIANT3:
            out.append(("/mods/originInfo/publisher", ["시험유통사", "시험유통사"]))
        if x == "/mods/originInfo/edition" and n_origin == 1 and VARIANT3:
            out[-1] = (x, ["개정판", "개정판"])
        if x == "/mods/note[@type]" and VARIANT3 and vals[0] == "acquisition":
            out.append(("/mods/note", ["문화체육관광부장관상, 2025(시험)", "문화체육관광부장관상, 2025(시험)"])); out.append(("/mods/note[@type]", ["awards", "awards"]))
        if x == "/mods/subject/topic" and n_sub_topic == 1:
            out[-1] = (x, ["만화[漢畵]".replace("漢", "漫"), "만화[漫畵]"])
            out.append(("/mods/subject[@ID]", ["KSH1998022212", "KSH1998022212"])); out.append(("/mods/subject[@authority]", [SUBJ_AUTH, SUBJ_AUTH]))
            out.append(("/mods/subject/", [None, None])); out.append(("/mods/subject/genre", ["웹툰[webtoon]", "웹툰[webtoon]"]))
            out.append(("/mods/subject[@ID]", ["KSH2016000049", "KSH2016000049"])); out.append(("/mods/subject[@authority]", [SUBJ_AUTH, SUBJ_AUTH]))
        if x == "/mods/identifier[@type]":
            n_ident_type += 1
            if n_ident_type == 1:
                out.append(("/mods/identifier", ["G903+NL-TEST-0001", "G903+NL-TEST-0002"])); out.append(("/mods/identifier[@type]", ["uci", "uci"]))
    # 다른이름(첫 저자 묶음의 기존 열에 값)
    for k, (x, vals) in enumerate(out):
        if x == "/mods/name/alternativeName[@type]":
            out[k] = (x, ["nickname", "nickname"])
        if x == "/mods/name/alternativeName/namePart":
            out[k] = (x, ["황재(시험 다른이름)", "황재(시험 다른이름)"])
    return out


def write(cols):
    shutil.copy(SRC, XLSX)
    wb = openpyxl.load_workbook(XLSX); ws = wb["Contents"]
    ws.delete_rows(1, ws.max_row)
    ws.append([x for x, _ in cols])
    for r in range(2):
        ws.append([vals[r] for _, vals in cols])
    wb.save(XLSX)
    return len(cols)


def analyze(xml: str) -> dict:
    import xml.etree.ElementTree as ET
    ns = {"mods": "http://www.loc.gov/mods/v3"}
    root = ET.fromstring(xml)
    names = [{"name": (n.findtext("mods:namePart", "", ns) or ""), "ID": n.get("ID"), "authority": n.get("authority"), "type": n.get("type"), "displayForm": n.findtext("mods:displayForm", "", ns),
              "alt": [(a.get("altType") or a.get("type"), (a.findtext("mods:namePart", "", ns) or "")) for a in n.findall("mods:alternativeName", ns)]} for n in root.findall("mods:name", ns)]
    subjects = [{"ID": s.get("ID"), "authority": s.get("authority"), "children": [(c.tag.split("}")[1], c.text) for c in s]} for s in root.findall("mods:subject", ns)]
    idents = [(i.get("type"), i.text) for i in root.findall("mods:identifier", ns)]
    origins = [{"eventType": o.get("eventType"), "type": o.get("type"), "places": len(o.findall("mods:place", ns)), "placeTerms": [(p.get("type"), p.text) for p in o.findall("mods:place/mods:placeTerm", ns)],
                "publisher": [p.text for p in o.findall("mods:publisher", ns)], "dateIssued": o.findtext("mods:dateIssued", "", ns), "edition": o.findtext("mods:edition", "", ns)} for o in root.findall("mods:originInfo", ns)]
    notes = [(e.get("type"), e.text) for e in root.findall("mods:note", ns)]
    return {"names": names, "subjects": subjects, "identifiers": idents, "originInfos": origins, "notes": notes}


def main():
    h, data = load(SRC)
    from kolis_tool import authority
    kac = (authority.lookup("황재", log=lambda m: None)["candidates"] or [{"AC_CONTROL_NO": "KAC201418251"}])[0]["AC_CONTROL_NO"]   # 헤더 수용 시험용(취소할 접수)
    cols = build(h, data, kac)
    n = write(cols)
    print(f"시험 엑셀: {XLSX} ({n}열, 2행). 끼운 열: name[@ID]={kac}, name[@authority], alternativeName(+[@type]), subject[@ID]/[@authority]×2 + subject/genre, identifier(uci), 출처정보 2묶음(각 발행지 부모 열 + 2쌍)")
    year = str(datetime.date.today().year)
    log = print
    c = kolis_http.Client(log)
    try:
        r = req.import_excel(c, year, NOTE, XLSX, 2, "YES", log)
        receipt = r["receipt"]
        print("접수번호", receipt, "건수", r["count"], "메시지", r.get("message"))
        kolis_flow.ledger_put(Path("work"), {"처리": kolis_flow.CANCEL, "작품": "원시인 삼촌(헤더 시험)", "접수번호": receipt, "가원부번호": "", "건수": r["count"],
                                           "콘텐츠ID": f"{r['ids'][0]} ~ {r['ids'][-1]}", "비고": NOTE, "결과": "반입 헤더 시험(전거·다른이름·주제명·UCI·출처정보 2묶음). 원문 등록 안 함 → 취소 요청", "반입 시각": f"{datetime.datetime.now():%Y-%m-%d %H:%M}", "실행 기록": "tools/import_probe.py"})
        cont = req.contents(c, year, receipt, log)
        results = []
        for it in cont["items"]:
            (OUT / f"{it['id']}.xml").write_text(it["mods"], encoding="utf-8")
            a = analyze(it["mods"]) if it["mods"] else {"error": "MODS 없음"}
            results.append({"id": it["id"], "vol": it["vol"], **a})
            print(f"\n== {it['id']} 권 {it['vol']}"); print(json.dumps(a, ensure_ascii=False, indent=1))
        (OUT / f"probe_result_{receipt}.json").write_text(json.dumps({"receipt": receipt, "message": r.get("message"), "columns": [x for x, _ in cols], "results": results}, ensure_ascii=False, indent=1), encoding="utf-8")
    finally:
        c.close()


if __name__ == "__main__":
    main()
